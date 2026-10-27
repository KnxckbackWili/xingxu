"""Real generated evidence, project-local weeks, source ACLs and atomic failure."""
import io,hashlib
from zipfile import ZipFile
from datetime import datetime,date
from uuid import uuid4
import pytest
from openpyxl import load_workbook
from docx import Document
from app import config,models
from app.db import SessionLocal
from app.services import document_deadlines as deadlines
from app.services import report_forms
from .test_project_documents import create,view,upload
from .test_confidential_documents_t26 import setup_project,policy


NOW=datetime(2025,12,28,11,30) # Auckland is already Monday 2026-W01, UTC still 2025-W52.


def weekly(client,headers,monkeypatch,pid=None):
    monkeypatch.setattr(deadlines,'utcnow',lambda:NOW)
    if pid is None:pid=create(client,headers,start_date='2025-01-01',timezone='Pacific/Auckland')['id']
    with SessionLocal() as db:
        project=db.get(models.Project,pid);project.start_date=date(2025,1,1);project.timezone='Pacific/Auckland'
        rule=db.query(models.ProjectDocumentDefinition).filter_by(project_id=pid,code='6.2').one();rule.created_at=NOW
        db.flush();deadlines.generate_weekly(db,project,now=NOW,actor_id=project.manager);db.commit()
    document=next(row for row in view(client,headers,pid)['documents'] if row['code']=='6.2' and row['period_key']=='2026-W01')
    return pid,document


def event(pid,day,title,attachment_id=None):
    with SessionLocal() as db:
        unit=db.query(models.VehicleUnit).filter_by(project_id=pid).first()
        if unit is None:
            unit=models.VehicleUnit(project_id=pid,asset_code='REAL-AGV-01',fleet_no='7',model_name='AGV真实型号');db.add(unit);db.flush()
        row=models.VehicleLifecycleEvent(project_id=pid,vehicle_unit_id=unit.id,event_type='fault',event_date=day,title=title,
            description=title+'真实原描述',action_taken='原解决方法',handler='原记录人',attachment_id=attachment_id,created_at=NOW)
        db.add(row);db.commit();return row.id


def generate(client,headers,pid,did,kind,**body):
    return client.post(f'/api/projects/{pid}/documents/{did}/generate-{kind}',headers=headers,json=body)


def workbook(client,headers,document):
    response=client.get('/api'+document['latest_version']['download_url'],headers=headers);assert response.status_code==200,response.text
    assert hashlib.sha256(response.content).hexdigest()==document['latest_version']['sha256']
    return load_workbook(io.BytesIO(response.content))


def test_current_week_excel_uses_site_iso_year_exact_original_dates_and_headers(client,pm_headers,admin_headers,monkeypatch):
    pid,doc=weekly(client,pm_headers,monkeypatch)
    before=event(pid,date(2025,12,28),'上周原问题');first=event(pid,date(2025,12,29),'=真实周一记录');last=event(pid,date(2026,1,4),'周日原记录');event(pid,date(2026,1,5),'下周原问题')
    result=generate(client,pm_headers,pid,doc['id'],'lifecycle');assert result.status_code==200,result.text
    result=result.json();ws=workbook(client,pm_headers,result).active
    assert '2026-W01' in ws['A1'].value
    assert [c.value for c in ws[2]]==report_forms.LIFECYCLE_HEADERS
    assert [ws.cell(row,6).value for row in (3,4)]==['=真实周一记录','周日原记录'] and ws['F3'].data_type=='s'
    assert [ws.cell(row,10).value for row in (3,4)]==['2025-12-29','2026-01-04'] and ws.max_row==4
    metadata=result['latest_version']['review_metadata']['generated']
    assert metadata['period_key']=='2026-W01' and metadata['source_event_ids']==[first,last]
    reviewed=client.post(f'/api/projects/{pid}/documents/{doc["id"]}/versions/{result["latest_version"]["id"]}/review',headers=admin_headers,json={'decision':'approved','note':'核对系统生成来源'})
    assert reviewed.status_code==200,reviewed.text
    assert reviewed.json()['review_metadata']['generated']['source_event_ids']==[first,last]
    with SessionLocal() as db:assert db.get(models.VehicleLifecycleEvent,before).event_date==date(2025,12,28)
    newer=generate(client,pm_headers,pid,doc['id'],'lifecycle');assert newer.status_code==200 and len(newer.json()['versions'])==2


def test_public_output_omits_even_owner_readable_private_sources_without_names_or_counts(client,pm_headers,worker_headers,monkeypatch):
    pid,users=setup_project(client,pm_headers);pid,doc=weekly(client,pm_headers,monkeypatch,pid)
    original=upload(client,pm_headers,pid,data=b'private source bytes').json()['file']
    assert policy(client,pm_headers,pid,'file',original,[users['worker1']]).status_code==200
    public_id=event(pid,date(2025,12,29),'普通公开记录')
    secret_id=event(pid,date(2025,12,30),'SECRET隐藏来源',original['id']);event(pid,date(2025,12,31),'UNKNOWN来源','missing-file')
    response=generate(client,pm_headers,pid,doc['id'],'lifecycle');assert response.status_code==200,response.text
    data=response.json();ws=workbook(client,worker_headers,data).active
    assert ws.max_row==3 and 'SECRET' not in str(list(ws.values)) and 'UNKNOWN' not in str(list(ws.values))
    source=data['latest_version']['review_metadata']['generated'];assert source['source_event_ids']==[public_id] and secret_id not in str(source)
    assert not any('excluded' in key or 'hidden' in key for key in source)
    narrow=generate(client,pm_headers,pid,doc['id'],'lifecycle',confidential=True,confidential_user_ids=[users['worker1']])
    assert narrow.status_code==200,narrow.text
    assert 'SECRET隐藏来源' in str(list(workbook(client,worker_headers,narrow.json()).active.values))
    other=client.post('/api/auth/login',json={'username':'worker2','password':'worker123'}).json()['access_token']
    other_headers={'Authorization':'Bearer '+other}
    assert client.get('/api'+narrow.json()['latest_version']['download_url'],headers=other_headers).status_code==403
    old=client.get(f'/api/projects/{pid}/export/lifecycle',headers=other_headers)
    assert old.status_code==200 and old.headers['cache-control']=='private, no-store'
    assert 'SECRET' not in str(list(load_workbook(io.BytesIO(old.content)).active.values))
    owner_copy=client.get(f'/api/projects/{pid}/export/lifecycle',headers=pm_headers)
    assert owner_copy.status_code==200 and 'SECRET' in str(list(load_workbook(io.BytesIO(owner_copy.content)).active.values))
    broad=generate(client,pm_headers,pid,doc['id'],'lifecycle',confidential=True,confidential_user_ids=[users['worker2']])
    assert broad.status_code==200 and 'SECRET' not in str(list(workbook(client,pm_headers,broad.json()).active.values))
    assert policy(client,pm_headers,pid,'document',broad.json(),[]).status_code==200
    latent=generate(client,pm_headers,pid,doc['id'],'lifecycle',confidential=True,confidential_user_ids=[users['worker2']])
    assert latent.status_code==200 and 'SECRET' not in str(list(workbook(client,pm_headers,latent.json()).active.values))


def test_historical_week_and_wrong_requirements_are_rejected_without_new_evidence(client,pm_headers,monkeypatch):
    pid,doc=weekly(client,pm_headers,monkeypatch)
    with SessionLocal() as db:
        db.get(models.ProjectDocument,doc['id']).period_key='2025-W52';db.commit()
    assert generate(client,pm_headers,pid,doc['id'],'lifecycle').status_code==409
    other=next(row for row in view(client,pm_headers,pid)['documents'] if row['code']=='4.1')
    assert generate(client,pm_headers,pid,other['id'],'lifecycle').status_code==400
    assert generate(client,pm_headers,pid,other['id'],'contact',application_id='none').status_code==400
    with SessionLocal() as db:assert db.query(models.DocumentVersion).count()==0 and db.query(models.Attachment).filter_by(project_id=pid).count()==0


def test_generated_contact_preserves_real_application_and_blocks_private_regeneration_to_public(client,pm_headers,worker_headers):
    pid,users=setup_project(client,pm_headers)
    doc=next(row for row in view(client,pm_headers,pid)['documents'] if row['code']=='6.9.3')
    created=client.post(f'/api/projects/{pid}/item-applications',headers=pm_headers,json={'title':'现场吊装申请','item_name':'吊车',
        'requirement':'原始吊装需求','reason':'原始申请原因','amount':'3000元','applicant':'原申请人','apply_date':'2026-02-03'})
    assert created.status_code==200,created.text
    aid=created.json()['id'];base=f'/api/projects/{pid}/documents/{doc["id"]}'
    assert client.get(base+'/generation-sources',headers=pm_headers).json()[0]['id']==aid
    result=generate(client,pm_headers,pid,doc['id'],'contact',application_id=aid,confidential=True,confidential_user_ids=[users['worker1']])
    assert result.status_code==200,result.text
    download=client.get('/api'+result.json()['latest_version']['download_url'],headers=worker_headers);assert download.status_code==200
    content='\n'.join(p.text for p in Document(io.BytesIO(download.content)).paragraphs)
    assert '需求：原始吊装需求' in content and '申请日期：2026-02-03' in content
    with ZipFile(io.BytesIO(download.content)) as package:assert b'<w:pBdr' not in package.read('word/styles.xml')
    assert result.json()['latest_version']['review_metadata']['generated']['source_application_ids']==[aid]
    assert client.get(base+'/generation-sources',headers=pm_headers).json()==[]
    assert generate(client,pm_headers,pid,doc['id'],'contact',application_id=aid).status_code==404
    assert client.get(base+'/generation-sources?confidential=true&confidential_user_ids='+users['worker1'],headers=pm_headers).json()[0]['id']==aid
    assert client.get(base+'/generation-sources?confidential=true&confidential_user_ids='+users['worker2'],headers=pm_headers).json()==[]
    denied_token=client.post('/api/auth/login',json={'username':'worker2','password':'worker123'}).json()['access_token']
    assert client.get(f'/api/projects/{pid}/item-applications/{aid}/export',headers={'Authorization':'Bearer '+denied_token}).status_code==403
    assert client.get(f'/api/projects/{pid}/item-applications/{aid}/export',headers=pm_headers).status_code==200
    with SessionLocal() as db:assert db.query(models.DocumentVersion).filter_by(project_document_id=doc['id']).count()==1


def test_cross_project_application_and_unauthorized_target_do_not_create_files(client,pm_headers,worker_headers):
    pid,users=setup_project(client,pm_headers);other=create(client,pm_headers)['id']
    app=client.post(f'/api/projects/{other}/item-applications',headers=pm_headers,json={'title':'其它项目真实申请'}).json()
    doc=next(row for row in view(client,pm_headers,pid)['documents'] if row['code']=='6.9.3')
    assert generate(client,pm_headers,pid,doc['id'],'contact',application_id=app['id']).status_code==404
    assert policy(client,pm_headers,pid,'document',doc,[]).status_code==200
    assert generate(client,worker_headers,pid,doc['id'],'contact',application_id=app['id']).status_code==403
    with SessionLocal() as db:assert db.query(models.Attachment).filter_by(project_id=pid).count()==0


def test_unknown_old_private_contact_origin_cannot_be_guessed_as_a_public_source(client,pm_headers):
    pid=create(client,pm_headers)['id'];document=next(row for row in view(client,pm_headers,pid)['documents'] if row['code']=='6.9.3')
    app=client.post(f'/api/projects/{pid}/item-applications',headers=pm_headers,json={'title':'不能凭名称推断公开范围'}).json()
    uploaded=upload(client,pm_headers,pid,[document['id']],data=b'old private contact with no reliable application id').json()
    assert policy(client,pm_headers,pid,'file',uploaded['file'],[]).status_code==200
    base=f'/api/projects/{pid}/documents/{document["id"]}'
    assert client.get(base+'/generation-sources',headers=pm_headers).json()==[]
    assert generate(client,pm_headers,pid,document['id'],'contact',application_id=app['id']).status_code==404


def test_generation_audit_failure_rolls_back_real_file_versions_status_and_source(client,pm_headers,monkeypatch):
    from app.services import generated_document_forms as forms
    pid,doc=weekly(client,pm_headers,monkeypatch);identity=event(pid,date(2025,12,29),'不可改写的原问题')
    before=set(config.UPLOAD_DIR.iterdir())
    def fail(*args,**kwargs):raise RuntimeError('generation audit offline')
    monkeypatch.setattr(forms,'log_change',fail)
    with pytest.raises(RuntimeError,match='generation audit offline'):generate(client,pm_headers,pid,doc['id'],'lifecycle')
    with SessionLocal() as db:
        assert db.query(models.DocumentVersion).count()==0 and db.query(models.Attachment).filter_by(project_id=pid).count()==0
        assert db.get(models.ProjectDocument,doc['id']).status=='pending' and db.get(models.VehicleLifecycleEvent,identity).title=='不可改写的原问题'
    assert set(config.UPLOAD_DIR.iterdir())==before


def test_archived_or_explicitly_denied_generation_creates_no_evidence(client,pm_headers,monkeypatch):
    pid,document=weekly(client,pm_headers,monkeypatch)
    with SessionLocal() as db:db.get(models.User,'user-pm-001').permission_overrides={'worker.deliver':False};db.commit()
    assert generate(client,pm_headers,pid,document['id'],'lifecycle').status_code==403
    with SessionLocal() as db:
        db.get(models.User,'user-pm-001').permission_overrides={};db.get(models.Project,pid).status='archived';db.commit()
    assert generate(client,pm_headers,pid,document['id'],'lifecycle').status_code==400
    with SessionLocal() as db:assert db.query(models.DocumentVersion).count()==0 and db.query(models.Attachment).filter_by(project_id=pid).count()==0


def test_inventory_shared_generation_preserves_original_movement_time_and_source(client,pm_headers):
    pid=create(client,pm_headers)['id']
    assert client.put(f'/api/projects/{pid}/events/internal_kickoff',headers=pm_headers,json={'actual_at':'2026-01-02T01:00:00Z'}).status_code==200
    document=next(row for row in view(client,pm_headers,pid)['documents'] if row['code']=='9.2')
    with SessionLocal() as db:
        item=models.ProjectInventoryItem(project_id=pid,name='真实工具',quantity=2,unit='个');db.add(item);db.flush()
        move=models.ProjectInventoryMovement(project_id=pid,item_id=item.id,movement_type='in',quantity=2,balance_after=2,
            occurred_at=datetime(2026,1,2,3,4,5),time_basis='utc');db.add(move)
        db.add(models.ProjectInventoryMovement(project_id=pid,item_id=item.id,movement_type='out',quantity=1,balance_after=1,
            occurred_at=datetime(2026,1,2,18,30),time_basis='legacy_local'))
        db.commit();mid=move.id
    result=generate(client,pm_headers,pid,document['id'],'inventory');assert result.status_code==200,result.text
    book=workbook(client,pm_headers,result.json());assert book['收发流水']['A2'].value=='2026-01-02T03:04:05+00:00'
    assert book['收发流水']['A3'].value=='2026-01-02T18:30:00' and book['收发流水']['B3'].value=='legacy_local'
    assert mid in result.json()['latest_version']['review_metadata']['generated']['source_movement_ids']
