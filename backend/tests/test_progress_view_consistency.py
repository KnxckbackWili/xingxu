"""管理视图与导出采用相同的当前进度依据。"""




def _minimal_import():
    return {'type': 'agvpm-project', 'project': {'name': 'round6 import', 'start_date': '2026-01-01', 'end_date': '2026-12-31'},
            'modules': [{'name': 'module', 'phase_name': 'phase', 'tasks': [{'name': 'task', 'quantity_type': 'fixed', 'target_quantity': 1}]}]}


def test_delivery_overview_respects_safety_gate(client, pm_headers):
    from app.db import SessionLocal
    from app.models import Phase, Module, Task, CheckItem, ScheduleEntry, DailyProgress, Project
    from datetime import date
    from openpyxl import load_workbook
    from io import BytesIO
    created = client.post('/api/projects', headers=pm_headers, json={
        'name': 'Export safety gate', 'start_date': '2026-01-01', 'end_date': '2026-01-02', 'new_vehicles': False}).json()
    uid = client.get('/api/auth/me', headers=pm_headers).json()['id']
    pid = created['id']
    with SessionLocal() as db:
        phase = Phase(project_id=pid, name='Safety phase'); db.add(phase); db.flush()
        module = Module(project_id=pid, phase_id=phase.id, name='Safety module'); db.add(module); db.flush()
        task = Task(project_id=pid, phase_id=phase.id, module_id=module.id, name='Safety task', quantity_type='fixed', target_quantity=1, manually_completed=True)
        db.add(task); db.flush()
        db.add(CheckItem(project_id=pid, task_id=task.id, name='Safety failed', is_safety_critical=True, is_done=False, failed=True))
        entry = ScheduleEntry(project_id=pid, task_id=task.id, work_date=date(2026,1,1), planned_quantity=1)
        db.add(entry); db.flush()
        db.add(DailyProgress(project_id=pid, task_id=task.id, schedule_entry_id=entry.id, user_id=uid, work_date=entry.work_date, completed_quantity=1, client_uuid='round6-excel-safety'))
        db.commit()
    from tests.lifecycle_helpers import complete_vehicle_progress
    from app.services.progress import project_progress_tree, project_completion_blockers
    with SessionLocal() as db:
        db.get(Project, pid).status = 'executing'; db.commit()
    complete_vehicle_progress(client, pm_headers, pid)
    with SessionLocal() as db:
        expected = project_progress_tree(db, pid)['percent']
        assert 0 <= expected < 100
        assert project_completion_blockers(db, pid)
    response = client.get(f'/api/projects/{pid}/export/excel', headers=pm_headers)
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content))
    values = {row[1]:row[2] for row in workbook['交付概览'].iter_rows(values_only=True) if row[0]=='项目概况'}
    assert values['整体进度 %'] == expected
    evidence = client.get(f'/api/projects/{pid}/export/json', headers=pm_headers).json()['execution_evidence']
    assert evidence['read_only'] and any(row['check_items'] and row['check_items'][0]['failed'] for row in evidence['tasks'])


def test_excel_overview_uses_same_weighted_progress_as_project(client, pm_headers):
    from io import BytesIO
    from openpyxl import load_workbook
    from app.db import SessionLocal
    from app.services.progress import project_progress_tree
    payload = _minimal_import()
    payload['modules'][0]['tasks'] = [
        {'name': 'large', 'quantity_type': 'fixed', 'target_quantity': 100},
        {'name': 'small', 'quantity_type': 'fixed', 'target_quantity': 1}]
    payload['progress'] = [{'task_name': 'small', 'work_date': '2026-01-01', 'completed_quantity': 1}]
    pid = client.post('/api/projects/import/json', headers=pm_headers, json=payload).json()['id']
    with SessionLocal() as db: expected = project_progress_tree(db,pid)['percent']
    workbook = load_workbook(BytesIO(client.get(f'/api/projects/{pid}/export/excel', headers=pm_headers).content))
    values = {row[1]: row[2] for row in workbook['交付概览'].iter_rows(values_only=True) if row[0]=='项目概况'}
    assert values['整体进度 %'] == expected
    assert expected != 50
