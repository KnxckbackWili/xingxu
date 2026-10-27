"""Project handover packages preserve closure evidence without replaying authority."""
import hashlib
import io
import json
import zipfile
from datetime import datetime, timedelta
from html.parser import HTMLParser
from urllib.parse import unquote

import pytest


def _new_project(client, headers, name="归档证据测试"):
    response = client.post("/api/projects", headers=headers, json={
        "name": name, "start_date": "2026-08-01", "end_date": "2026-12-31", "new_vehicles": False,
    })
    assert response.status_code == 200, response.text
    return response.json()["id"]


def _archived_project(client, headers):
    from app.db import SessionLocal
    from app.models import Project, AcceptanceRecord, ArchiveChecklist
    from app.services.acceptance import acceptance_scope_fingerprint
    from app.time_utils import project_today

    pid = _new_project(client, headers, '归档 <script>alert("project")</script>')
    attachments = []
    for content in (b"%PDF-1.4\nQA acceptance one\n%%EOF", b"%PDF-1.4\nQA acceptance two\n%%EOF"):
        response = client.post(f"/api/projects/{pid}/documents/files", headers=headers,
                               files={"file": ("验收#资料.pdf", content, "application/pdf")})
        assert response.status_code == 200, response.text
        attachments.append((response.json()['file']['id'], content))
    content = '完整复盘\n<script>alert("retro")</script>\n' + "长原文" * 1000
    today = project_today("Asia/Shanghai")
    with SessionLocal() as db:
        project = db.get(Project, pid)
        project.status = "archived"
        project.retrospective = {"content": content, "author": "user-pm-001", "written_at": "2026-10-04T21:58:00"}
        fingerprint = acceptance_scope_fingerprint(db, pid)
        db.add_all([
            AcceptanceRecord(id="archive-old", project_id=pid, accepted_by_name="旧客户代表",
                             accepted_at=today - timedelta(days=1), conclusion="rejected", notes="旧整改意见",
                             scope_fingerprint="old-scope", created_by="user-pm-001", created_at=datetime(2026, 10, 4, 12)),
            AcceptanceRecord(id="archive-current", project_id=pid, accepted_by_name='客户 <img src=x onerror="bad()">',
                             accepted_at=today, conclusion="accepted", notes='终稿通过\n<script>alert("notes")</script>',
                             scope_fingerprint=fingerprint, attachment_id=attachments[0][0],
                             created_by="user-pm-001", created_at=datetime(2026, 10, 4, 21, 52, 0, 123456)),
            AcceptanceRecord(id="archive-future", project_id=pid, accepted_by_name="历史未来日期",
                             accepted_at=today + timedelta(days=1), conclusion="accepted", scope_fingerprint=fingerprint,
                             created_at=datetime(2026, 10, 4, 21, 53)),
            ArchiveChecklist(id="archive-check", project_id=pid, category="验收文件", item_name="终验报告",
                             is_required=True, is_done=True, done_by="user-pm-001", done_at=datetime(2026, 10, 4, 21, 54),
                             attachment_id=attachments[1][0], requirements={"doc_type": "PDF", "version_requirement": "终稿",
                             "signoff_required": True, "notes": '<svg onload="bad()">', "unknown_field": "原扩展要求"}),
        ])
        db.commit()
    return pid, content, attachments


def _package(client, headers, pid):
    response = client.get(f"/api/projects/{pid}/export/package", headers=headers)
    assert response.status_code == 200, response.text
    archive = zipfile.ZipFile(io.BytesIO(response.content))
    return archive, json.loads(archive.read("归档记录.json")), archive.read("归档记录.html").decode("utf-8")


def test_package_preserves_full_closure_history_original_times_and_manifest_links(client, pm_headers):
    from app.db import SessionLocal
    from app.models import Project, AcceptanceRecord
    from openpyxl import load_workbook

    pid, content, attachments = _archived_project(client, pm_headers)
    other = _new_project(client, pm_headers, "独立无关项目")
    with SessionLocal() as db:
        db.add(AcceptanceRecord(project_id=other, accepted_by_name="不可混入的其他项目",
                               accepted_at=datetime(2026, 1, 1).date(), notes="私有其他项目签字"))
        db.commit()
    archive, evidence, html = _package(client, pm_headers, pid)
    manifest = json.loads(archive.read("manifest.json"))
    assert manifest["archive_records"] == {"json": "归档记录.json", "html": "归档记录.html", "read_only": True}
    assert evidence["read_only"] is True and evidence["import_supported"] is False
    assert evidence["generated_at"] == manifest["generated_at"]
    assert evidence["retrospective"]["content"] == content
    assert evidence["retrospective"]["author_name"] == "项目经理"
    assert evidence["retrospective"]["written_at"] == "2026-10-04T21:58:00"
    assert evidence["retrospective"]["written_at_local"] == "2026-10-05T05:58:00+08:00"
    assert evidence["current_acceptance_id"] == "archive-current"
    rows = {row["id"]: row for row in evidence["acceptance_records"]}
    assert set(rows) == {"archive-current", "archive-old", "archive-future"}
    assert rows["archive-current"]["scope_current"] is True
    assert rows["archive-old"]["scope_current"] is False and rows["archive-old"]["scope_verified"] is True
    assert rows["archive-future"]["effective_date_reached"] is False and rows["archive-future"]["is_effective"] is False
    files = {item["id"]: item for item in manifest["files"]}
    assert rows["archive-current"]["attachment"] == files[attachments[0][0]]
    check = next(row for row in evidence["archive_checklist"] if row["id"] == "archive-check")
    assert check["attachment"] == files[attachments[1][0]]
    assert check["done_at_local"] == "2026-10-05T05:54:00+08:00"
    for attachment_id, original in attachments:
        item = files[attachment_id]
        assert archive.read(item["archive_path"]) == original
        assert item["sha256"] == hashlib.sha256(original).hexdigest()
    assert "私有其他项目签字" not in html
    assert "password_hash" not in json.dumps(evidence)
    assert "范围已变更" in html and "日期尚未生效" in html
    assert "暂无验收记录" not in html
    # Portable planning remains separate from closure evidence; T27 uses six delivery sheets.
    snapshot = json.loads(archive.read("项目备份.json"))
    assert "retrospective" not in snapshot["project"] and "acceptance_records" not in snapshot
    assert load_workbook(io.BytesIO(archive.read("项目数据.xlsx"))).sheetnames == ["交付概览", "生命周期表", "问题清单", "物资清单", "收发流水", "交付资料"]
    with SessionLocal() as db:
        assert db.get(Project, pid).retrospective["content"] == content


class _ParsedHtml(HTMLParser):
    def __init__(self):
        super().__init__()
        self.tags = []
        self.links = []
        self.events = []
        self.texts = []
        self.visible_texts = []
        self.details_depth = 0
        self.open_details = []

    def handle_starttag(self, tag, attrs):
        self.tags.append(tag)
        self.events.extend(key for key, _ in attrs if key.lower().startswith("on"))
        if tag == "details":
            self.details_depth += 1
            if "open" in dict(attrs):
                self.open_details.append(tag)
        if tag == "a":
            self.links.append(dict(attrs)["href"])

    def handle_data(self, data):
        self.texts.append(data)
        if self.details_depth == 0:
            self.visible_texts.append(data)

    def handle_endtag(self, tag):
        if tag == "details":
            self.details_depth -= 1


def test_package_html_escapes_all_record_text_and_encodes_relative_file_links(client, pm_headers):
    pid, content, _ = _archived_project(client, pm_headers)
    archive, evidence, html = _package(client, pm_headers, pid)
    parser = _ParsedHtml()
    parser.feed(html)
    assert "script" not in parser.tags and "img" not in parser.tags and "svg" not in parser.tags
    assert parser.events == []
    assert content in "".join(parser.texts)
    assert "<meta charset=\"UTF-8\">" in html and "default-src 'none'" in html
    assert "white-space:pre-wrap" in html and "overflow-wrap:anywhere" in html
    file_links = [link for link in parser.links if link.startswith("files/")]
    assert file_links and all("%" in link and "#" not in link for link in file_links)
    assert all(unquote(link) in archive.namelist() for link in parser.links)
    assert "归档记录.json" in {unquote(link) for link in parser.links}
    assert "新项目导入及克隆不会继承" in html
    visible = "".join(parser.visible_texts)
    assert "2026-10-05 05:58:00" in visible and "2026-10-05 05:52:00" in visible
    assert "文件格式：PDF" in visible and "版本要求：终稿" in visible and "客户签字：需要" in visible
    assert "补充说明：" in visible and '<svg onload="bad()">' in visible
    assert "原扩展要求" not in visible and "原扩展要求" in html
    assert "archive-current" not in visible and "archive-current" in html
    assert "SHA-256" not in visible and "SHA-256" in html
    assert "2026-10-04T21:52:00.123456" not in visible and "2026-10-04T21:52:00.123456" in html
    assert "原始记录与校验信息" in html and parser.open_details == []
    from app.services.import_export import _archive_evidence_html
    evidence["archive_checklist"] = [next(row for row in evidence["archive_checklist"] if row["id"] == "archive-check")]
    evidence["archive_checklist"][0]["requirements"] = {}
    empty = _ParsedHtml()
    empty.feed(_archive_evidence_html(evidence))
    assert "无额外要求" in "".join(empty.visible_texts)


def test_missing_and_unrelated_attachment_refs_match_manifest_without_fake_links(client, pm_headers):
    from app.config import UPLOAD_DIR
    from app.db import SessionLocal
    from app.models import Attachment, ArchiveChecklist

    pid, _, attachments = _archived_project(client, pm_headers)
    with SessionLocal() as db:
        stored = db.get(Attachment, attachments[0][0]).stored_name
        db.get(ArchiveChecklist, "archive-check").attachment_id = "outside-project-attachment"
        db.commit()
    (UPLOAD_DIR / stored).unlink()
    archive, evidence, html = _package(client, pm_headers, pid)
    manifest = json.loads(archive.read("manifest.json"))
    missing = next(row for row in manifest["files"] if row["id"] == attachments[0][0])
    current = next(row for row in evidence["acceptance_records"] if row["is_effective"])
    assert current["attachment"] == missing and missing["missing"] is True
    assert missing["archive_path"] not in archive.namelist()
    check = next(row for row in evidence["archive_checklist"] if row["id"] == "archive-check")
    assert check["attachment"] == {"id": "outside-project-attachment", "missing": True, "reason": "attachment_not_in_project_manifest"}
    assert "附件缺失，未包含文件内容" in html
    parser = _ParsedHtml()
    parser.feed(html)
    assert not any(link.startswith("files/") for link in parser.links)


@pytest.mark.parametrize("retrospective,expected", [
    ("旧版原文\n第二行", "旧版原文\n第二行"),
    ({"success": "成功经验", "problems": "问题教训", "improvements": "改进建议", "author": "missing-author", "written_at": "无法解析的历史时间"}, "成功经验"),
    ({}, "尚未填写项目复盘"),
])
def test_package_handles_legacy_or_absent_retrospective_and_no_acceptance(client, pm_headers, retrospective, expected):
    from app.db import SessionLocal
    from app.models import Project

    pid = _new_project(client, pm_headers)
    with SessionLocal() as db:
        project = db.get(Project, pid)
        project.retrospective = retrospective
        project.timezone = "old-unknown-zone"
        db.commit()
    _, evidence, html = _package(client, pm_headers, pid)
    assert expected in html
    assert evidence["current_acceptance_id"] is None and evidence["acceptance_records"] == []
    assert evidence["display_timezone"] == "Asia/Shanghai"
    assert evidence["retrospective"]["author_name"] == ""
    assert "原作者已不存在或未记录" in html and "暂无验收记录" in html
    if isinstance(retrospective, dict) and retrospective.get("author"):
        assert evidence["retrospective"]["written_at"] == "无法解析的历史时间"
        assert evidence["retrospective"]["written_at_local"] == ""
        parser = _ParsedHtml()
        parser.feed(html)
        assert "missing-author" not in "".join(parser.visible_texts) and "missing-author" in html
    with SessionLocal() as db:
        assert db.get(Project, pid).retrospective == retrospective


def test_export_sidecar_cannot_confer_acceptance_or_retrospective_on_import_or_clone(client, pm_headers):
    from app.db import SessionLocal
    from app.models import Project, AcceptanceRecord

    pid, _, _ = _archived_project(client, pm_headers)
    archive, evidence, _ = _package(client, pm_headers, pid)
    snapshot = json.loads(archive.read("项目备份.json"))
    snapshot["project"]["status"] = "draft"
    snapshot["project"]["name"] = "明确新草稿"
    # Even explicitly supplied archival evidence is outside the new-project allowlist.
    snapshot["archive_evidence"] = evidence
    response = client.post("/api/projects/import/json", headers=pm_headers, json=snapshot)
    assert response.status_code == 200, response.text
    imported = response.json()["id"]
    # Service API clone is independently fresh even when its source is archived.
    response = client.post(f"/api/projects/{pid}/clone", headers=pm_headers, json={
        "name": "明确新克隆", "start_date": "2026-08-01", "end_date": "2026-12-31",
    })
    assert response.status_code == 200, response.text
    cloned = response.json()["id"]
    with SessionLocal() as db:
        for new_id in (imported, cloned):
            assert db.get(Project, new_id).retrospective == {}
            assert db.query(AcceptanceRecord).filter_by(project_id=new_id).count() == 0


def test_record_added_at_service_read_has_source_fields_without_export_failure(client, pm_headers, monkeypatch):
    from app.db import SessionLocal
    from app.models import AcceptanceRecord
    from app.services import acceptance
    from app.time_utils import project_today

    pid = _new_project(client, pm_headers)
    original_evidence = acceptance.acceptance_evidence

    def insert_before_selected_read(db, project_id):
        assert project_id == pid
        with SessionLocal() as writer:
            writer.add(AcceptanceRecord(id="concurrently-recorded", project_id=pid,
                                       accepted_by_name="同步录入客户代表", accepted_at=project_today("Asia/Shanghai"),
                                       conclusion="accepted", notes="新增记录完整保留",
                                       created_by="user-pm-001", created_at=datetime(2026, 10, 4, 22)))
            writer.commit()
        return original_evidence(db, project_id)

    monkeypatch.setattr(acceptance, "acceptance_evidence", insert_before_selected_read)
    _, evidence, html = _package(client, pm_headers, pid)
    assert evidence["current_acceptance_id"] == "concurrently-recorded"
    assert evidence["acceptance_records"][0]["created_by"] == "user-pm-001"
    assert "同步录入客户代表" in html and "新增记录完整保留" in html
