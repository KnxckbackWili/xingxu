"""Real weekly plans are separate from actual work and original daily sources."""
from datetime import date
from app.db import SessionLocal
from app.models import Project, User, Task, WeeklyPlanCell, Notification, VehicleProgress
from tests.test_vehicle_progress import setup_fleet, write
from tests.test_work_items_t34 import phase, create as create_work


def setup(client, headers, *, vehicle=False):
    pid, units = setup_fleet(client, headers, 2)
    ph = phase(client, headers, pid)
    task = create_work(client, headers, pid, ph, **({'kind':'vehicle','vehicle_step_key':'arrival_inspection'} if vehicle else {}))
    return pid, units, ph, task


def create_cell(client, headers, pid, task, **values):
    response = client.post(f'/api/projects/{pid}/weekly-plans', headers=headers,
        json={'task_id':task['id'],'week_start':'2001-01-01',**values})
    assert response.status_code == 200, response.text
    return response.json()


def publish(client, headers, pid, *cells):
    return client.post(f'/api/projects/{pid}/weekly-plans/publish', headers=headers,
        json={'cells':[{'id':row['id'],'version':row['version']} for row in cells]})


def test_multi_people_scoped_publish_edit_cas_and_no_autoassign(client, pm_headers, worker_headers):
    pid, units, ph, task = setup(client, pm_headers)
    with SessionLocal() as db:
        worker = db.query(User).filter_by(username='worker1').one(); wid=worker.id
        manager=db.get(Project,pid).manager;p=db.get(Project,pid);p.members=[wid];db.commit()
    row=create_cell(client,pm_headers,pid,task,assignee_ids=[wid,manager],mode='finish')
    draft=create_cell(client,pm_headers,pid,task,week_start='2001-01-08')
    assert row['assignee_ids']==sorted([wid,manager]) and not row['published']
    assert client.get('/api/worker/weekly-plans?week_start=2001-01-01',headers=worker_headers).json()['cells']==[]
    assert publish(client,pm_headers,pid,row).status_code==200
    result=client.get('/api/worker/weekly-plans?week_start=2001-01-01&weeks=2',headers=worker_headers).json()
    assert [cell['id'] for cell in result['cells']]==[row['id']]
    url=f"/api/projects/{pid}/weekly-plans/{row['id']}"
    changed=client.put(url,headers=pm_headers,json={'version':1,'week_start':'2001-01-15','notes':'直接调整已发布计划'})
    assert changed.status_code==200 and changed.json()['published'],changed.text
    assert client.put(url,headers=pm_headers,json={'version':1,'notes':'迟到旧表单'}).status_code==409
    with SessionLocal() as db:
        assert not db.get(WeeklyPlanCell,draft['id']).published
        assert db.query(Notification).filter_by(user_id=wid,project_id=pid,ntype='schedule_confirmed').count()==2
        assert not db.get(Task,task['id']).manually_completed


def test_finish_and_in_progress_have_distinct_overdue_and_other_work_is_boolean(client,pm_headers,monkeypatch):
    pid, units, ph, task=setup(client,pm_headers)
    pending=create_cell(client,pm_headers,pid,task)
    finish=create_cell(client,pm_headers,pid,task,week_start='2001-01-08',mode='finish')
    assert publish(client,pm_headers,pid,pending,finish).status_code==200
    rows=client.get(f'/api/projects/{pid}/weekly-plans?week_start=2001-01-01&weeks=2',headers=pm_headers).json()['cells']
    by_id={row['id']:row for row in rows}
    assert not by_id[pending['id']]['completion']['overdue']
    assert by_id[finish['id']]['completion']['overdue']
    assert client.put(f"/api/projects/{pid}/work-items/{task['id']}/completion",headers=pm_headers,json={'version':0,'is_completed':True,'reason':'非车辆工作已完成'}).status_code==200
    rows=client.get(f'/api/projects/{pid}/weekly-plans?week_start=2001-01-01&weeks=2',headers=pm_headers).json()['cells']
    assert all(row['completion']['is_complete'] and not row['completion']['overdue'] for row in rows)
    from app.services import weekly_plans as service
    monkeypatch.setattr(service,'local_today',lambda project: date(2001,1,14))
    with SessionLocal() as db:
        db.get(Task,task['id']).manually_completed=False;db.commit()
    rows=client.get(f'/api/projects/{pid}/weekly-plans?week_start=2001-01-08&weeks=1',headers=pm_headers).json()['cells']
    assert not rows[0]['completion']['overdue']  # Sunday is part of the planned week.


def test_vehicle_finish_counts_are_cumulative_not_reused_for_two_weeks(client,pm_headers):
    pid, units, ph, task=setup(client,pm_headers,vehicle=True)
    first=create_cell(client,pm_headers,pid,task,mode='finish',target_count=1)
    second=create_cell(client,pm_headers,pid,task,mode='finish',target_count=1,week_start='2001-01-08')
    assert publish(client,pm_headers,pid,first,second).status_code==200
    assert write(client,pm_headers,pid,[units[0]]).status_code==200
    rows=client.get(f'/api/projects/{pid}/weekly-plans?week_start=2001-01-01&weeks=2',headers=pm_headers).json()['cells']
    by_id={row['id']:row for row in rows}
    assert by_id[first['id']]['completion']['required_count']==1 and by_id[first['id']]['completion']['is_complete']
    assert by_id[second['id']]['completion']['required_count']==2 and by_id[second['id']]['completion']['overdue']
    with SessionLocal() as db:assert db.query(VehicleProgress).filter_by(project_id=pid,is_done=True).count()==1


def test_explicit_claim_has_one_cas_winner_and_never_claims_somebody_elses_cell(client,pm_headers,worker_headers):
    pid, units, ph, task=setup(client,pm_headers)
    with SessionLocal() as db:
        person=db.query(User).filter_by(username='worker1').one();wid=person.id
        db.get(Project,pid).members=[wid];db.commit()
    row=create_cell(client,pm_headers,pid,task)
    assert publish(client,pm_headers,pid,row).status_code==200
    result=client.get('/api/worker/weekly-plans?week_start=2001-01-01',headers=worker_headers).json()
    assert not result['cells'] and [cell['id'] for cell in result['unassigned']]==[row['id']]
    url=f"/api/projects/{pid}/weekly-plans/{row['id']}/claim"
    accepted=client.post(url,headers=worker_headers,json={'version':1})
    assert accepted.status_code==200 and accepted.json()['assignee_ids']==[wid],accepted.text
    assert client.post(url,headers=pm_headers,json={'version':1}).status_code==409
    assert client.get('/api/worker/weekly-plans?week_start=2001-01-01',headers=worker_headers).json()['cells'][0]['id']==row['id']


def test_target_scope_and_publication_are_atomic_and_do_not_create_zero_goals(client,pm_headers):
    pid, units, ph, task=setup(client,pm_headers,vehicle=True)
    for values in [{'mode':'finish','target_count':0},{'mode':'finish','target_count':1.5},{'mode':'in_progress','target_count':1}]:
        assert client.post(f'/api/projects/{pid}/weekly-plans',headers=pm_headers,json={'task_id':task['id'],'week_start':'2001-01-01',**values}).status_code in (400,422)
    first=create_cell(client,pm_headers,pid,task)
    second=create_cell(client,pm_headers,pid,task,week_start='2001-01-08')
    bad={**second,'version':9}
    assert publish(client,pm_headers,pid,first,bad).status_code==409
    with SessionLocal() as db:assert not db.get(WeeklyPlanCell,first['id']).published and db.get(WeeklyPlanCell,first['id']).version==0
    other_pid, other_units, other_ph, other_task=setup(client,pm_headers,vehicle=True)
    response=client.post(f'/api/projects/{pid}/weekly-plans',headers=pm_headers,json={'task_id':task['id'],'week_start':'2001-01-15','scope':{'vehicle_unit_ids':other_units}})
    assert response.status_code==404


def test_holiday_missing_year_and_known_workday_are_explicit(client,pm_headers):
    pid, units, ph, task=setup(client,pm_headers)
    result=client.get(f'/api/projects/{pid}/weekly-plans?week_start=2027-01-04&weeks=1',headers=pm_headers)
    assert result.status_code==200,result.text
    assert 2027 in result.json()['weeks'][0]['missing_holiday_years']
    known=client.get(f'/api/projects/{pid}/weekly-plans?week_start=2026-02-09&weeks=1',headers=pm_headers).json()
    assert '2026-02-14' in known['weeks'][0]['makeup_workdays']


def test_scoped_scheduler_can_edit_settings_without_core_project_update_and_source_days_remain(client,pm_headers,worker_headers):
    from app.models import ScheduleEntry
    pid,units,ph,task=setup(client,pm_headers)
    with SessionLocal() as db:
        worker=db.query(User).filter_by(username='worker1').one();wid=worker.id
        p=db.get(Project,pid);p.members=[wid];p.member_permissions=[{'user_id':wid,'capabilities':['project.manage_schedule']}]
        old=ScheduleEntry(project_id=pid,task_id=task['id'],work_date=date(2001,1,1),assignee_id=wid,planned_quantity=0,is_confirmed=False)
        db.add(old);db.commit();oldid=old.id
    url=f'/api/projects/{pid}/weekly-plans/settings'
    current=client.get(url,headers=worker_headers)
    assert current.status_code==200,current.text
    version=current.json()['version']
    body={'version':version,'rest_mode':'single','rest_days':['2026-10-01'],'skip_holidays':True}
    changed=client.put(url,headers=worker_headers,json=body)
    assert changed.status_code==200 and changed.json()['settings']['rest_mode']=='one',changed.text
    assert client.put(url,headers=worker_headers,json=body).status_code==409
    assert client.put(f'/api/projects/{pid}',headers=pm_headers,json={'rest_mode':'two'}).status_code==400
    row=create_cell(client,worker_headers,pid,task)
    assert row['permissions']['edit']
    with SessionLocal() as db:
        person=db.get(User,wid);person.permission_overrides={'project.manage_schedule':False};person.permission_version+=1;db.commit()
    assert client.put(url,headers=worker_headers,json={**body,'version':version+1}).status_code==403
    with SessionLocal() as db:assert db.get(ScheduleEntry,oldid).assignee_id==wid


def test_draft_preview_is_readonly_exact_events_and_stale_scope_requires_repreview(client,pm_headers):
    from datetime import datetime
    from app.models import ProjectEvent
    pid,units,ph,task=setup(client,pm_headers,vehicle=True)
    with SessionLocal() as db:
        project=db.get(Project,pid);project.timezone='Pacific/Honolulu'
        db.add(ProjectEvent(project_id=pid,event_key='ship',planned_date=date(2001,1,15),actual_at=datetime(2001,1,8,5)))
        db.commit()
    body={'phase_rules':[{'phase_id':ph,'anchor':'before_ship','offset_weeks':0,'mode':'finish'}]}
    base=f'/api/projects/{pid}/weekly-plans/draft'
    preview=client.post(base+'/preview',headers=pm_headers,json=body)
    assert preview.status_code==200,preview.text
    item=preview.json()['proposals'][0]
    assert item['week_start']=='2001-01-01' and item['target_count']==2 and not item['assignee_ids']
    with SessionLocal() as db:assert db.query(WeeklyPlanCell).filter_by(project_id=pid).count()==0
    with SessionLocal() as db:
        event=db.query(ProjectEvent).filter_by(project_id=pid,event_key='ship').one();event.actual_at=datetime(2001,1,9,5);db.commit()
    assert client.post(base,headers=pm_headers,json={**body,'preview_token':preview.json()['preview_token']}).status_code==409
    fresh=client.post(base+'/preview',headers=pm_headers,json=body).json()
    assert fresh['proposals'][0]['week_start']=='2001-01-08'
    saved=client.post(base,headers=pm_headers,json={**body,'preview_token':fresh['preview_token']})
    assert saved.status_code==200 and saved.json()['count']==1,saved.text
    with SessionLocal() as db:
        row=db.query(WeeklyPlanCell).filter_by(project_id=pid).one();assert row.target_count==2 and not row.published and not row.assignee_ids
    assert not client.post(base+'/preview',headers=pm_headers,json=body).json()['proposals']


def test_missing_event_no_guessed_dates_and_no_zero_goal(client,pm_headers):
    pid,units,ph,task=setup(client,pm_headers)
    body={'phase_rules':[{'phase_id':ph,'anchor':'after_arrival','offset_weeks':0}]}
    preview=client.post(f'/api/projects/{pid}/weekly-plans/draft/preview',headers=pm_headers,json=body).json()
    assert not preview['proposals'] and preview['warnings'][0]['code']=='missing_anchor'
    result=client.post(f'/api/projects/{pid}/weekly-plans/draft',headers=pm_headers,json={**body,'preview_token':preview['preview_token']})
    assert result.status_code==200 and result.json()['count']==0


def test_cancelled_rows_are_reviewable_and_restore_requires_reason_and_fresh_version(client,pm_headers,worker_headers):
    pid,units,ph,task=setup(client,pm_headers)
    with SessionLocal() as db:
        worker=db.query(User).filter_by(username='worker1').one();p=db.get(Project,pid);p.members=[worker.id];db.commit()
    row=create_cell(client,pm_headers,pid,task)
    url=f"/api/projects/{pid}/weekly-plans/{row['id']}"
    assert client.put(url,headers=pm_headers,json={'version':0,'is_cancelled':True}).status_code==400
    cancelled=client.put(url,headers=pm_headers,json={'version':0,'is_cancelled':True,'reason':'本周范围取消'})
    assert cancelled.status_code==200,cancelled.text
    assert not client.get(f'/api/projects/{pid}/weekly-plans?week_start=2001-01-01',headers=pm_headers).json()['cells']
    assert client.get(f'/api/projects/{pid}/weekly-plans?week_start=2001-01-01&include_cancelled=true',headers=pm_headers).json()['cells'][0]['is_cancelled']
    assert client.get(f'/api/projects/{pid}/weekly-plans?week_start=2001-01-01&include_cancelled=true',headers=worker_headers).status_code==403
    assert client.post(f'/api/projects/{pid}/weekly-plans',headers=pm_headers,json={'task_id':task['id'],'week_start':'2001-01-01'}).status_code==409
    assert client.put(url,headers=pm_headers,json={'version':1,'is_cancelled':False,'reason':'已明确恢复本周安排'}).status_code==200


def test_legacy_days_never_count_or_block_handover_and_published_weekly_obligation_does(client,pm_headers,worker_headers):
    from app.models import ScheduleEntry
    from app.services.member_handover import ensure_member_handover
    from fastapi import HTTPException
    import pytest
    pid,units,ph,task=setup(client,pm_headers)
    with SessionLocal() as db:
        worker=db.query(User).filter_by(username='worker1').one();wid=worker.id
        p=db.get(Project,pid);p.members=[wid];db.add(ScheduleEntry(project_id=pid,task_id=task['id'],work_date=date(2001,1,1),assignee_id=wid,planned_quantity=0,is_confirmed=True));db.commit()
        ensure_member_handover(db,p,{wid})
    data=client.get(f'/api/projects/{pid}/overview',headers=pm_headers)
    assert data.status_code==200,data.text
    assert 'schedule' not in data.json() and data.json()['weekly_plans']['total']==0
    row=create_cell(client,pm_headers,pid,task,assignee_ids=[wid],mode='finish')
    assert publish(client,pm_headers,pid,row).status_code==200
    with SessionLocal() as db:
        with pytest.raises(HTTPException,match='已发布周计划'):ensure_member_handover(db,db.get(Project,pid),{wid})
    detail=client.get(f'/api/projects/{pid}',headers=pm_headers).json()
    assert detail['weekly_plan_summary']['published_count']==1 and detail['weekly_plan_summary']['unassigned_count']==0
    assert client.put(f"/api/projects/{pid}/weekly-plans/{row['id']}",headers=pm_headers,json={'version':1,'is_cancelled':True,'reason':'交接后取消原安排'}).status_code==200
    with SessionLocal() as db:ensure_member_handover(db,db.get(Project,pid),{wid})


def test_excel_preserves_text_and_portable_import_only_quarantines_original_plans(client,pm_headers):
    from io import BytesIO
    from openpyxl import load_workbook
    from app.models import ProjectSchedulePolicy,TaskDurationOverride,ScheduleEntry
    pid,units,ph,task=setup(client,pm_headers)
    row=create_cell(client,pm_headers,pid,task,notes='=HYPERLINK("https://untrusted.invalid")')
    with SessionLocal() as db:
        db.add(ProjectSchedulePolicy(project_id=pid,duration_scale=1.25,version=4))
        db.add(TaskDurationOverride(task_id=task['id'],project_id=pid,workdays=3,baseline_quantity=2,baseline_unit='项',is_active=True,version=7))
        db.commit()
    response=client.get(f'/api/projects/{pid}/weekly-plans/export/excel?week_start=2001-01-01&weeks=1',headers=pm_headers)
    assert response.status_code==200,response.text
    workbook=load_workbook(BytesIO(response.content));assert workbook.sheetnames==['周计划','核对提示','原日程来源']
    assert workbook['周计划']['D4'].data_type=='s' and '=HYPERLINK' in workbook['周计划']['D4'].value
    payload=client.get(f'/api/projects/{pid}/export/json',headers=pm_headers)
    assert payload.status_code==200,payload.text
    original=payload.json();source=original['weekly_plan_source']['cells'][0]
    assert source['id']==row['id'] and source['notes'].startswith('=HYPERLINK')
    assert original['legacy_planning_source']['task_duration_overrides'][0]['version']==7
    original['project']['status']='preparing'
    imported=client.post('/api/projects/import/json',headers=pm_headers,json=original)
    assert imported.status_code==200,imported.text
    with SessionLocal() as db:
        fresh=db.get(Project,imported.json()['id']);assert fresh.portable_unverified_snapshot['weekly_plan_source']['cells'][0]==source
        assert db.query(WeeklyPlanCell).filter_by(project_id=fresh.id).count()==0
        assert db.query(ScheduleEntry).filter_by(project_id=fresh.id).count()==0
        assert db.query(ProjectSchedulePolicy).filter_by(project_id=fresh.id).count()==0
        assert db.query(TaskDurationOverride).filter_by(project_id=fresh.id).count()==0


def test_old_day_routes_and_range_duration_are_explicitly_retired(client,pm_headers):
    pid,units,ph,task=setup(client,pm_headers)
    for path in [f'/api/projects/{pid}/schedule',f'/api/projects/{pid}/schedule/generate',f'/api/projects/{pid}/schedule/reschedule',
                 f"/api/projects/{pid}/schedule/tasks/{task['id']}/range",f"/api/projects/{pid}/schedule/tasks/{task['id']}/duration/auto",'/api/schedule/calendar','/api/schedule/arbitrary-old-id']:
        assert client.post(path,headers=pm_headers,json={}).status_code==410


def test_default_rules_follow_exact_standard_origins_and_calendar_window_not_sort_numbers(client,pm_headers):
    from app.models import Phase,ProjectEvent
    pid,units,_,_=setup(client,pm_headers)
    with SessionLocal() as db:
        for index,origin in enumerate(['phase-01','phase-02','phase-05','phase-08']):
            ph=Phase(project_id=pid,name='任意改名'+str(index),sort_order=index+1);db.add(ph);db.flush()
            db.add(Task(project_id=pid,phase_id=ph.id,name='可核对来源'+str(index),quantity_type='fixed',target_quantity=None,
                legacy_work_snapshot={'kind':'template_definition_source','original':{'kind':'t34_legacy_work','phase':{'id':origin}}}))
        db.add(ProjectEvent(project_id=pid,event_key='ship',planned_date=date(2001,1,15)))
        db.add(ProjectEvent(project_id=pid,event_key='arrival_all',planned_date=date(2001,1,22)))
        db.add(ProjectEvent(project_id=pid,event_key='acceptance',planned_date=date(2001,2,5)))
        db.commit()
    result=client.get(f'/api/projects/{pid}/weekly-plans/draft-rules',headers=pm_headers)
    assert result.status_code==200,result.text
    data=result.json();by_source={row['source_phase_id']:row for row in data['suggestions']}
    assert by_source['phase-01']['anchor']=='before_ship' and by_source['phase-01']['offset_weeks']==1
    assert by_source['phase-01']['week_start']=='2001-01-08'
    assert [by_source[key]['offset_weeks'] for key in ['phase-02','phase-05','phase-08']]==[0,1,2]
    assert by_source['phase-05']['anchor']=='after_arrival'  # Old sort/id 5 is system deployment, not document stage 5.
    assert by_source['phase-08']['week_start']=='2001-02-05'
    assert any(item['code']=='unknown_phase_origin' for item in data['warnings'])
    preview=client.post(f'/api/projects/{pid}/weekly-plans/draft/preview',headers=pm_headers,json={'phase_rules':data['rules']})
    assert preview.status_code==200 and len(preview.json()['proposals'])==4,preview.text
    with SessionLocal() as db:assert db.query(WeeklyPlanCell).filter_by(project_id=pid).count()==0


def test_default_rules_missing_dates_never_guess_and_short_window_allows_overlap(client,pm_headers):
    from app.models import Phase,ProjectEvent
    pid,units,ph,task=setup(client,pm_headers)
    with SessionLocal() as db:
        db.get(Task,task['id']).legacy_work_snapshot={'phase':{'id':'phase-02'}}
        other=Phase(project_id=pid,name='第二现场阶段',sort_order=9);db.add(other);db.flush()
        db.add(Task(project_id=pid,phase_id=other.id,name='另一个工作',quantity_type='fixed',legacy_work_snapshot={'phase':{'id':'phase-03'}}));db.commit()
    path=f'/api/projects/{pid}/weekly-plans/draft-rules'
    missing=client.get(path,headers=pm_headers).json();assert not missing['rules'] and missing['warnings']
    with SessionLocal() as db:
        db.add(ProjectEvent(project_id=pid,event_key='arrival_all',planned_date=date(2001,1,22)))
        db.add(ProjectEvent(project_id=pid,event_key='acceptance',planned_date=date(2001,1,23)));db.commit()
    data=client.get(path,headers=pm_headers).json();assert len(data['rules'])==2 and all(rule['offset_weeks']==0 for rule in data['rules'])
