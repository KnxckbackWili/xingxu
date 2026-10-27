"""Real six-sheet delivery workbooks preserve facts and both evidence ACL layers."""
from datetime import date, datetime
from io import BytesIO
import json
import zipfile

import pytest
from openpyxl import load_workbook
from fastapi import HTTPException

from app.db import SessionLocal
from app import models
from app.services.delivery_workbook import SHEET_NAMES, INVENTORY_HEADERS, MOVEMENT_HEADERS
from app.services.import_export import project_to_excel, project_package_zip
from app.services.report_forms import LIFECYCLE_HEADERS


def create_project(client, headers):
    result = client.post('/api/projects', headers=headers, json={'name':'交付核对项目','start_date':'2026-01-01','end_date':'2026-12-31','members':['user-worker-001']})
    assert result.status_code == 200, result.text
    project = result.json()
    with SessionLocal() as db:
        row = db.get(models.Project,project['id']);row.members=['user-worker-001'];db.commit()
    return project['id']


def workbook(client, headers, project_id):
    result = client.get(f'/api/projects/{project_id}/export/excel', headers=headers)
    assert result.status_code == 200, result.text
    return load_workbook(BytesIO(result.content),data_only=False)


def records(sheet, header=2):
    keys = [cell.value for cell in sheet[header]]
    return [dict(zip(keys,row)) for row in sheet.iter_rows(min_row=header+1,values_only=True)]


def test_exact_six_sheet_contract_headers_styles_and_empty_project_are_real_xlsx(client, admin_headers, tmp_path):
    pid = create_project(client,admin_headers)
    book = workbook(client,admin_headers,pid)
    assert tuple(book.sheetnames) == SHEET_NAMES
    assert [cell.value for cell in book['生命周期表'][2]][:11] == LIFECYCLE_HEADERS
    assert [cell.value for cell in book['物资清单'][1]][:10] == INVENTORY_HEADERS
    assert [cell.value for cell in book['收发流水'][1]][:10] == MOVEMENT_HEADERS
    for sheet in book:
        assert sheet.auto_filter.ref and sheet.freeze_panes
        assert sheet.print_title_rows and sheet.page_setup.fitToWidth == 1
    assert book.active.title == '交付概览'
    assert len(records(book['交付资料'])) > 0  # Actual instantiated T22 requirements.
    path = tmp_path / 'six-sheet-delivery.xlsx';book.save(path)
    with zipfile.ZipFile(path) as bundle:
        assert len([name for name in bundle.namelist() if name.startswith('xl/worksheets/sheet') and name.endswith('.xml')]) == 6


def test_lifecycle_keeps_template_prefix_all_fault_fields_and_no_issue_vehicle(client, admin_headers):
    pid = create_project(client,admin_headers)
    with SessionLocal() as db:
        unit=models.VehicleUnit(project_id=pid,asset_code='AGV-001',model_name='整机型号',fleet_no='007',serial_number='SERIAL-A',current_stage='maintenance',current_status='attention',location='现场A')
        empty=models.VehicleUnit(project_id=pid,asset_code='AGV-002',model_name='无故障车辆',fleet_no='008')
        db.add_all([unit,empty]);db.flush()
        event=models.VehicleLifecycleEvent(project_id=pid,vehicle_unit_id=unit.id,event_type='fault',event_date=date(2026,2,3),title='接插件间歇异常',description='原始故障现象',fault_part='传感器',cause='接触不良',action_taken='重新压接',replaced_parts=[{'name':'接插件','model':'M12','quantity':2,'original_extra':'不能丢'}],downtime_hours=0,verification='连续运行通过',handler='处理人',is_closed=True,stage_after='running',status_after='normal',location_after='现场B',operator_id='user-admin-001',extra={'category':'硬件','custom':'扩展故障事实'})
        db.add(event);db.commit();eid=event.id
    rows=records(workbook(client,admin_headers,pid)['生命周期表'])
    row=next(row for row in rows if row['事件ID']==eid)
    assert row['车号']=='007' and row['出厂序列号']=='SERIAL-A'
    assert row['故障部位']=='传感器' and row['原因判断']=='接触不良' and row['解决方法']=='重新压接'
    assert row['停机小时']==0 and row['验证结果']=='连续运行通过' and row['是否关闭']=='是'
    assert row['事件后阶段']=='running' and row['事件后状态']=='normal' and row['事件后位置']=='现场B'
    assert json.loads(row['更换配件'])[0]['original_extra']=='不能丢'
    assert json.loads(row['扩展信息'])['custom']=='扩展故障事实'
    assert any(row['车辆编号']=='AGV-002' and not row['出现过的问题'] for row in rows)


def test_issue_states_quality_terms_and_extra_are_not_collapsed_or_truncated(client,admin_headers):
    pid=create_project(client,admin_headers)
    with SessionLocal() as db:
        for status in ['open','processing','resolved','closed','cancelled','legacy-review']:
            db.add(models.Issue(project_id=pid,title='状态-'+status,description='问题现象',status=status,issue_type='hardware',severity='high',reporter_id='user-admin-001',resolution='保留解决/作废理由',extra={'vehicle_model':'MODEL','vehicle_no':'001','source':'现场','part':'驱动轮','feedback_date':'2026/02/03','reporter_name':'反馈人原件','nature':'quality','nested':{'parts':['A','B']}}))
        db.commit()
    book=workbook(client,admin_headers,pid);rows=records(book['问题清单'])
    assert {row['状态代码'] for row in rows}=={'open','processing','resolved','closed','cancelled','legacy-review'}
    for row in rows:
        assert row['已闭环']==('是' if row['状态代码'] in {'closed','cancelled'} else '否')
        assert row['车辆编号 No.']=='001' and row['问题来源 From']=='现场' and row['问题部件 Parts']=='驱动轮'
        assert row['反馈人 Name']=='反馈人原件' and json.loads(row['扩展字段'])['nested']=={'parts':['A','B']}
        assert row['解决方案/关闭或作废说明']=='保留解决/作废理由'
    overview=records(book['交付概览'])
    assert next(row['内容'] for row in overview if row['项目/指标']=='未闭环问题')==4


def test_all_inventory_movements_and_legacy_time_provenance_survive_with_numeric_quantities(client,admin_headers):
    pid=create_project(client,admin_headers)
    with SessionLocal() as db:
        item=models.ProjectInventoryItem(project_id=pid,item_code='I-01',name='工具',specification='M12',ownership='client',quantity=0,minimum_quantity=0,unit='件',supplier='甲方',extra={'serial_fact':'保留'})
        db.add(item);db.flush()
        for index in range(75):
            db.add(models.ProjectInventoryMovement(project_id=pid,item_id=item.id,movement_type='transfer' if index%2 else 'receive',quantity=index-.5,balance_after=index,from_location='A',to_location='B',handler='经手人',operator_id='user-admin-001',occurred_at=datetime(2026,2,3,10,0),time_basis=None if index==0 else 'utc',note='未知来源时间' if index==0 else '完整流水-'+str(index)))
        db.flush()
        from sqlalchemy import update
        db.execute(update(models.ProjectInventoryMovement).where(models.ProjectInventoryMovement.note=='未知来源时间').values(time_basis=None))
        db.commit()
    book=workbook(client,admin_headers,pid)
    assert records(book['物资清单'],1)[0]['现存数量']==0
    assert json.loads(records(book['物资清单'],1)[0]['扩展字段'])=={'serial_fact':'保留'}
    rows=records(book['收发流水'],1);assert len(rows)==75
    assert all(isinstance(row['数量'],(int,float)) for row in rows)
    assert sum(row['数量'] for row in rows)==sum(index-.5 for index in range(75))
    old=next(row for row in rows if row['说明']=='未知来源时间')
    assert old['时间依据']=='原记录时区未知' and old['发生时间（UTC）'].startswith('原记录：')
    assert any(row['发生时间（UTC）'].endswith('+00:00') for row in rows if row['时间依据']=='utc')


def test_long_fields_continue_without_excel_cell_truncation_or_double_count_and_formula_text_is_literal(client,admin_headers):
    pid=create_project(client,admin_headers)
    description='第一段😀'*9000
    note='=HYPERLINK("https://unsafe.invalid","公式不得执行")'
    with SessionLocal() as db:
        db.get(models.Project,pid).name=note+'\x01'
        issue=models.Issue(project_id=pid,title=note,description=description,reporter_id='user-admin-001',extra={'long':description})
        item=models.ProjectInventoryItem(project_id=pid,name=note,quantity=3,note=description)
        db.add_all([issue,item]);db.commit();iid=issue.id
    book=workbook(client,admin_headers,pid)
    rows=[row for row in records(book['问题清单']) if row['问题ID']==iid]
    assert len(rows)>1 and ''.join(row['问题简述（或现象）Description'] or '' for row in rows)==description
    assert json.loads(''.join(row['扩展字段'] or '' for row in rows))=={'long':description}
    stock=records(book['物资清单'],1)
    assert sum(row['现存数量'] or 0 for row in stock)==3
    for sheet in book:
        for row in sheet:
            for cell in row:
                assert cell.data_type!='f'
    assert any(cell.value==note and cell.data_type=='s' for row in book['问题清单'] for cell in row)


def test_actor_required_and_confidential_document_or_file_never_leaks_through_another_sheet(client,admin_headers,worker_headers):
    pid=create_project(client,admin_headers)
    with SessionLocal() as db:
        worker=db.get(models.User,'user-worker-001');worker.permission_overrides={'project.export':True};db.commit()
        secret=models.Attachment(project_id=pid,filename='PRIVATE-FILE-NAME.bin',stored_name='private-delivery.bin',file_size=1,confidential=True,uploaded_by='user-admin-001')
        public=models.Attachment(project_id=pid,filename='PRIVATE-DOC-PUBLIC-FILE.bin',stored_name='private-doc-public-file.bin',file_size=1,confidential=False,uploaded_by='user-admin-001')
        allowed=models.Attachment(project_id=pid,filename='VISIBLE-VERSION.bin',stored_name='visible-delivery.bin',file_size=1,uploaded_by='user-admin-001')
        db.add_all([secret,public,allowed]);db.flush()
        hidden=models.ProjectDocument(project_id=pid,instance_key='test-hidden',stage_key='6',code='secret',name='PRIVATE-DOCUMENT-TITLE',confidential=True,status='submitted',submitted_by='user-admin-001')
        shown=models.ProjectDocument(project_id=pid,instance_key='test-shown',stage_key='6',code='public',name='公开资料',status='submitted',submitted_by='user-admin-001')
        db.add_all([hidden,shown]);db.flush()
        for doc,asset,note in [(hidden,public,'PRIVATE-VERSION-NOTE'),(shown,secret,'PRIVATE-FILE-VERSION-NOTE'),(shown,allowed,'VISIBLE-VERSION-NOTE')]:
            db.add(models.DocumentVersion(project_document_id=doc.id,attachment_id=asset.id,uploaded_by='user-admin-001',note=note,review_metadata={'note':note}))
        db.commit()
        with pytest.raises(ValueError,match='当前账号'):project_to_excel(db,pid)
    book=workbook(client,worker_headers,pid)
    visible='\n'.join(str(cell.value or '') for sheet in book for row in sheet for cell in row)
    assert 'VISIBLE-VERSION.bin' in visible and 'VISIBLE-VERSION-NOTE' in visible
    assert 'PRIVATE-' not in visible
    packed=client.get(f'/api/projects/{pid}/export/package',headers=worker_headers)
    assert packed.status_code==200,packed.text
    with zipfile.ZipFile(BytesIO(packed.content)) as bundle:
        text=bundle.read('manifest.json').decode()+bundle.read('项目备份.json').decode()
        assert 'PRIVATE-' not in text and not any('PRIVATE-' in name for name in bundle.namelist())
    admin=workbook(client,admin_headers,pid)
    assert 'PRIVATE-DOCUMENT-TITLE' in '\n'.join(str(cell.value or '') for sheet in admin for row in sheet for cell in row)


def test_workbook_uses_project_scope_and_complete_package_keeps_canonical_json(client,admin_headers):
    pid=create_project(client,admin_headers)
    other=create_project(client,admin_headers)
    with SessionLocal() as db:
        db.add(models.Issue(project_id=other,title='OTHER-PROJECT-FACT',reporter_id='user-admin-001'));db.commit()
        actor=db.get(models.User,'user-admin-001')
        book=load_workbook(BytesIO(project_to_excel(db,pid,actor=actor)))
        assert 'OTHER-PROJECT-FACT' not in '\n'.join(str(cell.value or '') for sheet in book for row in sheet for cell in row)
        data=project_package_zip(db,pid,actor=actor)
    with zipfile.ZipFile(BytesIO(data)) as bundle:
        assert tuple(load_workbook(BytesIO(bundle.read('项目数据.xlsx'))).sheetnames)==SHEET_NAMES
        snapshot=json.loads(bundle.read('项目备份.json'))
        assert 'modules' in snapshot and 'inventory_items' in snapshot and 'document_versions' in snapshot


def test_complete_json_preserves_exact_readonly_execution_evidence_and_reimport_cannot_forge_completion(client,admin_headers):
    pid=create_project(client,admin_headers)
    with SessionLocal() as db:
        from app.services.execution_scope import REMINDER_ONLY_PHASE_NAMES
        phase=db.query(models.Phase).filter(models.Phase.project_id.is_(None),models.Phase.name.notin_(REMINDER_ONLY_PHASE_NAMES)).first()
        module=models.Module(project_id=pid,phase_id=phase.id,name='执行事实模块');db.add(module);db.flush()
        task=models.Task(project_id=pid,module_id=module.id,name='原直接完成任务',target_quantity=1,unit='台',manually_completed=True,
            manual_completed_by='user-admin-001',manual_completed_role='admin',manual_completed_at=datetime(2026,2,3,4,5),manual_completed_reason='精确保留覆盖理由')
        db.add(task);db.flush()
        check=models.CheckItem(task_id=task.id,project_id=pid,name='保留失败检查',failed=True,failed_reason='原失败事实',is_done=False,version=7)
        db.add(check);db.flush()
        exact=[{'check_id':check.id,'version':7,'name':check.name,'failed_reason':check.failed_reason}]
        task.manual_check_overrides=exact
        db.add(models.ChangeLog(project_id=pid,entity_type='task',entity_id=check.id,action='update',operator_id='user-admin-001',detail='检查项「保留失败检查」不通过：原失败事实',created_at=datetime(2026,2,3,3,0)))
        db.commit();tid=task.id;cid=check.id
    response=client.get(f'/api/projects/{pid}/export/json',headers=admin_headers)
    assert response.status_code==200,response.text
    payload=response.json();source=payload['execution_evidence']
    assert source['read_only']
    task=next(row for row in source['tasks'] if row['source_id']==tid)
    assert task['manually_completed'] and task['manual_completed_reason']=='精确保留覆盖理由'
    assert task['manual_check_overrides']==exact and task['manual_completed_by_name']
    check=next(row for row in task['check_items'] if row['id']==cid)
    assert check['version']==7 and check['failed'] and check['failed_reason']=='原失败事实'
    assert check['last_operation']['recorded_at'].startswith('2026-02-03T03:00') and check['last_operation']['actor_name']
    payload['project']['name']='重导只读执行事实'
    imported=client.post('/api/projects/import/json',headers=admin_headers,json=payload)
    assert imported.status_code==200,imported.text
    with SessionLocal() as db:
        tasks=db.query(models.Task).filter_by(project_id=imported.json()['id']).all()
        assert tasks and all(not row.manually_completed and not row.manual_check_overrides for row in tasks)
        checks=db.query(models.CheckItem).filter(models.CheckItem.task_id.in_([row.id for row in tasks])).all()
        assert checks and all(not row.is_done and not row.failed for row in checks)


def test_vehicle_matrix_and_unknown_identity_credits_fit_six_sheets_and_stay_unverified_on_import(client,admin_headers):
    from app.services.vehicle_progress import matrix
    from tests.test_project_events import unit
    from tests.test_vehicle_progress import write
    pid=create_project(client,admin_headers)
    response=client.post(f'/api/projects/{pid}/vehicles',headers=admin_headers,json={'vehicle_type':'潜伏式','model_name':'原实车','quantity':2})
    assert response.status_code==200,response.text
    uid=unit(client,admin_headers,pid,response.json()['id'],'AGV-原实车')
    with SessionLocal() as db:
        db.get(models.Project,pid).status='executing'
        db.get(models.VehicleUnit,uid).legacy_current_stage='production'
        db.add(models.ProjectVehicleStep(project_id=pid,key='legacy_unknown',name='历史环节：原调试',expected_quantity=2,sort_order=5))
        db.add(models.VehicleProgressMigration(project_id=pid,step_key='legacy_unknown',total_count=1,unassigned_count=1,
            source_metadata={'task_id':'original-task','target_quantity':2.4,'completed_quantity':1.2,'manually_completed':False,'identity':'unknown'}))
        db.commit()
    assert write(client,admin_headers,pid,[uid]).status_code==200
    book=workbook(client,admin_headers,pid)
    assert tuple(book.sheetnames)==SHEET_NAMES
    overview=records(book['交付概览'])
    assert any(row['类别']=='车辆环节事实' and 'AGV-原实车' in str(row) for row in overview)
    assert any(row['类别']=='历史完成凭证' and '1.2' in str(row) and '原调试' in str(row) for row in overview)
    assert records(book['生命周期表'])[0]['升级前九阶段（只读）']=='production'
    payload=client.get(f'/api/projects/{pid}/export/json',headers=admin_headers).json()
    facts=payload['execution_evidence']['vehicle_progress_matrix']
    assert facts['read_only'] and facts['summary']['completed']==2
    payload['project']['status']='draft'
    imported=client.post('/api/projects/import/json',headers=admin_headers,json=payload)
    assert imported.status_code==200,imported.text
    imported_id=imported.json()['id']
    with SessionLocal() as db:
        assert matrix(db,imported_id)['summary']['completed']==0
        assert db.query(models.VehicleProgress).filter_by(project_id=imported_id).count()==0
        assert db.query(models.VehicleProgressMigration).filter_by(project_id=imported_id).count()==0
        origin=db.get(models.Project,imported_id).portable_unverified_snapshot
        assert origin['vehicle_progress_matrix']['legacy']==payload['vehicle_progress_matrix']['legacy']
        assert origin['execution_evidence']==payload['execution_evidence']
    source=client.get(f'/api/projects/{imported_id}/vehicle-progress/source',headers=admin_headers)
    assert source.status_code==200 and source.json()['verified'] is False


def test_portable_quantity_history_keeps_exact_original_fields_without_completing_new_work(client,admin_headers):
    from app.services.progress import project_progress_tree
    pid=create_project(client,admin_headers)
    with SessionLocal() as db:
        phase=models.Phase(project_id=pid,name='非车辆工作');db.add(phase);db.flush()
        module=models.Module(project_id=pid,phase_id=phase.id,name='原模块');db.add(module);db.flush()
        task=models.Task(project_id=pid,module_id=module.id,name='原1000米工作',quantity_type='fixed',target_quantity=1000,unit='米');db.add(task);db.flush()
        report=models.DailyProgress(project_id=pid,task_id=task.id,user_id='user-admin-001',completed_quantity=1000,
            work_date=date(2026,2,3),created_at=datetime(2026,2,5,8,9,10),client_uuid='portable-original-uuid',source='admin',note='原项目完整数量与备注')
        db.add(report);db.commit()
    payload=client.get(f'/api/projects/{pid}/export/json',headers=admin_headers).json()
    original=payload['progress']
    assert original[0]['client_uuid']=='portable-original-uuid'
    assert original[0]['user_id']=='user-admin-001' and original[0]['created_at'].startswith('2026-02-05T08:09:10')
    assert original[0]['work_date']=='2026-02-03' and original[0]['completed_quantity']==1000
    imported=client.post('/api/projects/import/json',headers=admin_headers,json=payload)
    assert imported.status_code==200,imported.text
    target=imported.json()['id']
    with SessionLocal() as db:
        assert db.query(models.DailyProgress).filter_by(project_id=target).count()==0
        assert project_progress_tree(db,target)['completed']==0
        origin=db.get(models.Project,target).portable_unverified_snapshot
        assert origin['legacy_quantity_reports']==original
    forbidden=client.put(f'/api/projects/{target}',headers=admin_headers,json={'portable_unverified_snapshot':{}})
    assert forbidden.status_code==422,forbidden.text
    assert client.get(f'/api/projects/{target}/vehicle-progress/source',headers=admin_headers).json()['source']['legacy_quantity_reports']==original
