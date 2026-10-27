"""T21 rules and portable files, through the interfaces used by users."""
import copy
import hashlib
import io
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from app.db import SessionLocal
from app.models import DocumentTemplateFile, Template, ChangeLog
from app.services.document_requirements import RequirementInput

BASE = '/api/templates/default/document-requirements'


def body(**changes):
    return {'stage_key': '6', 'code': '6.99', 'name': '调试证明', 'scope': 'project', 'timing_type': 'on_demand', **changes}


def list_rows(client, headers, tid='default'):
    result = client.get(f'/api/templates/{tid}/document-requirements', headers=headers)
    assert result.status_code == 200, result.text
    return result.json()


def clean_export(bundle):
    bundle = copy.deepcopy(bundle)
    files = bundle['document_template_files']
    for row in bundle['document_requirements']:
        row['file'] = files.get(row.pop('blank_file_key'))
    return bundle['document_requirements']


def test_default_exact_frozen_inventory_and_custom_copy(client, admin_headers):
    rows = list_rows(client, admin_headers)
    seeds = json.loads((Path(__file__).parents[1] / 'migrations/v3_document_requirements_seed.json').read_text())
    standard=[row for row in rows if row['stage_key']!='legacy']
    assert len(standard) == 56 and sum(not row['is_group'] for row in standard) == 53
    assert all(row['source_metadata'] for row in rows if row['stage_key']=='legacy')
    for seed, row in zip(seeds, standard):
        expected={**seed}
        if seed['code'] in {'6.11.6','6.11.6a','6.11.6b','6.11.6c'}:
            expected.update(confidential_default=True,external_visible_default=False,external_never_visible=True)
        assert all(row[key] == value for key, value in expected.items())
        assert row['blank_file'] is None
    assert {row['stage_key'] for row in standard} == set('456789')
    assert not any(row['external_visible_default'] for row in rows)
    result = client.post('/api/templates', headers=admin_headers, json={'name': '资料完整副本'})
    assert result.status_code == 200, result.text
    copied = list_rows(client, admin_headers, result.json()['id'])
    assert len(copied) == len(rows)
    assert not {row['id'] for row in copied} & {row['id'] for row in rows}
    assert {row['code']: row['parent_code'] for row in copied} == {row['code']: row['parent_code'] for row in rows}


@pytest.mark.parametrize('changes', [
    {'timing_type': 'weekly'}, {'timing_type': 'weekly', 'timing_weekday': 0},
    {'timing_type': 'weekly', 'timing_weekday': 8}, {'timing_type': 'weekly', 'timing_weekday': 7, 'timing_event': 'ship'},
    {'timing_type': 'after_event', 'timing_event': 'vehicle_arrival', 'timing_offset_hours': 24},
    {'timing_type': 'before_event', 'timing_event': 'unknown', 'timing_offset_hours': 0},
    {'timing_type': 'after_event', 'timing_event': 'ship', 'timing_offset_hours': -1},
    {'timing_type': 'after_event', 'timing_event': 'ship', 'timing_offset_hours': 1.5},
    {'timing_type': 'on_trigger', 'trigger_type': 'bad', 'timing_offset_hours': 0},
    {'timing_type': 'on_demand', 'timing_weekday': 7}, {'stage_key': '3'}, {'code': '5.9'},
    {'confidential_default': True, 'external_visible_default': True},
    {'code': '6.11.6b', 'external_visible_default': True},
    {'is_group': True, 'required': False, 'scope': None, 'timing_type': None, 'timing_event': 'ship'},
])
def test_invalid_rules_fail_before_writes(changes):
    with pytest.raises(ValidationError):
        RequirementInput.model_validate(body(**changes))


def test_crud_parent_rules_permissions_and_audit(client, admin_headers, worker_headers):
    assert client.post(BASE, headers=worker_headers, json=body()).status_code == 403
    assert client.post(BASE, headers=admin_headers, json=body(timing_type='weekly')).status_code == 422
    created = client.post(BASE, headers=admin_headers, json=body(timing_type='weekly', timing_weekday=7))
    assert created.status_code == 200, created.text
    rid = created.json()['id']
    assert client.post(BASE, headers=admin_headers, json=body()).status_code == 409
    leaf = next(row for row in list_rows(client, admin_headers) if row['code'] == '6.2')
    assert client.put(f'{BASE}/{rid}', headers=admin_headers, json=body(parent_id=leaf['id'])).status_code == 400
    group = next(row for row in list_rows(client, admin_headers) if row['code'] == '6.11')
    assert client.put(f'{BASE}/{rid}', headers=admin_headers, json=body(parent_id=group['id'])).status_code == 200
    assert client.delete(f'{BASE}/{group["id"]}', headers=admin_headers).status_code == 400
    assert client.delete(f'{BASE}/{rid}', headers=admin_headers).status_code == 200
    assert client.get('/api/templates/missing/document-requirements', headers=admin_headers).status_code == 404
    with SessionLocal() as db:
        assert db.query(ChangeLog).filter_by(entity_type='document_requirement', entity_id=rid).count() == 3


def test_blank_form_copy_detach_snapshot_and_tree_save(client, admin_headers):
    rid = next(row for row in list_rows(client, admin_headers) if row['code'] == '6.1')['id']
    data = b'blank form\x00' * 7000
    uploaded = client.post(f'{BASE}/{rid}/blank-file', headers=admin_headers, files={'file': ('点检表.xlsx', data, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert uploaded.status_code == 200, uploaded.text
    assert uploaded.json()['sha256'] == hashlib.sha256(data).hexdigest()
    tid = client.post('/api/templates', headers=admin_headers, json={'name': '表格副本'}).json()['id']
    copied = next(row for row in list_rows(client, admin_headers, tid) if row['code'] == '6.1')
    with SessionLocal() as db:
        assert db.get(DocumentTemplateFile, uploaded.json()['id']).stored_name != db.get(DocumentTemplateFile, copied['blank_file_id']).stored_name
    assert client.post(f'/api/templates/{tid}/publish', headers=admin_headers).status_code == 200
    with SessionLocal() as db:
        assert len(db.get(Template, tid).snapshots[-1]['document_requirements']) == len(list_rows(client,admin_headers,tid))
    assert client.delete(f'{BASE}/{rid}/blank-file', headers=admin_headers).status_code == 200
    assert client.get(f'{BASE}/{rid}/blank-file', headers=admin_headers).status_code == 404
    downloaded = client.get(f'/api/templates/{tid}/document-requirements/{copied["id"]}/blank-file', headers=admin_headers)
    assert downloaded.status_code == 200 and downloaded.content == data
    tree = client.get(f'/api/templates/{tid}/tree', headers=admin_headers).json()
    assert client.put(f'/api/templates/{tid}/tree', headers=admin_headers, json={'phases': tree}).status_code == 200
    assert len([row for row in list_rows(client, admin_headers, tid) if row['stage_key']!='legacy']) == 56
    assert client.post(f'{BASE}/{rid}/blank-file', headers=admin_headers, files={'file': ('../escape', b'x')}).status_code == 400


def test_blank_upload_audit_failure_rolls_back_file_and_pointer(client, admin_headers, monkeypatch):
    from app.config import UPLOAD_DIR
    import app.routers.document_requirements as routes
    rid = next(row for row in list_rows(client, admin_headers) if row['code'] == '4.1')['id']
    assert client.post(f'{BASE}/{rid}/blank-file', headers=admin_headers, files={'file': ('old.txt', b'old')}).status_code == 200
    before_files = set(UPLOAD_DIR.iterdir())
    before_row = next(row for row in list_rows(client, admin_headers) if row['id'] == rid)
    def fail(*args, **kwargs): raise RuntimeError('audit unavailable')
    monkeypatch.setattr(routes, 'log_change', fail)
    with pytest.raises(RuntimeError, match='audit unavailable'):
        client.post(f'{BASE}/{rid}/blank-file', headers=admin_headers, files={'file': ('new.txt', b'new')})
    assert set(UPLOAD_DIR.iterdir()) == before_files
    assert next(row for row in list_rows(client, admin_headers) if row['id'] == rid)['blank_file_id'] == before_row['blank_file_id']


def test_json_excel_portable_binary_roundtrip_and_legacy_preservation(client, admin_headers):
    rid = next(row for row in list_rows(client, admin_headers) if row['code'] == '6.1')['id']
    data = bytes(range(256)) * 300
    assert client.post(f'{BASE}/{rid}/blank-file', headers=admin_headers, files={'file': ('template.bin', data)}).status_code == 200
    source = client.get('/api/templates/default/export/json', headers=admin_headers).json()
    tid = client.post('/api/templates', headers=admin_headers, json={'name': '导入目标'}).json()['id']
    result = client.post(f'/api/templates/{tid}/import/json', headers=admin_headers, json=source)
    assert result.status_code == 200, result.text
    assert clean_export(source) == clean_export(client.get(f'/api/templates/{tid}/export/json', headers=admin_headers).json())
    workbook = client.get('/api/templates/export/excel', params={'template_id': 'default'}, headers=admin_headers)
    assert workbook.status_code == 200
    wb = load_workbook(io.BytesIO(workbook.content))
    assert wb['空白模板文件'].sheet_state == 'hidden' and wb['空白模板文件'].max_row > 3
    wb.active = 1
    modified = io.BytesIO(); wb.save(modified)
    imported = client.post('/api/templates/import/excel', params={'template_id': tid}, headers=admin_headers, files={'file': ('roundtrip.xlsx', modified.getvalue())})
    assert imported.status_code == 200, imported.text
    after = client.get(f'/api/templates/{tid}/export/json', headers=admin_headers).json()
    assert clean_export(source) == clean_export(after)
    del wb['资料要求']; del wb['空白模板文件']; old = io.BytesIO(); wb.save(old)
    assert client.post('/api/templates/import/excel', params={'template_id': tid}, headers=admin_headers, files={'file': ('old.xlsx', old.getvalue())}).status_code == 200
    assert client.post(f'/api/templates/{tid}/import/json', headers=admin_headers, json={'phases': after['phases']}).status_code == 200
    assert clean_export(after) == clean_export(client.get(f'/api/templates/{tid}/export/json', headers=admin_headers).json())


def test_invalid_import_atomic_and_no_startup_resurrection(client, admin_headers):
    from app.db import init_db
    before = client.get('/api/templates/default/export/json', headers=admin_headers).json()
    changed = copy.deepcopy(before)
    changed['phases'][0]['name'] = 'must not commit'
    changed['document_requirements'][0].update(timing_type='weekly', timing_event=None, timing_offset_hours=None)
    assert client.post('/api/templates/default/import/json', headers=admin_headers, json=changed).status_code == 400
    assert client.get('/api/templates/default/export/json', headers=admin_headers).json() == before
    rid = next(row for row in list_rows(client, admin_headers) if row['code'] == '4.1')['id']
    assert client.delete(f'{BASE}/{rid}', headers=admin_headers).status_code == 200
    init_db()
    assert len([row for row in list_rows(client, admin_headers) if row['stage_key']!='legacy']) == 55
    assert not any(row['code'] == '4.1' for row in list_rows(client, admin_headers))


def test_bad_file_hash_and_cross_template_parent_reject_atomically(client, admin_headers):
    from app.config import UPLOAD_DIR
    rid = next(row for row in list_rows(client, admin_headers) if row['code'] == '4.1')['id']
    assert client.post(f'{BASE}/{rid}/blank-file', headers=admin_headers, files={'file': ('原件.txt', b'correct bytes')}).status_code == 200
    before = client.get('/api/templates/default/export/json', headers=admin_headers).json()
    files_before = set(UPLOAD_DIR.iterdir())
    broken = copy.deepcopy(before)
    next(iter(broken['document_template_files'].values()))['sha256'] = '0' * 64
    broken['phases'][0]['name'] = 'should stay unchanged'
    rejected = client.post('/api/templates/default/import/json', headers=admin_headers, json=broken)
    assert rejected.status_code == 400 and 'SHA256' in rejected.text
    assert set(UPLOAD_DIR.iterdir()) == files_before
    assert client.get('/api/templates/default/export/json', headers=admin_headers).json() == before
    workbook = client.get('/api/templates/export/excel', headers=admin_headers).content
    wb = load_workbook(io.BytesIO(workbook)); wb['空白模板文件']['E2'] = '0' * 64
    invalid = io.BytesIO(); wb.save(invalid)
    assert client.post('/api/templates/import/excel', headers=admin_headers, files={'file': ('bad.xlsx', invalid.getvalue())}).status_code == 400
    assert set(UPLOAD_DIR.iterdir()) == files_before
    tid = client.post('/api/templates', headers=admin_headers, json={'name': '隔离父级'}).json()['id']
    parent = next(row for row in list_rows(client, admin_headers, tid) if row['code'] == '6.11')
    assert client.post(BASE, headers=admin_headers, json=body(parent_id=parent['id'])).status_code == 404
    # Delete a source template after cloning; the destination form stays intact.
    clone = client.post('/api/templates', headers=admin_headers, json={'name': '独立保全副本', 'copy_from': tid}).json()['id']
    copied = next(row for row in list_rows(client, admin_headers, clone) if row['code'] == '4.1')
    assert client.delete(f'/api/templates/{tid}', headers=admin_headers).status_code == 200
    assert client.get(f'/api/templates/{clone}/document-requirements/{copied["id"]}/blank-file', headers=admin_headers).content == b'correct bytes'


def test_declared_permissions_apply_denials_and_explicit_grants(client, admin_headers, worker_headers):
    from app.models import User
    with SessionLocal() as db:
        worker = db.query(User).filter_by(username='worker1').one()
        worker.permission_overrides = {'system.templates.read': False, 'system.templates.manage': True}
        db.commit()
    assert client.get(BASE, headers=worker_headers).status_code == 403
    allowed = client.post(BASE, headers=worker_headers, json=body())
    assert allowed.status_code == 200, allowed.text
    with SessionLocal() as db:
        worker = db.query(User).filter_by(username='worker1').one()
        worker.permission_overrides = {'system.templates.manage': False}
        db.commit()
    assert client.put(f'{BASE}/{allowed.json()["id"]}', headers=worker_headers, json=body(name='禁止修改')).status_code == 403
    exported = client.get('/api/templates/default/export/json', headers=admin_headers).json()
    with SessionLocal() as db:
        worker = db.query(User).filter_by(username='worker1').one()
        worker.permission_overrides = {'system.template_io.manage': True, 'system.templates.manage': False}
        db.commit()
    # Import authority is a separate action, never an accidental AND with editor authority.
    assert client.post('/api/templates/default/import/json', headers=worker_headers, json=exported).status_code == 200



@pytest.mark.parametrize('custom', [False, True])
def test_concurrent_parent_edits_cannot_form_a_cycle(client, admin_headers, custom):
    import threading
    from fastapi import HTTPException
    from sqlalchemy import event
    from app.db import engine
    from app.models import DocumentRequirement, User
    from app.services.document_requirements import lock_templates, save_requirement
    tid = client.post('/api/templates', headers=admin_headers, json={'name': '并发模板'}).json()['id'] if custom else 'default'
    endpoint = f'/api/templates/{tid}/document-requirements'
    group = lambda code: body(code=code, is_group=True, required=False, scope=None, timing_type=None)
    first = client.post(endpoint, headers=admin_headers, json=group('6.90')).json()['id']
    second = client.post(endpoint, headers=admin_headers, json=group('6.91')).json()['id']
    with SessionLocal() as db:
        actor = db.query(User.id).filter_by(username='admin').scalar()
    acquired, attempted, release = threading.Event(), threading.Event(), threading.Event()
    results = []
    def observe(conn, cursor, statement, parameters, context, executemany):
        if threading.current_thread().name == 'T21-parent-second' and ('UPDATE ' in statement.upper() or 'PG_ADVISORY_XACT_LOCK' in statement.upper()):
            attempted.set()
    event.listen(engine, 'before_cursor_execute', observe)
    def writer_one():
        with SessionLocal() as db:
            try:
                lock_templates(db, tid)
                acquired.set()
                assert release.wait(5)
                save_requirement(db, tid, RequirementInput.model_validate(group('6.90') | {'parent_id': second}), actor, first)
                db.commit(); results.append(('first', 200))
            except Exception as error:
                db.rollback(); results.append(('first', repr(error)))
    def writer_two():
        with SessionLocal() as db:
            try:
                save_requirement(db, tid, RequirementInput.model_validate(group('6.91') | {'parent_id': first}), actor, second)
                db.commit(); results.append(('second', 200))
            except HTTPException as error:
                db.rollback(); results.append(('second', error.status_code))
            except Exception as error:
                db.rollback(); results.append(('second', repr(error)))
    a = threading.Thread(target=writer_one, name='T21-parent-first')
    b = threading.Thread(target=writer_two, name='T21-parent-second')
    try:
        a.start(); assert acquired.wait(5)
        b.start(); assert attempted.wait(5)
        release.set(); a.join(8); b.join(8)
        assert not a.is_alive() and not b.is_alive()
    finally:
        release.set()
        a.join(8)
        if b.ident is not None: b.join(8)
        event.remove(engine, 'before_cursor_execute', observe)
    assert ('first', 200) in results, results
    assert any(name == 'second' and result in (400, 409) for name, result in results), results
    with SessionLocal() as db:
        assert db.get(DocumentRequirement, first).parent_id == second
        assert db.get(DocumentRequirement, second).parent_id is None
    assert client.put(f'{endpoint}/{second}', headers=admin_headers, json=group('6.91') | {'parent_id': first}).status_code == 400


def test_copy_and_import_audit_failure_remove_only_new_files(client, admin_headers, monkeypatch):
    from app.config import UPLOAD_DIR
    import app.routers.templates as templates
    import app.services.document_requirements as docs
    rid = next(row for row in list_rows(client, admin_headers) if row['code'] == '4.1')['id']
    assert client.post(f'{BASE}/{rid}/blank-file', headers=admin_headers, files={'file': ('原模板.txt', b'original')}).status_code == 200
    original = client.get('/api/templates/default/export/json', headers=admin_headers).json()
    before_files = set(UPLOAD_DIR.iterdir())
    def fail(*args, **kwargs): raise RuntimeError('audit refused')
    with monkeypatch.context() as patch:
        patch.setattr(templates, 'log_change', fail)
        with pytest.raises(RuntimeError, match='audit refused'):
            client.post('/api/templates', headers=admin_headers, json={'name': '不应留下失败副本'})
    assert set(UPLOAD_DIR.iterdir()) == before_files
    with SessionLocal() as db:
        assert db.query(Template).filter_by(name='不应留下失败副本').count() == 0
    with monkeypatch.context() as patch:
        patch.setattr(docs, 'log_change', fail)
        with pytest.raises(RuntimeError, match='audit refused'):
            client.post('/api/templates/default/import/json', headers=admin_headers, json=original)
    assert set(UPLOAD_DIR.iterdir()) == before_files
    assert client.get('/api/templates/default/export/json', headers=admin_headers).json() == original


@pytest.mark.parametrize('field', ['parent_code', 'blank_file_key'])
@pytest.mark.parametrize('value', [[], ['bad'], {}, {'bad': 'ref'}, False, 0, '', ' ', 'x' * 41])
def test_portable_reference_shape_rejects_bad_types(field, value):
    from app.services.document_requirements import parse_bundle
    definition = body() | {field: value}
    with pytest.raises(ValueError, match='父编号|文件键'):
        parse_bundle({'document_requirements': [definition], 'document_template_files': {}})


def test_malformed_portable_references_leave_tree_files_and_audit_unchanged(client, admin_headers):
    from app.config import UPLOAD_DIR
    rid = next(row for row in list_rows(client, admin_headers) if row['code'] == '4.1')['id']
    assert client.post(f'{BASE}/{rid}/blank-file', headers=admin_headers, files={'file': ('keep.txt', b'original bytes')}).status_code == 200
    original = client.get('/api/templates/default/export/json', headers=admin_headers).json()
    workbook = client.get('/api/templates/export/excel', headers=admin_headers).content
    files_before = {path.name: hashlib.sha256(path.read_bytes()).hexdigest() for path in UPLOAD_DIR.iterdir() if path.is_file()}
    def audit():
        with SessionLocal() as db:
            return [(row.id, row.detail, row.old_value, row.new_value) for row in db.query(ChangeLog).order_by(ChangeLog.id)]
    audit_before = audit()
    for field in ('parent_code', 'blank_file_key'):
        for value in ([], ['bad'], {}, False, 0, '', ' ', 'x' * 41):
            broken = copy.deepcopy(original)
            broken['phases'][0]['name'] = 'must remain unchanged'
            broken['document_requirements'][0][field] = value
            response = client.post('/api/templates/default/import/json', headers=admin_headers, json=broken)
            assert response.status_code in (400, 422), (field, value, response.text)
    # Malformed spreadsheet references must not be coerced into logical file/parent IDs.
    for column, value in (('F', False), ('F', 'x' * 41), ('T', False), ('T', 'x' * 37)):
        wb = load_workbook(io.BytesIO(workbook))
        wb['实施模板']['A3'] = 'must remain unchanged'
        wb['资料要求'][f'{column}3'] = value
        packed = io.BytesIO(); wb.save(packed)
        response = client.post('/api/templates/import/excel', headers=admin_headers, files={'file': ('invalid.xlsx', packed.getvalue())})
        assert response.status_code in (400, 422), (column, value, response.text)
    assert client.get('/api/templates/default/export/json', headers=admin_headers).json() == original
    assert files_before == {path.name: hashlib.sha256(path.read_bytes()).hexdigest() for path in UPLOAD_DIR.iterdir() if path.is_file()}
    assert audit() == audit_before
