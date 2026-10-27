"""Deadline facts, site calendars, durable schedules and real API boundaries."""
from datetime import date, datetime, timedelta, timezone
from types import SimpleNamespace as Obj

import pytest

from app.db import SessionLocal
from app.models import (Project, ProjectDocument, ProjectDocumentDefinition, ProjectEvent, User,
                        DocumentNoticeReceipt, ProjectReportReceipt, Notification, NotificationOutbox)
from app.services.document_deadlines import calculate, refresh_project, generate_weekly, remind_project, action_summary
from .test_project_documents import create, view, upload


def rule(**fields):
    return Obj(stage_key='6',timing_type='after_event',timing_event='ship',timing_offset_hours=24,
               period_key=None,timing_weekday=None,vehicle_unit_id=None,triggered_at=None,trigger_ref=None,**fields)


def test_actual_planned_and_date_precision_have_distinct_bases():
    project=Obj(timezone='America/New_York')
    document=rule()
    planned=Obj(planned_date=date(2026,3,8),actual_at=None)
    result=calculate(document,project,{'ship':planned},{})
    assert result.due_at==datetime(2026,3,9,5) and result.estimated
    assert result.basis['kind']=='planned_date_anchor' and 'actual_at' not in result.basis
    planned.actual_at=datetime(2026,3,8,12,30)
    result=calculate(document,project,{'ship':planned},{})
    assert result.due_at==datetime(2026,3,9,12,30) and not result.estimated
    document.timing_event='vehicle_arrival';document.vehicle_unit_id='car'
    result=calculate(document,project,{}, {'car':Obj(id='car',arrived_at=date(2026,3,8))})
    assert result.due_at==datetime(2026,3,9,5)
    assert result.basis['precision']=='date' and result.basis['actual_date']=='2026-03-08'
    assert 'actual_at' not in result.basis


@pytest.mark.parametrize('key,expected',[
    ('2026-W53',datetime(2027,1,4,4,59,59,999999)),
    ('2026-W10',datetime(2026,3,9,3,59,59,999999)),
    ('2026-W44',datetime(2026,11,2,4,59,59,999999)),
])
def test_weekly_iso_year_and_dst_end_of_local_sunday(key,expected):
    document=rule();document.timing_type='weekly';document.period_key=key;document.timing_weekday=7
    result=calculate(document,Obj(timezone='America/New_York'),{},{})
    assert result.due_at==expected and not result.estimated
    assert result.basis['period_key']==key


def test_plan_change_actual_clear_and_submitted_snapshot_via_real_api(client,admin_headers):
    pid=create(client,admin_headers)['id']
    path=f'/api/projects/{pid}/events/ship'
    assert client.put(path,headers=admin_headers,json={'planned_date':'2026-11-20'}).status_code==200
    document=next(row for row in view(client,admin_headers,pid)['documents'] if row['code']=='5.1.1')
    assert document['due_at']=='2026-11-12T16:00:00+00:00' and document['due_is_estimated']
    assert client.put(path,headers=admin_headers,json={'actual_at':'2026-11-21T15:30:00+08:00'}).status_code==200
    current=next(row for row in view(client,admin_headers,pid)['documents'] if row['id']==document['id'])
    assert current['due_at']=='2026-11-14T07:30:00+00:00' and not current['due_is_estimated']
    assert client.put(path,headers=admin_headers,json={'actual_at':None,'reason':'真实登记错误，恢复待确认'}).status_code==200
    current=next(row for row in view(client,admin_headers,pid)['documents'] if row['id']==document['id'])
    assert current['due_at']==document['due_at'] and current['due_is_estimated']
    assert upload(client,admin_headers,pid,[current['id']]).status_code==200
    frozen=next(row for row in view(client,admin_headers,pid)['documents'] if row['id']==document['id'])
    assert client.put(path,headers=admin_headers,json={'planned_date':'2026-12-01'}).status_code==200
    kept=next(row for row in view(client,admin_headers,pid)['documents'] if row['id']==document['id'])
    assert (kept['due_at'],kept['due_basis'],kept['version'])==(frozen['due_at'],frozen['due_basis'],frozen['version'])


def weekly_fixture(client,headers,timezone='UTC'):
    pid=create(client,headers,timezone=timezone)['id']
    with SessionLocal() as db:
        project=db.get(Project,pid);project.status='preparing';project.created_at=datetime(2026,10,27)
        # Preparing projects now instantiate company weekly rules immediately.
        # This fixture tests a single explicitly activated weekly rule.
        db.query(ProjectDocument).filter_by(project_id=pid,timing_type='weekly').delete(synchronize_session=False)
        db.query(ProjectDocumentDefinition).filter_by(project_id=pid,timing_type='weekly').delete(synchronize_session=False)
        definition=ProjectDocumentDefinition(project_id=pid,source_key='test-real-weekly',stage_key='6',code='6.99',name='实际周资料',
            scope='project',timing_type='weekly',timing_weekday=7,created_at=datetime(2026,10,27))
        db.add(definition);db.commit();return pid,definition.id


def test_missed_weeks_do_not_invent_pre_activation_obligations(client,admin_headers):
    pid,rid=weekly_fixture(client,admin_headers,'Pacific/Kiritimati')
    with SessionLocal() as db:
        project=db.get(Project,pid)
        assert generate_weekly(db,project,datetime(2026,11,8,10))==3
        refresh_project(db,project,now=datetime(2026,11,8,10));db.commit()
        rows=db.query(ProjectDocument).filter_by(definition_id=rid).order_by(ProjectDocument.period_key).all()
        assert [row.period_key for row in rows]==['2026-W44','2026-W45','2026-W46']
        assert rows[0].due_at==datetime(2026,11,1,9,59,59,999999)
        assert generate_weekly(db,project,datetime(2026,11,8,10))==0
        summary=action_summary(db,project,None,datetime(2026,11,8,10))
        assert summary['overdue_count']==2 and summary['period_key']=='2026-W46'


def test_custom_weekday_never_backfills_a_deadline_before_rule_activation(client,admin_headers):
    pid,rid=weekly_fixture(client,admin_headers)
    with SessionLocal() as db:
        definition=db.get(ProjectDocumentDefinition,rid);definition.timing_weekday=1;db.flush()
        project=db.get(Project,pid)
        assert generate_weekly(db,project,datetime(2026,10,27,12))==0
        assert generate_weekly(db,project,datetime(2026,11,2,12))==1
        refresh_project(db,project,now=datetime(2026,11,2,12));db.commit()
        row=db.query(ProjectDocument).filter_by(definition_id=rid).one()
        assert row.period_key=='2026-W45' and row.due_at==datetime(2026,11,2,23,59,59,999999)


def test_owner_assignment_and_durable_reminder_escalation(client,admin_headers,worker_headers,monkeypatch):
    pid=create(client,admin_headers)['id'];did=next(row for row in view(client,admin_headers,pid)['documents'] if row['code']=='4.5')['id']
    with SessionLocal() as db:
        project=db.get(Project,pid);project.status='preparing';project.members=['user-worker-001'];project.technical_manager='user-worker-001'
        # Isolate this one non-periodic reminder from the newly active company's
        # current-week rules; their escalation is tested separately.
        db.query(ProjectDocument).filter_by(project_id=pid,timing_type='weekly').delete(synchronize_session=False)
        db.query(ProjectDocumentDefinition).filter_by(project_id=pid,timing_type='weekly').delete(synchronize_session=False)
        row=db.get(ProjectDocument,did);row.due_at=datetime(2026,11,2,12);row.due_is_estimated=False;db.commit()
        assert remind_project(db,project,datetime(2026,11,1,11,59))['notified']==0
        assert remind_project(db,project,datetime(2026,11,1,12))['notified']==1;db.commit()
        assert remind_project(db,project,datetime(2026,11,1,13))['notified']==0
        assert remind_project(db,project,datetime(2026,11,2,12,1))['notified']==2;db.commit()
        assert remind_project(db,project,datetime(2026,11,3))['notified']==0
        assert db.query(DocumentNoticeReceipt).filter_by(project_document_id=did).count()==3
        assert db.query(Notification).filter_by(entity_id=did,ntype='document_deadline').count()==3
        assert db.query(NotificationOutbox).filter_by(ntype='document_deadline').count()==3
    current=next(row for row in view(client,admin_headers,pid)['documents'] if row['id']==did)
    assert current['effective_owner_id']=='user-worker-001'
    with SessionLocal() as db:
        db.get(User,'user-worker-001').permission_overrides={'project.manage_plan':False};db.commit()
    assert client.put(f'/api/projects/{pid}/documents/{did}/owner',headers=worker_headers,json={'owner_id':None}).status_code==403
    assert client.put(f'/api/projects/{pid}/documents/{did}/owner',headers=admin_headers,json={'owner_id':'user-worker-002'}).status_code==400
    assigned=client.put(f'/api/projects/{pid}/documents/{did}/owner',headers=admin_headers,json={'owner_id':'user-worker-001','version':current['version']})
    assert assigned.status_code==200 and assigned.json()['owner_id']=='user-worker-001',assigned.text


def test_reports_run_on_local_sunday_once_and_scheduler_catchup_uses_real_db(client,admin_headers,monkeypatch):
    from app.services import report_scheduler as schedule
    from app.services import document_deadlines as deadlines
    from apscheduler.schedulers.background import BackgroundScheduler
    pid,rid=weekly_fixture(client,admin_headers)
    monkeypatch.setattr(schedule.config,'REPORT_WEEKLY_TIME','20:00')
    with SessionLocal() as db:
        assert schedule.generate_and_notify_weekly_report(db,now=datetime(2026,11,1,19,59))['notified']==0
        assert schedule.generate_and_notify_weekly_report(db,now=datetime(2026,11,1,20))['notified']==1;db.commit()
        assert schedule.generate_and_notify_weekly_report(db,now=datetime(2026,11,1,20,1))['notified']==0
        receipt=db.query(ProjectReportReceipt).filter_by(project_id=pid).one()
        assert receipt.period_key=='2026-W44'
        assert {'document_summary','document_submissions','vehicle_progress_changes'}<=receipt.summary.keys()
    monkeypatch.setattr(deadlines,'utcnow',lambda:datetime(2026,11,2,1))
    monkeypatch.setattr(BackgroundScheduler,'start',lambda self:None)
    scheduler=schedule.start_scheduler()
    jobs={job.id:job for job in scheduler.get_jobs()}
    assert {'weekly_report','document_deadlines'}<=jobs.keys()
    assert str(jobs['weekly_report'].trigger)=='interval[0:01:00]'
    with SessionLocal() as db:
        assert db.query(ProjectDocument).filter_by(definition_id=rid).count()==2
        assert db.query(ProjectReportReceipt).filter_by(project_id=pid).count()==1


def test_quantity_backlog_does_not_become_overdue_and_on_demand_is_separate(client,admin_headers):
    from .test_startup_targets import setup_project
    from app.models import ScheduleEntry
    pid,tasks=setup_project(client,admin_headers)
    with SessionLocal() as db:
        db.add(ScheduleEntry(project_id=pid,task_id=tasks[0]['id'],work_date=date(2020,1,1),planned_quantity=100))
        db.commit()
    result=client.get(f'/api/projects/{pid}/overdue',headers=admin_headers)
    assert result.status_code==200 and result.json()['count']==0,result.text
    assert view(client,admin_headers,pid)['document_actions']['on_demand']


@pytest.mark.parametrize('second',[datetime(2026,11,1,12),datetime(2026,10,1,12)])
def test_frozen_or_regressing_record_clock_keeps_latest_file_and_order(client,admin_headers,monkeypatch,second):
    import app.services.project_documents as service
    pid=create(client,admin_headers)['id'];did=view(client,admin_headers,pid)['documents'][0]['id']
    monkeypatch.setattr(service,'utcnow',lambda:datetime(2026,11,1,12))
    first=upload(client,admin_headers,pid,[did],data=b'first recording');assert first.status_code==200
    monkeypatch.setattr(service,'utcnow',lambda:second)
    last=upload(client,admin_headers,pid,[did],data=b'latest recording');assert last.status_code==200,last.text
    row=next(row for row in view(client,admin_headers,pid)['documents'] if row['id']==did)
    assert [version['uploaded_at'] for version in row['versions']]==['2026-11-01T12:00:00+00:00','2026-11-01T12:00:00.000001+00:00']
    assert client.get('/api'+row['latest_version']['download_url'],headers=admin_headers).content==b'latest recording'
    assert client.get('/api'+row['versions'][0]['download_url'],headers=admin_headers).content==b'first recording'


def test_new_version_follows_an_undated_legacy_original_on_both_databases(client,admin_headers):
    from app.models import Attachment,DocumentVersion
    pid=create(client,admin_headers)['id'];did=view(client,admin_headers,pid)['documents'][0]['id']
    assert upload(client,admin_headers,pid,[did],data=b'undated original').status_code==200
    with SessionLocal() as db:
        version=db.query(DocumentVersion).filter_by(project_document_id=did).one();version.uploaded_at=None
        db.get(Attachment,version.attachment_id).uploaded_at=None;db.get(ProjectDocument,did).submitted_at=None;db.commit()
    assert upload(client,admin_headers,pid,[did],data=b'new dated recording').status_code==200
    row=next(row for row in view(client,admin_headers,pid)['documents'] if row['id']==did)
    assert row['versions'][0]['uploaded_at'] is None and row['versions'][1]['uploaded_at'] is not None
    assert client.get('/api'+row['latest_version']['download_url'],headers=admin_headers).content==b'new dated recording'
    import hashlib,io,json
    from zipfile import ZipFile
    from openpyxl import load_workbook
    packaged=client.post(f'/api/projects/{pid}/documents/package',headers=admin_headers,json={'document_ids':[did]})
    assert packaged.status_code==200,packaged.text
    with ZipFile(io.BytesIO(packaged.content)) as bundle:
        manifest=json.loads(bundle.read('manifest.json'))
        assert bundle.read(manifest['files'][0]['archive_path'])==b'new dated recording'
    exported=client.get(f'/api/projects/{pid}/export/excel',headers=admin_headers)
    assert exported.status_code==200,exported.text
    sheet=load_workbook(io.BytesIO(exported.content))['交付资料']
    keys=[cell.value for cell in sheet[2]]
    versions=[dict(zip(keys,values)) for values in sheet.iter_rows(min_row=3,values_only=True)]
    versions=[item for item in versions if item['资料ID']==did]
    assert len(versions)==2 and versions[0]['提交时间（UTC）'] is None
    assert versions[-1]['SHA-256']==hashlib.sha256(b'new dated recording').hexdigest()


def test_confidential_reminder_respects_visibility_and_owner_failure_is_visible(client,admin_headers):
    pid=create(client,admin_headers)['id'];did=view(client,admin_headers,pid)['documents'][0]['id']
    with SessionLocal() as db:
        project=db.get(Project,pid);project.status='preparing';project.technical_manager='user-worker-001';project.members=['user-worker-001']
        row=db.get(ProjectDocument,did);row.owner_id='user-worker-001';row.confidential=True;row.due_at=datetime(2026,11,1)
        remind_project(db,project,datetime(2026,10,31,12));db.commit()
        recipients={item.user_id for item in db.query(Notification).filter_by(entity_id=did,ntype='document_deadline')}
        assert recipients=={project.manager}
        from app.services.document_deadlines import owner_details,_users
        effective=owner_details(project,row,_users(db,project,[row]))
        assert effective['effective_owner_id']==project.manager
        assert '保密' in effective['owner_warning']
        db.get(User,'user-worker-001').is_active=False;db.commit()
    item=next(row for row in view(client,admin_headers,pid)['documents'] if row['id']==did)
    assert item['effective_owner_id']=='user-admin-001' and item['owner_warning']


def test_two_scheduler_transactions_cannot_duplicate_reminder_receipts(client,admin_headers):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier
    from app.services.record_lock import lock_record
    pid=create(client,admin_headers)['id'];did=view(client,admin_headers,pid)['documents'][0]['id']
    with SessionLocal() as db:
        p=db.get(Project,pid);p.status='preparing';row=db.get(ProjectDocument,did);row.due_at=datetime(2026,11,1);db.commit()
    barrier=Barrier(2)
    def run(_):
        with SessionLocal() as db:
            project=db.get(Project,pid);barrier.wait(timeout=10)
            lock_record(db,project,conflict_detail='test project lock')
            result=remind_project(db,project,datetime(2026,10,31,12));db.commit();return result['notified']
    with ThreadPoolExecutor(max_workers=2) as pool:assert sum(pool.map(run,range(2)))==1
    with SessionLocal() as db:
        assert db.query(DocumentNoticeReceipt).filter_by(project_document_id=did).count()==1
        assert db.query(Notification).filter_by(entity_id=did,ntype='document_deadline').count()==1


def test_unassigned_reminder_never_bypasses_explicit_read_denial(client,admin_headers):
    pid=create(client,admin_headers)['id'];did=view(client,admin_headers,pid)['documents'][0]['id']
    with SessionLocal() as db:
        project=db.get(Project,pid);project.status='preparing';project.manager=None;project.technical_manager=None;project.members=[]
        row=db.get(ProjectDocument,did);row.owner_id=None;row.due_at=datetime(2026,11,1)
        db.get(User,'user-admin-001').permission_overrides={'project.read':False}
        db.flush()
        assert remind_project(db,project,datetime(2026,10,31,12))=={'notified':0,'owner_missing':1}
        db.commit()
        assert db.query(Notification).filter_by(entity_id=did,ntype='document_deadline').count()==0
        receipt=db.query(DocumentNoticeReceipt).filter_by(project_document_id=did).one()
        assert receipt.recipient_id is None and receipt.kind=='owner_missing'


def test_missing_report_manager_warns_once_and_retries_after_assignment(client,admin_headers):
    from app.services.report_scheduler import generate_and_notify_weekly_report
    pid,_=weekly_fixture(client,admin_headers)
    with SessionLocal() as db:
        project=db.get(Project,pid);project.manager='missing-original-manager';db.commit()
        assert generate_and_notify_weekly_report(db,now=datetime(2026,11,1,20))['notified']==0;db.commit()
        pending=db.query(ProjectReportReceipt).filter_by(project_id=pid).one()
        assert pending.delivered_at is None and pending.summary['delivery_pending']
        assert generate_and_notify_weekly_report(db,now=datetime(2026,11,1,20,1))['notified']==0;db.commit()
        assert db.query(Notification).filter(Notification.project_id==pid,Notification.title.contains('周报负责人需核对')).count()==1
        project.manager='user-admin-001';db.commit()
        assert generate_and_notify_weekly_report(db,now=datetime(2026,11,1,20,2))['notified']==1;db.commit()
        delivered=db.query(ProjectReportReceipt).filter_by(project_id=pid).one()
        assert delivered.delivered_at==datetime(2026,11,1,20,2)
        assert delivered.summary['previous_delivery_warning']['delivery_pending']


def test_weekly_report_restart_recovers_every_period_after_its_real_cursor(client,admin_headers):
    from app.services.report_scheduler import generate_and_notify_weekly_report
    pid,_=weekly_fixture(client,admin_headers)
    with SessionLocal() as db:
        assert generate_and_notify_weekly_report(db,now=datetime(2026,11,1,20))['notified']==1;db.commit()
        assert generate_and_notify_weekly_report(db,now=datetime(2026,11,22,20))['notified']==3;db.commit()
        assert {row.period_key for row in db.query(ProjectReportReceipt).filter_by(project_id=pid)}=={'2026-W44','2026-W45','2026-W46','2026-W47'}
        assert generate_and_notify_weekly_report(db,now=datetime(2026,11,22,21))['notified']==0


def test_project_plan_edit_updates_document_deadline_without_moving_schedule(client,admin_headers):
    from .test_startup_targets import setup_project
    from app.models import ScheduleEntry
    pid,tasks=setup_project(client,admin_headers)
    with SessionLocal() as db:
        entry=ScheduleEntry(project_id=pid,task_id=tasks[0]['id'],work_date=date(2026,11,1),planned_quantity=1)
        db.add(entry);db.commit();identity=entry.id
    result=client.put('/api/projects/'+pid,headers=admin_headers,json={'event_plans':{'ship':'2026-11-20'}})
    assert result.status_code==200,result.text
    document=next(row for row in view(client,admin_headers,pid)['documents'] if row['code']=='5.1.1')
    assert document['due_at']=='2026-11-12T16:00:00+00:00' and document['due_is_estimated']
    with SessionLocal() as db:
        entry=db.get(ScheduleEntry,identity)
        assert entry.work_date==date(2026,11,1) and entry.planned_quantity==1


def test_lifetime_waiting_for_actual_kickoff_blocks_only_archive(client,admin_headers):
    from app.services.lifecycle_gates import _pending_required_documents
    pid=create(client,admin_headers)['id']
    with SessionLocal() as db:
        for row in db.query(ProjectDocument).filter_by(project_id=pid):row.status='waived';row.waived_reason='聚焦全周期门禁'
        db.commit()
        assert _pending_required_documents(db,pid)==[]
        assert any('9.2' in item and '内部启动会' in item for item in _pending_required_documents(db,pid,archiving=True))
        assert not db.query(ProjectDocument).filter_by(project_id=pid,timing_type='lifetime').first()


def test_real_vehicle_arrival_api_keeps_date_precision_and_recomputes_clear(client,admin_headers):
    pid=create(client,admin_headers,timezone='America/New_York')['id']
    vehicle=client.post(f'/api/projects/{pid}/vehicles',headers=admin_headers,json={'vehicle_type':'潜伏式','model_name':'AGV','quantity':1}).json()['id']
    unit=client.post(f'/api/projects/{pid}/vehicle-units',headers=admin_headers,json={'vehicle_id':vehicle,'asset_code':'NY-01'}).json()['id']
    assert client.put(f'/api/projects/{pid}/vehicle-arrivals',headers=admin_headers,json={'unit_ids':[unit],'arrived_at':'2026-11-01'}).status_code==200
    row=next(row for row in view(client,admin_headers,pid)['documents'] if row['code']=='6.1')
    assert row['due_at']=='2026-11-02T04:00:00+00:00' and not row['due_is_estimated']
    assert row['due_basis']['actual_date']=='2026-11-01' and row['due_basis']['precision']=='date'
    assert 'actual_at' not in row['due_basis']
    assert client.put(f'/api/projects/{pid}/vehicle-arrivals',headers=admin_headers,json={'unit_ids':[unit],'arrived_at':None,'reason':'到场登记待核对'}).status_code==200
    cleared=next(item for item in view(client,admin_headers,pid)['documents'] if item['id']==row['id'])
    assert cleared['due_at'] is None and cleared['due_state']=='awaiting_event'


def test_document_get_does_not_generate_missed_periods_or_business_writes(client,admin_headers):
    from app.models import ChangeLog
    pid,rid=weekly_fixture(client,admin_headers)
    with SessionLocal() as db:
        before=(db.query(ProjectDocument).count(),db.query(ChangeLog).filter(ChangeLog.action!='confidential_view').count(),db.query(DocumentNoticeReceipt).count())
    assert view(client,admin_headers,pid)['document_actions']
    with SessionLocal() as db:
        assert (db.query(ProjectDocument).count(),db.query(ChangeLog).filter(ChangeLog.action!='confidential_view').count(),db.query(DocumentNoticeReceipt).count())==before
        assert not db.query(ProjectDocument).filter_by(definition_id=rid).first()


def test_scheduler_receipts_survive_encrypted_backup_and_are_purged_with_project(client,admin_headers):
    import json
    from .test_project_recycle_bin import _delete
    from .test_retention_unification import prepare,purge
    from app.services.project_deletion import backup_zip,_backup_path
    from app.models import ProjectPermanentBackup
    project=create(client,admin_headers);pid=project['id'];did=view(client,admin_headers,pid)['documents'][0]['id']
    with SessionLocal() as db:
        notice=DocumentNoticeReceipt(project_id=pid,project_document_id=did,recipient_id='user-admin-001',
            recipient_key='user-admin-001',kind='due_soon',due_at=datetime(2026,11,1),due_is_estimated=False)
        report=ProjectReportReceipt(project_id=pid,period_key='2026-W44',recipient_id='user-admin-001',
            delivery_pending=False,delivered_at=datetime(2026,11,1,20),summary={'documents_submitted':1})
        db.add_all([notice,report]);db.commit();notice_id,report_id=notice.id,report.id
    deleted=_delete(client,admin_headers,project)
    assert deleted.status_code==200,deleted.text
    row=client.get(f'/api/project-recycle-bin/projects/{pid}',headers=admin_headers).json()
    prepared=prepare(client,admin_headers,row)
    assert prepared.status_code==200,prepared.text
    receipt=prepared.json()
    with SessionLocal() as db:
        path=_backup_path(db.get(ProjectPermanentBackup,receipt['backup_id']).backup_name)
        with backup_zip(path) as bundle:
            records=json.loads(bundle.read('project-backup.json'))['records']
        assert [item['id'] for item in records['document_notice_receipts']]==[notice_id]
        assert [item['id'] for item in records['project_report_receipts']]==[report_id]
    result=purge(client,admin_headers,row,receipt)
    assert result.status_code==200,result.text
    with SessionLocal() as db:
        assert db.get(Project,pid) is None
        assert db.get(DocumentNoticeReceipt,notice_id) is None and db.get(ProjectReportReceipt,report_id) is None
