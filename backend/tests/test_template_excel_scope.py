import io
from openpyxl import Workbook, load_workbook
from app.db import SessionLocal
from app.models import Phase, Module, Task


def test_template_excel_selected_scope(client, admin_headers):
    created = client.post('/api/templates', headers=admin_headers, json={'name': '自定义范围', 'copy_from': 'default'})
    assert created.status_code == 200
    tid = created.json()['id']
    with SessionLocal() as db:
        task = db.query(Task).filter(Task.template_id == tid).first()
        task.name = '只在自定义模板'; db.commit()
        default_ids = {x.id for x in db.query(Task).filter(Task.project_id.is_(None), Task.template_id.is_(None)).all()}
    custom = client.get('/api/templates/export/excel', params={'template_id': tid}, headers=admin_headers)
    default = client.get('/api/templates/export/excel', headers=admin_headers)
    assert custom.status_code == default.status_code == 200
    text = lambda b: str(list(load_workbook(io.BytesIO(b)).active.values))
    assert '只在自定义模板' in text(custom.content)
    assert '只在自定义模板' not in text(default.content)
    imported = client.post('/api/templates/import/excel', params={'template_id': tid}, headers=admin_headers, files={'file': ('template.xlsx', custom.content)})
    assert imported.status_code == 200, imported.text
    with SessionLocal() as db:
        assert default_ids == {x.id for x in db.query(Task).filter(Task.project_id.is_(None), Task.template_id.is_(None)).all()}
        assert db.query(Task).filter(Task.template_id == tid, Task.name == '只在自定义模板').count() == 1
        assert db.query(Phase).filter(Phase.template_id == tid).count() > 0
        assert db.query(Module).filter(Module.template_id == tid).count() > 0


def test_empty_excel_cannot_wipe_existing_template(client, admin_headers):
    with SessionLocal() as db:
        before = {x.id for x in db.query(Task).filter(Task.project_id.is_(None)).all()}
    wb = Workbook(); content = io.BytesIO(); wb.save(content)
    rejected = client.post('/api/templates/import/excel', headers=admin_headers, files={'file': ('empty.xlsx', content.getvalue())})
    assert rejected.status_code == 400
    assert '未修改' in rejected.json()['detail']
    with SessionLocal() as db:
        assert before == {x.id for x in db.query(Task).filter(Task.project_id.is_(None)).all()}
    assert client.get('/api/templates/export/excel', params={'template_id': 'missing'}, headers=admin_headers).status_code == 404
