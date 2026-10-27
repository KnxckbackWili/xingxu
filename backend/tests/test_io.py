"""Import/export + clone tests."""

from tests.lifecycle_helpers import versioned
import io
import json


def test_export_template_excel(client, pm_headers):
    r = client.get("/api/templates/export/excel", headers=pm_headers)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert len(r.content) > 1000


def test_import_export_template_roundtrip(client, pm_headers, admin_headers):
    """Export → import → tree should survive."""
    r = client.get("/api/templates/export/excel", headers=pm_headers)
    assert r.status_code == 200
    data = r.content

    r = client.post("/api/templates/import/excel", headers=admin_headers,
                    files={"file": ("template.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.json()
    result = r.json()
    assert result["tasks"] > 30

    # Tree still has content
    r = client.get("/api/templates/default/tree", headers=pm_headers)
    assert r.status_code == 200
    assert len(r.json()) == 9  # C2-2：模板树完整返回 9 阶段（含项目准备）


def test_project_export_json(client, pm_headers):
    # Build a project
    r = client.post("/api/projects", headers=pm_headers, json={
        "name": "导出测试项目", "client": "客户X",
        "start_date": "2026-08-10", "new_vehicles": False, "end_date": "2026-09-30"})
    pid = r.json()["id"]
    client.post(f"/api/projects/{pid}/vehicles", headers=pm_headers,
                json={"vehicle_type": "潜伏式", "model_name": "AGV-500", "quantity": 3})
    client.post(f"/api/projects/{pid}/tasks/generate", headers=pm_headers,
                json={"module_ids": ["mod-02", "mod-03"]})

    r = client.get(f"/api/projects/{pid}/export/json", headers=pm_headers)
    assert r.status_code == 200
    snapshot = r.json()
    assert snapshot["type"] == "agvpm-project"
    assert len(snapshot["vehicles"]) == 1
    assert len(snapshot["modules"]) == 2

    # Import into a new project
    r = client.post("/api/projects/import/json", headers=pm_headers, json=snapshot)
    assert r.status_code == 200
    new_id = r.json()["id"]

    # Verify new project has vehicles + tasks
    r = client.get(f"/api/projects/{new_id}", headers=pm_headers)
    assert r.json()["vehicle_total"] == 3
    r = client.get(f"/api/projects/{new_id}/tree", headers=pm_headers)
    assert len(r.json()) == 2


def test_clone_project(client, pm_headers):
    r = client.post("/api/projects", headers=pm_headers, json={
        "name": "克隆源项目", "client": "客户Y",
        "start_date": "2026-08-10", "new_vehicles": False, "end_date": "2026-09-30"})
    pid = r.json()["id"]
    client.post(f"/api/projects/{pid}/vehicles", headers=pm_headers,
                json={"vehicle_type": "堆高车", "model_name": "Fork-1500", "quantity": 2})
    client.post(f"/api/projects/{pid}/tasks/generate", headers=pm_headers,
                json={"module_ids": ["mod-03"]})

    r = client.post(f"/api/projects/{pid}/clone", headers=pm_headers, json={
        "name": "克隆目标项目（二期）",
        "start_date": "2026-10-01", "new_vehicles": False,
        "end_date": "2026-11-30",
    })
    assert r.status_code == 200
    new_id = r.json()["id"]

    r = client.get(f"/api/projects/{new_id}", headers=pm_headers)
    data = r.json()
    assert data["name"] == "克隆目标项目（二期）"
    assert data["vehicle_total"] == 2
    assert data["start_date"] == "2026-10-01"

    # Tasks copied
    r = client.get(f"/api/projects/{new_id}/tree", headers=pm_headers)
    tasks = [t for ph in r.json() for m in ph["modules"] for t in m["tasks"]]
    assert len(tasks) == 5  # 单车调试 has 5 tasks


def test_project_export_excel(client, pm_headers):
    r = client.post("/api/projects", headers=pm_headers, json={
        "name": "Excel导出项目",
        "start_date": "2026-08-10", "new_vehicles": False, "end_date": "2026-09-30"})
    pid = r.json()["id"]
    client.post(f"/api/projects/{pid}/tasks/generate", headers=pm_headers,
                json={"module_ids": ["mod-03"]})
    r = client.get(f"/api/projects/{pid}/export/excel", headers=pm_headers)
    assert r.status_code == 200
    assert len(r.content) > 500


def test_project_export_excel_sheets(client, pm_headers, worker_headers):
    """导出完整项目工作簿，包括现场物品。"""
    from openpyxl import load_workbook

    # 建项目 + 车型 + 任务 + 排程
    r = client.post("/api/projects", headers=pm_headers, json={
        "name": "7Sheet导出项目",
        "start_date": "2026-08-10", "new_vehicles": False, "end_date": "2026-09-30"})
    pid = r.json()["id"]
    assert client.post(f"/api/projects/{pid}/status", headers=pm_headers,
                       json={"status": "preparing"}).status_code == 200
    client.post(f"/api/projects/{pid}/vehicles", headers=pm_headers,
                json={"vehicle_type": "潜伏式", "model_name": "AGV-500", "quantity": 2})
    client.post(f"/api/projects/{pid}/tasks/generate", headers=pm_headers,
                json={"module_ids": ["mod-03"]})
    client.post(f"/api/projects/{pid}/schedule/generate", headers=pm_headers)

    # worker 加入项目（成员校验）后填报一条（填报流水 + 日程已完成数）
    r = client.get("/api/auth/me", headers=worker_headers)
    wid = r.json()["id"]
    r = client.get(f"/api/projects/{pid}", headers=pm_headers)
    client.put(f"/api/projects/{pid}", headers=pm_headers,
               json=versioned(pid, {"members": [*r.json()["members"], wid]}))
    # 下达后现场成员填报，导出断言仍验证真实填报与日程关联。
    r = client.post(f"/api/projects/{pid}/schedule/confirm", headers=pm_headers)
    assert r.status_code == 200, r.json()
    r = client.get(f"/api/projects/{pid}/schedule", headers=pm_headers)
    entry = r.json()[0]
    r = client.post("/api/worker/reports", headers=worker_headers, json={
        "task_id": entry["task_id"], "schedule_entry_id": entry["id"],
        "work_date": entry["work_date"], "completed_quantity": 2,
        "client_uuid": "io-export-sheet-0001"})
    assert r.status_code == 200, r.json()

    # 一个问题（含分类专属字段）
    client.post(f"/api/projects/{pid}/issues", headers=pm_headers, json={
        "title": "充电桩故障", "issue_type": "hardware",
        "extra": {"device": "AGV-03", "affects_schedule": "影响"}})

    r = client.get(f"/api/projects/{pid}/export/excel", headers=pm_headers)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ['交付概览','生命周期表','问题清单','物资清单','收发流水','交付资料']
    overview = {row[1]:row[2] for row in wb['交付概览'].iter_rows(min_row=3, values_only=True) if row[0]=='项目概况'}
    assert {'项目状态','执行任务数','未闭环问题','整体进度 %'} <= set(overview)
    from app.services.report_forms import LIFECYCLE_HEADERS
    assert [cell.value for cell in wb['生命周期表'][2]][:len(LIFECYCLE_HEADERS)] == LIFECYCLE_HEADERS
    headers = [cell.value for cell in wb['问题清单'][2]]
    assert '负责人' in headers and '解决方案/关闭或作废说明' in headers and '扩展字段' in headers
    assert any('AGV-03' in str(cell.value) for row in wb['问题清单'].iter_rows(min_row=3) for cell in row if cell.value)
    assert '更换配件' in [cell.value for cell in wb['生命周期表'][2]]
    assert [cell.value for cell in wb['物资清单'][1]][:10] == ['物品编码','名称','规格','序列号','归属','现存数量','单位','位置','保管人','说明']
    assert [cell.value for cell in wb['收发流水'][1]][:10] == ['发生时间（UTC）','时间依据','物品','动作','数量','结余','原位置','新位置','经手人','说明']
    assert {'文件名','资料状态','提交人','提交时间（UTC）','审核信息'} <= set(cell.value for cell in wb['交付资料'][2])


def test_project_export_package_zip(client, pm_headers, worker_headers):
    """HYPE-58：项目整包 ZIP（xlsx + json + files/ + manifest，含去重/SHA-256/缺失标记），worker 无导出权限 403。"""
    import zipfile
    import json as stdjson

    pid = client.post("/api/projects", headers=pm_headers, json={
        "name": "整包导出项目",
        "start_date": "2026-08-10", "new_vehicles": False, "end_date": "2026-09-30"}).json()["id"]
    client.post(f"/api/projects/{pid}/vehicles", headers=pm_headers,
                json={"vehicle_type": "潜伏式", "model_name": "AGV-500", "quantity": 2})
    client.post(f"/api/projects/{pid}/tasks/generate", headers=pm_headers,
                json={"module_ids": ["mod-03"]})

    # 同名附件 × 2 → ZIP 内 files/ 自动去重（-1 序号），manifest 记录 sha256/size
    for content in (b"hello-pkg-1", b"hello-pkg-2"):
        r = client.post(f"/api/projects/{pid}/documents/files", headers=pm_headers,
                        files={"file": ("same.txt", content, "text/plain")})
        assert r.status_code == 200, r.json()

    r = client.get(f"/api/projects/{pid}/export/package", headers=pm_headers)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/zip")

    zf = zipfile.ZipFile(io.BytesIO(r.content))
    names = zf.namelist()
    assert "项目数据.xlsx" in names
    assert "项目备份.json" in names
    assert "manifest.json" in names
    files_in = [n for n in names if n.startswith("files/")]
    assert len(files_in) == 2, files_in

    manifest = stdjson.loads(zf.read("manifest.json"))
    assert manifest["type"] == "agvpm-project-package"
    assert manifest["project"]["id"] == pid
    assert manifest["file_count"] == 2
    assert manifest["missing_count"] == 0
    paths = [f["archive_path"] for f in manifest["files"]]
    assert len(set(paths)) == len(paths)  # 去重：archive_path 唯一
    for f in manifest["files"]:
        assert f["id"]  # 清单记录附件 id
        assert f["sha256"]
        assert f["size"] > 0
        assert f["archive_path"].startswith("files/")
        # 文件本体与 manifest 的 SHA-256 一致
        disk = zf.read(f["archive_path"])
        import hashlib
        assert hashlib.sha256(disk).hexdigest() == f["sha256"]

    # 同一项目 Excel 接口仍兼容（不删旧接口）
    assert client.get(f"/api/projects/{pid}/export/excel", headers=pm_headers).status_code == 200
    assert client.get(f"/api/projects/{pid}/export/json", headers=pm_headers).status_code == 200

    # worker 加入项目（可见）但无 project.export 能力 → 403
    me = client.get("/api/auth/me", headers=worker_headers).json()
    detail = client.get(f"/api/projects/{pid}", headers=pm_headers).json()
    client.put(f"/api/projects/{pid}", headers=pm_headers,
               json=versioned(pid, {"members": [*detail["members"], me["id"]]}))
    assert client.get(f"/api/projects/{pid}/export/package", headers=worker_headers).status_code == 403


def test_project_export_package_zip_missing_file(client, pm_headers):
    """磁盘附件缺失时 ZIP 不中断，manifest 显式标错（missing=true），不能静默漏掉。"""
    import zipfile
    import json as stdjson
    from app.config import UPLOAD_DIR

    pid = client.post("/api/projects", headers=pm_headers, json={
        "name": "缺文件导出项目",
        "start_date": "2026-08-10", "new_vehicles": False, "end_date": "2026-09-30"}).json()["id"]
    r = client.post(f"/api/projects/{pid}/documents/files", headers=pm_headers,
                    files={"file": ("关键点位表.xlsx", b"\x50\x4b\x03\x04fake-xlsx", "application/vnd.ms-excel")})
    assert r.status_code == 200, r.json()

    # 删除磁盘文件本体（模拟备份缺失），数据库记录仍在
    for f in UPLOAD_DIR.iterdir():
        f.unlink()

    r = client.get(f"/api/projects/{pid}/export/package", headers=pm_headers)
    assert r.status_code == 200
    zf = zipfile.ZipFile(io.BytesIO(r.content))
    manifest = stdjson.loads(zf.read("manifest.json"))
    assert manifest["file_count"] == 1
    assert manifest["missing_count"] == 1
    entry = manifest["files"][0]
    assert entry["name"] == "关键点位表.xlsx"
    assert entry["missing"] is True
    assert "sha256" not in entry  # 文件缺失无法计算哈希
    assert not any(n.startswith("files/") for n in zf.namelist())  # 缺失文件不进包体
    # xlsx / json / manifest 本体仍在，导出可用
    assert "项目数据.xlsx" in zf.namelist()
    assert "项目备份.json" in zf.namelist()


def test_template_json_export_route(client, pm_headers):
    """模板中心调用的 /templates/export/json（修复 404）。"""
    r = client.get("/api/templates/export/json", headers=pm_headers)
    assert r.status_code == 200
    assert len(r.json()["phases"]) == 9
    assert any(p["name"] == "项目准备" for p in r.json()["phases"])  # C2-2：完整返回含项目准备


def test_template_tree_route(client, pm_headers):
    """模板中心加载用的 /templates/tree（修复 404）。"""
    r = client.get("/api/templates/tree", headers=pm_headers)
    assert r.status_code == 200
    assert len(r.json()) == 9
    assert any(p["name"] == "项目准备" for p in r.json())  # C2-2：完整返回含项目准备


def test_ai_quick_import_template_preview_and_full_import(client, pm_headers):
    template_response = client.get("/api/projects/import/template", headers=pm_headers)
    assert template_response.status_code == 200
    payload = template_response.json()
    assert payload["type"] == "agvpm-project"
    assert payload["version"] == "2.3"
    assert "client_contacts" in payload["project"]
    assert "vehicle_units" in payload
    payload["project"].update({
        "name": "AI 导入测试项目", "client": "测试客户", "start_date": "2026-09-01", "end_date": "2026-12-31",
        "client_contacts": [{"name": "张经理", "department": "设备部", "is_primary": True}],
        "entry_plan": {**payload["project"]["entry_plan"], "mode": "collective", "network_access": "甲方分配 IP"},
    })
    payload["vehicles"] = [{"vehicle_type": "潜伏式", "model_name": "AGV-500", "quantity": 1, "configuration": {"navigation": "SLAM"}}]
    payload["vehicle_units"] = [{
        "asset_code": "AGV-001", "vehicle_type": "潜伏式", "model_name": "AGV-500", "current_stage": "on_site", "current_status": "attention",
        "events": [{"event_type": "fault", "event_date": "2026-09-20", "title": "传感器报警", "action_taken": "重新压接", "verification": "运行正常", "is_closed": True, "replaced_parts": [{"name": "M12 接插件", "quantity": 1}]}],
    }]
    payload["modules"] = [{"phase_name": "现场实施", "name": "网络联调", "tasks": [{"name": "完成现场网络联调", "quantity_type": "manual", "target_quantity": 1, "unit": "项", "deliverables": [{"name": "测试记录", "required": True}], "check_items": [{"name": "网络连通", "acceptance_standard": "无丢包"}]}]}]
    payload["schedule"] = [{"task_name": "完成现场网络联调", "work_date": "2026-09-21", "planned_quantity": 1, "status": "planned"}]
    payload["inventory_items"] = [{
        "item_code": "CLIENT-PC-01", "name": "甲方工控机", "category": "equipment",
        "ownership": "client", "tracking_mode": "serial", "serial_number": "PC-SN-01",
        "quantity": 1, "unit": "台", "status": "on_site", "location": "控制室", "custodian": "张经理",
    }]

    preview = client.post("/api/projects/import/preview", headers=pm_headers, json=payload)
    assert preview.status_code == 200
    assert preview.json()["valid"] is True
    assert preview.json()["summary"]["vehicle_issues"] == 1
    assert preview.json()["summary"]["tasks"] == 1
    assert preview.json()["summary"]["inventory_items"] == 1

    imported = client.post("/api/projects/import/json", headers=pm_headers, json=payload)
    assert imported.status_code == 200, imported.text
    project_id = imported.json()["id"]
    project = client.get(f"/api/projects/{project_id}", headers=pm_headers).json()
    assert project["client_contacts"][0]["name"] == "张经理"
    assert project["entry_plan"]["mode"] == "collective"
    assert project["entry_plan"]["network_access"] == "甲方分配 IP"

    units = client.get(f"/api/projects/{project_id}/vehicle-units", headers=pm_headers).json()
    assert units[0]["asset_code"] == "AGV-001"
    vehicle_issues = client.get(f"/api/projects/{project_id}/vehicle-issues", headers=pm_headers).json()
    assert vehicle_issues[0]["status"] == "resolved"
    assert vehicle_issues[0]["replaced_parts"][0]["name"] == "M12 接插件"

    tree = client.get(f"/api/projects/{project_id}/tree", headers=pm_headers).json()
    tasks = [task for phase in tree for module in phase["modules"] for task in module["tasks"]]
    assert any(task["name"] == "完成现场网络联调" for task in tasks)
    assert 'deliverables' not in next(task for task in tasks if task['name'] == '完成现场网络联调')
    source = client.get(f'/api/projects/{project_id}/documents', headers=pm_headers).json()
    assert any(row['name'] == '测试记录' and row['stage_key'] == 'legacy' for row in source['definitions'])
    assert all(row['status'] != 'submitted' for row in source['documents'])
    schedule = client.get(f"/api/projects/{project_id}/schedule", headers=pm_headers).json()
    assert any(item["planned_quantity"] == 1 for item in schedule)
    inventory = client.get(f"/api/projects/{project_id}/inventory", headers=pm_headers).json()
    assert inventory["items"][0]["ownership"] == "client"
    assert inventory["items"][0]["serial_number"] == "PC-SN-01"


def test_ai_quick_import_preview_reports_structural_errors(client, pm_headers):
    payload = {
        "type": "wrong-type",
        "project": {"name": "", "start_date": "2026-09-10", "end_date": "2026-09-01"},
        "vehicle_units": [{"asset_code": "AGV-001"}, {"asset_code": "AGV-001"}],
    }
    preview = client.post("/api/projects/import/preview", headers=pm_headers, json=payload).json()
    assert preview["valid"] is False
    assert len(preview["errors"]) >= 4


def test_ai_assistant_status_skill_and_mocked_organize(client, pm_headers, monkeypatch):
    status = client.get("/api/ai-projects/status", headers=pm_headers)
    assert status.status_code == 200
    assert "available" in status.json()
    assert {item["id"] for item in status.json()["providers"]} == {"deepseek", "mimo", "minimax", "custom"}
    assert {item["id"] for item in status.json()["skills"]} == {"project", "planning", "inventory", "issues", "handover"}
    assert not {"supported_files", "max_file_mb", "max_files"} & set(status.json())

    bundle = client.get("/api/ai-projects/skill-bundle", headers=pm_headers)
    assert bundle.status_code == 200
    assert bundle.content.startswith(b"PK")
    inventory_bundle = client.get("/api/ai-projects/skill-bundle?skill=inventory", headers=pm_headers)
    assert inventory_bundle.status_code == 200
    assert b"agv-site-inventory" in inventory_bundle.content

    valid_draft = client.get("/api/projects/import/template", headers=pm_headers).json()
    valid_draft["project"].update({"name": "AI 草稿", "start_date": "2026-09-01", "end_date": "2026-10-01"})
    selected = {}

    def mocked_organize(*_args, **_kwargs):
        selected["provider"] = _args[3]
        return {
        "draft": valid_draft, "missing_fields": [], "conflicts": [],
        "assumptions": [], "change_summary": ["创建项目基础信息"],
        }

    monkeypatch.setattr("app.routers.ai_projects.organize_project", mocked_organize)
    organized = client.post("/api/ai-projects/organize", headers=pm_headers, json={
        "source_text": "项目资料", "provider": "mimo",
    })
    assert organized.status_code == 200, organized.text
    assert selected["provider"] == "mimo"
    assert organized.json()["validation"]["valid"] is True
    assert organized.json()["draft"]["version"] == "2.3"

    job = client.post("/api/ai-projects/organize/jobs", headers=pm_headers, json={
        "source_text": "项目资料", "provider": "mimo", "skill": "project",
    })
    assert job.status_code == 200, job.text
    job_status = client.get(f"/api/ai-projects/organize/jobs/{job.json()['id']}", headers=pm_headers)
    assert job_status.status_code == 200
    assert job_status.json()["status"] == "completed"
    assert job_status.json()["percent"] == 100

    monkeypatch.setattr("app.routers.ai_projects.test_connection", lambda provider: {
        "ok": True, "provider": provider, "model": "mock-model",
    })
    tested = client.post("/api/ai-projects/test-connection", headers=pm_headers, json={"provider": "minimax"})
    assert tested.status_code == 200
    assert tested.json()["provider"] == "minimax"


def test_ai_response_parser_accepts_reasoning_wrapper():
    from app.services.ai_project_assistant import _extract_json

    result = _extract_json('<think>内部推理</think>\n说明文字\n{"draft": {"project": {}}, "missing_fields": []}\n完成')
    assert result["draft"] == {"project": {}}


def test_ai_request_falls_back_for_compatible_providers(monkeypatch):
    import io as stdlib_io
    import urllib.error
    from app import config
    from app.services import ai_project_assistant as assistant

    provider = {
        "id": "minimax", "name": "MiniMax", "base_url": "https://example.invalid/v1",
        "model": "MiniMax-M2.7", "api_key": "server-only", "allow_no_key": False,
    }
    monkeypatch.setitem(config.AI_PROVIDERS, "minimax", provider)
    request_bodies = []

    def mocked_request(_provider, body):
        request_bodies.append(body)
        if len(request_bodies) < 3:
            raise urllib.error.HTTPError(
                "https://example.invalid", 422, "unsupported response_format", {},
                stdlib_io.BytesIO(b'{"error":"unsupported response_format"}'),
            )
        return {"choices": [{"message": {"content": '<think>reasoning</think>{"draft":{"project":{}},"missing_fields":[],"conflicts":[],"assumptions":[],"change_summary":[]}'}}]}

    monkeypatch.setattr(assistant, "_perform_request", mocked_request)
    result = assistant._request_ai([{"role": "user", "content": "test"}], "minimax")
    assert result["draft"] == {"project": {}}
    assert request_bodies[0]["response_format"]["type"] == "json_schema"
    assert request_bodies[1]["response_format"]["type"] == "json_object"
    assert "response_format" not in request_bodies[2]


def test_ai_preview_rejects_nested_malformed_payload_without_500(client, pm_headers):
    template = client.get("/api/projects/import/template", headers=pm_headers).json()
    template["project"].update({"name": "异常结构", "start_date": "2026-09-01", "end_date": "2026-10-01"})
    template["modules"] = [{"name": "联调", "tasks": "not-an-array"}]
    template["vehicle_units"] = [{"asset_code": "AGV-01", "events": "not-an-array"}]
    template["schedule"] = ["not-an-object"]
    response = client.post("/api/projects/import/preview", headers=pm_headers, json=template)
    assert response.status_code == 200
    assert response.json()["valid"] is False
    assert any("必须是数组" in item for item in response.json()["errors"])


def test_quick_import_rejects_invalid_inventory_movements(client, pm_headers):
    template = client.get("/api/projects/import/template", headers=pm_headers).json()
    template["project"].update({"name": "坏流水", "start_date": "2026-09-01", "end_date": "2026-10-01"})
    template["inventory_items"] = [{
        "name": "测试工具", "quantity": 1, "minimum_quantity": -1,
        "movements": [{"movement_type": "magic", "quantity": -2, "occurred_at": "not-a-date"}],
    }]
    result = client.post("/api/projects/import/preview", headers=pm_headers, json=template).json()
    assert result["valid"] is False
    assert any("minimum_quantity" in error for error in result["errors"])
    assert any("movement_type" in error for error in result["errors"])
    assert any("occurred_at" in error for error in result["errors"])


def test_ai_domain_job_and_issue_bulk_apply(client, pm_headers, monkeypatch):
    project = client.post("/api/projects", headers=pm_headers, json={
        "name": "问题整理项目", "start_date": "2026-09-01", "end_date": "2026-10-01",
    }).json()

    def mocked_domain(skill, *_args, **_kwargs):
        assert skill == "issues"
        return {
            "draft": {"issues": [{
                "title": "充电桩无法握手", "description": "两台车均无法启动充电",
                "issue_type": "hardware", "severity": "high", "extra": {"scope": "2 台车"},
            }]},
            "missing_fields": [], "conflicts": [], "assumptions": [], "change_summary": [],
        }

    monkeypatch.setattr("app.routers.ai_projects.organize_domain", mocked_domain)
    created = client.post("/api/ai-projects/organize/jobs", headers=pm_headers, json={
        "source_text": "现场反馈充电异常", "provider": "deepseek", "skill": "issues",
        "project_id": project["id"],
    })
    assert created.status_code == 200, created.text
    job = client.get(f"/api/ai-projects/organize/jobs/{created.json()['id']}", headers=pm_headers).json()
    assert job["status"] == "completed"
    assert job["result"]["validation"]["valid"] is True
    applied = client.post(
        f"/api/projects/{project['id']}/issues/import", headers=pm_headers,
        json={"items": job["result"]["draft"]["issues"]},
    )
    assert applied.status_code == 200, applied.text
    rows = client.get(f"/api/projects/{project['id']}/issues", headers=pm_headers).json()
    assert rows[0]["title"] == "充电桩无法握手"
    assert rows[0]["severity"] == "high"


def test_quick_import_requires_initial_status_and_preserves_original_payload(client, pm_headers):
    """New project import cannot restore lifecycle authority without its evidence."""
    for status in ("draft", "completed", "archived"):
        payload = {
            "type": "agvpm-project", "version": "2.3",
            "project": {"name": f"导入恢复-{status}", "start_date": "2026-08-01", "end_date": "2026-09-30",
                        "status": status, "conditions": {},
                        "client_contacts": [], "entry_plan": {}},
            "vehicles": [], "vehicle_units": [], "modules": [], "schedule": [],
            "preparations": [], "inventory_items": [], "progress": [], "issues": [],
        }
        r = client.post("/api/projects/import/json", headers=pm_headers, json=payload)
        if status == "draft":
            assert r.status_code == 200, r.text
            project = client.get(f"/api/projects/{r.json()['id']}", headers=pm_headers).json()
            assert project["status"] == status
        else:
            assert r.status_code == 400, r.text
            assert "新建项目导入" in r.json()["detail"]
        assert payload["project"]["status"] == status
