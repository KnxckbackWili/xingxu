"""One real-file evidence stream replaces the former document sources."""
from tests.lifecycle_helpers import versioned
import hashlib
import copy
import io
import json
from zipfile import ZipFile
import pytest


def create(client,headers,**fields):
    result=client.post('/api/projects',headers=headers,json={'name':'资料项目','start_date':'2026-01-01','end_date':'2026-12-31',**fields})
    assert result.status_code==200,result.text
    return result.json()


def view(client,headers,pid):
    result=client.get(f'/api/projects/{pid}/documents',headers=headers)
    assert result.status_code==200,result.text
    return result.json()


def upload(client,headers,pid,ids=(),data=b'evidence',**fields):
    import json
    return client.post(f'/api/projects/{pid}/documents/files',headers=headers,files={'file':('evidence.txt',data,'text/plain')},data={'document_ids':json.dumps(list(ids)),**fields})


def test_new_project_snapshot_platform_and_real_event_seams(client,admin_headers):
    p=create(client,admin_headers,dispatch_platform='NDC');pid=p['id']
    state=view(client,admin_headers,pid)
    assert len([row for row in state['definitions'] if row['stage_key']!='legacy'])==53 and not any(row.get('is_group') for row in state['definitions'])
    assert not any(row['stage_key']=='legacy' for row in state['documents'])
    assert not any(row['code'] in {'8.3','8.4'} for row in state['documents'])
    assert any(row['code']=='8.2' for row in state['documents'])
    assert not any(row['timing_type'] in {'weekly','on_trigger','lifetime'} for row in state['documents'])
    assert any(row['timing_type']=='weekly' and row['waiting_reason'] for row in state['definitions'])
    assert client.put(f'/api/projects/{pid}/events/internal_kickoff',headers=admin_headers,json=versioned(pid, {'actual_at':'2026-10-01T01:02:03Z'})).status_code==200
    assert {row['code'] for row in view(client,admin_headers,pid)['documents']} >= {'9.1','9.2'}
    assert client.get(f'/api/projects/{pid}/archive-checklist',headers=admin_headers).status_code==404
    assert client.get(f'/api/projects/{pid}/attachments',headers=admin_headers).status_code==404
    assert client.get(f'/api/projects/{pid}/preparations',headers=admin_headers).status_code==404
    assert client.get(f'/api/projects/{pid}/procurements',headers=admin_headers).json()['items']==[]


def test_real_file_multi_submission_versions_waive_and_optional_review(client,admin_headers):
    p=create(client,admin_headers);pid=p['id'];docs=view(client,admin_headers,pid)['documents'][:2];ids=[row['id'] for row in docs]
    assert client.post(f'/api/projects/{pid}/documents/submit',headers=admin_headers,json={'document_ids':ids}).status_code==422
    assert client.post(f'/api/projects/{pid}/documents/submit',headers=admin_headers,json={'document_ids':ids,'attachment_id':'missing'}).status_code==400
    result=upload(client,admin_headers,pid,ids)
    assert result.status_code==200,result.text
    assert len(result.json()['documents'])==2 and all(row['status']=='submitted' for row in result.json()['documents'])
    asset=result.json()['file'];assert asset['sha256']==hashlib.sha256(b'evidence').hexdigest()
    assert view(client,admin_headers,pid)['unclassified_files']==[]
    document=result.json()['documents'][0];old_version=document['versions'][0]
    assert client.get('/api'+old_version['download_url'],headers=admin_headers).content==b'evidence'
    assert upload(client,admin_headers,pid,[document['id']],data=b'next').status_code==200
    current=next(row for row in view(client,admin_headers,pid)['documents'] if row['id']==document['id'])
    assert len(current['versions'])==2
    assert client.delete('/api/files/'+asset['id'],headers=admin_headers).status_code==400
    path=f'/api/projects/{pid}/documents/{current["id"]}'
    assert client.put(path,headers=admin_headers,json={'status':'waived','reason':' '}).status_code==400
    waived=client.put(path,headers=admin_headers,json={'status':'waived','reason':'客户明确无需此项','version':current['version']})
    assert waived.status_code==200 and waived.json()['status']=='waived'
    second=next(row for row in view(client,admin_headers,pid)['documents'] if row['id']==ids[1])
    review=client.post(f'/api/projects/{pid}/documents/{second["id"]}/versions/{second["versions"][0]["id"]}/review',headers=admin_headers,json={'decision':'rejected','note':'补充说明'})
    assert review.status_code==400 and '自己审核' in review.text
    created=client.post('/api/users',headers=admin_headers,json={'username':'reviewer','display_name':'另一审核人','password':'Review123!','role':'admin'})
    assert created.status_code==200,created.text
    token=client.post('/api/auth/login',json={'username':'reviewer','password':'Review123!'}).json()['access_token']
    reviewer={'Authorization':'Bearer '+token}
    review=client.post(f'/api/projects/{pid}/documents/{second["id"]}/versions/{second["versions"][0]["id"]}/review',headers=reviewer,json={'decision':'rejected','note':'补充说明'})
    assert review.status_code==200,review.text
    assert next(row for row in view(client,admin_headers,pid)['documents'] if row['id']==ids[1])['status']=='submitted'


def test_unclassified_files_entity_filter_and_cross_project_atomicity(client,admin_headers):
    p=create(client,admin_headers);q=create(client,admin_headers);pid=p['id'];qid=q['id']
    result=upload(client,admin_headers,pid)
    assert result.status_code==200,result.text
    assert len(view(client,admin_headers,pid)['unclassified_files'])==1
    assert client.get(f'/api/projects/{pid}/documents/files',headers=admin_headers).json()[0]['id']==result.json()['file']['id']
    other_doc=view(client,admin_headers,qid)['documents'][0]['id']
    assert upload(client,admin_headers,pid,[other_doc]).status_code==404
    assert len(client.get(f'/api/projects/{pid}/documents/files',headers=admin_headers).json())==1
    own_doc=view(client,admin_headers,pid)['documents'][0]['id']
    assert client.post(f'/api/projects/{qid}/documents/submit',headers=admin_headers,json={'document_ids':[other_doc],'attachment_id':result.json()['file']['id']}).status_code==400
    assert client.delete('/api/files/'+result.json()['file']['id'],headers=admin_headers).status_code==200


def test_vehicle_document_instances_and_inventory_generation(client,admin_headers):
    p=create(client,admin_headers);pid=p['id']
    v=client.post(f'/api/projects/{pid}/vehicles',headers=admin_headers,json={'vehicle_type':'潜伏式','model_name':'AGV','quantity':1}).json()['id']
    unit=client.post(f'/api/projects/{pid}/vehicle-units',headers=admin_headers,json={'vehicle_id':v,'asset_code':'AGV-01'})
    assert unit.status_code==200,unit.text
    rows=[row for row in view(client,admin_headers,pid)['documents'] if row['code']=='6.1']
    assert len(rows)==1 and rows[0]['vehicle_unit_id']==unit.json()['id']
    client.put(f'/api/projects/{pid}/events/internal_kickoff',headers=admin_headers,json=versioned(pid, {'actual_at':'2026-10-01T00:00:00Z'}))
    inventory=next(row for row in view(client,admin_headers,pid)['documents'] if row['code']=='9.2')
    generated=client.post(f'/api/projects/{pid}/documents/{inventory["id"]}/generate-inventory',headers=admin_headers,json={})
    assert generated.status_code==200,generated.text
    assert generated.json()['status']=='submitted'
    downloaded=client.get('/api'+generated.json()['latest_version']['download_url'],headers=admin_headers)
    assert downloaded.status_code==200 and downloaded.content.startswith(b'PK')


def test_project_snapshot_blank_ownership_clone_portable_history_and_excel(client,admin_headers):
    from openpyxl import load_workbook
    from app.db import SessionLocal
    from app.models import Project,User
    from app.services.import_export import project_package_zip
    tid=client.post('/api/templates',headers=admin_headers,json={'name':'项目独立模板'}).json()['id']
    requirement=next(row for row in client.get(f'/api/templates/{tid}/document-requirements',headers=admin_headers).json() if row['code']=='4.1')
    blank=b'original blank form\x00'
    assert client.post(f'/api/templates/{tid}/document-requirements/{requirement["id"]}/blank-file',headers=admin_headers,files={'file':('blank.bin',blank)}).status_code==200
    p=create(client,admin_headers,template_id=tid);pid=p['id']
    definition=next(row for row in view(client,admin_headers,pid)['definitions'] if row['code']=='4.1')
    document=next(row for row in view(client,admin_headers,pid)['documents'] if row['code']=='4.1')
    uploaded=upload(client,admin_headers,pid,[document['id']],data=b'original evidence')
    assert uploaded.status_code==200,uploaded.text
    assert client.delete(f'/api/templates/{tid}',headers=admin_headers).status_code==200
    path=f'/api/projects/{pid}/documents/requirements/{definition["id"]}/blank-file'
    assert client.get(path,headers=admin_headers).content==blank
    clone=client.post(f'/api/projects/{pid}/clone',headers=admin_headers,json={'name':'新项目副本'})
    assert clone.status_code==200,clone.text
    cloned=view(client,admin_headers,clone.json()['id'])
    assert all(row['status']=='pending' and not row['versions'] for row in cloned['documents'])
    clone_definition=next(row for row in cloned['definitions'] if row['code']=='4.1')
    assert client.get(f'/api/projects/{clone.json()["id"]}/documents/requirements/{clone_definition["id"]}/blank-file',headers=admin_headers).content==blank
    exported=client.get(f'/api/projects/{pid}/export/json',headers=admin_headers)
    assert exported.status_code==200,exported.text
    payload=exported.json()
    assert len(payload['document_versions'])==1 and payload['document_versions'][0]['file']['sha256']==hashlib.sha256(b'original evidence').hexdigest()
    preview=client.post('/api/projects/import/preview',headers=admin_headers,json=payload)
    assert preview.status_code==200 and preview.json()['valid'],preview.text
    assert preview.json()['summary']['document_definitions']==len(payload['document_definitions'])
    imported=client.post('/api/projects/import/json',headers=admin_headers,json=payload)
    assert imported.status_code==200,imported.text
    target=view(client,admin_headers,imported.json()['id'])
    assert all(row['status']=='pending' and not row['versions'] for row in target['documents'])
    imported_definition=next(row for row in target['definitions'] if row['code']=='4.1')
    facts=imported_definition['source_metadata']
    assert facts['kind']=='portable_unverified_snapshot' and facts['document_versions'][0]['file']['sha256']==hashlib.sha256(b'original evidence').hexdigest()
    assert client.get(f'/api/projects/{imported.json()["id"]}/documents/requirements/{imported_definition["id"]}/blank-file',headers=admin_headers).content==blank
    with SessionLocal() as db:
        archive=project_package_zip(db,pid,actor=db.query(User).filter_by(username='admin').one())
    with ZipFile(io.BytesIO(archive)) as package:
        assert json.loads(package.read('项目备份.json'))['document_versions']
        assert any(package.read(name)==b'original evidence' for name in package.namelist() if name.startswith('files/'))
        wb=load_workbook(io.BytesIO(package.read('项目数据.xlsx')))
        assert {'交付资料','物资清单','收发流水','生命周期表'} <= set(wb.sheetnames)
        assert any(cell.value=='evidence.txt' for row in wb['交付资料'].iter_rows(min_row=3) for cell in row)
        assert '交付归档' not in wb.sheetnames


@pytest.mark.parametrize('field,value',[('source_id',[]),('source_id',''),('blank_file_key',[]),('blank_file_key',['bad'])])
def test_malformed_document_import_is_atomic(client,admin_headers,field,value):
    from app.db import SessionLocal
    from app.models import Project,ChangeLog
    from app.config import UPLOAD_DIR
    p=create(client,admin_headers)
    payload=client.get(f'/api/projects/{p["id"]}/export/json',headers=admin_headers).json()
    payload['document_definitions'][0][field]=value
    files=set(UPLOAD_DIR.iterdir())
    with SessionLocal() as db:before=(db.query(Project).count(),db.query(ChangeLog).count())
    result=client.post('/api/projects/import/json',headers=admin_headers,json=payload)
    assert result.status_code==400,result.text
    with SessionLocal() as db:assert (db.query(Project).count(),db.query(ChangeLog).count())==before
    assert set(UPLOAD_DIR.iterdir())==files


def test_submit_audit_failure_rolls_back_file_versions_and_status(client,admin_headers,monkeypatch):
    from app.config import UPLOAD_DIR
    from app.db import SessionLocal
    from app.models import Attachment,DocumentVersion,ChangeLog
    import app.services.project_documents as service
    pid=create(client,admin_headers)['id'];document=view(client,admin_headers,pid)['documents'][0]
    files=set(UPLOAD_DIR.iterdir())
    with SessionLocal() as db:before=(db.query(Attachment).count(),db.query(DocumentVersion).count(),db.query(ChangeLog).count())
    original=service.log_change
    def fail(db,pid,entity,*args,**kwargs):
        if entity=='project_document':raise RuntimeError('document audit unavailable')
        return original(db,pid,entity,*args,**kwargs)
    monkeypatch.setattr(service,'log_change',fail)
    with pytest.raises(RuntimeError,match='document audit unavailable'):upload(client,admin_headers,pid,[document['id']])
    assert set(UPLOAD_DIR.iterdir())==files
    # Check rollback before a new authorized metadata view adds its own T26 audit.
    with SessionLocal() as db:assert (db.query(Attachment).count(),db.query(DocumentVersion).count(),db.query(ChangeLog).count())==before
    assert next(row for row in view(client,admin_headers,pid)['documents'] if row['id']==document['id'])['status']=='pending'


def test_procurement_extra_patch_preserves_unknown_metadata_and_invalidates_changed_review(client,admin_headers):
    pid=create(client,admin_headers)['id'];base=f'/api/projects/{pid}/procurements'
    created=client.post(base,headers=admin_headers,json={'name':'充电桩','extra':{'quantity':1,'spec':'原规格','unit':'套','historical_reference':{'keep':'原订单'}}})
    assert created.status_code==200,created.text
    item=created.json();path=base+'/'+item['id']
    assert client.put(path,headers=admin_headers,json={'status':'passed','actual_date':'2026-01-01'}).status_code==200
    updated=client.put(path,headers=admin_headers,json={'extra':{'quantity':2}})
    assert updated.status_code==200,updated.text
    assert updated.json()['extra']=={'quantity':2,'spec':'原规格','unit':'套','historical_reference':{'keep':'原订单'}}
    assert updated.json()['status']=='pending' and updated.json()['completed_at'] is None


def test_readonly_project_member_can_read_documents_without_writing(client,admin_headers):
    from app.db import SessionLocal
    from app.models import User,Project
    pid=create(client,admin_headers)['id'];document=view(client,admin_headers,pid)['documents'][0]
    assert upload(client,admin_headers,pid,[document['id']]).status_code==200
    result=client.post('/api/users',headers=admin_headers,json={'username':'reader','display_name':'只读成员','password':'Review123!','role':'readonly'})
    assert result.status_code==200,result.text
    with SessionLocal() as db:
        reader=db.query(User).filter_by(username='reader').one();p=db.get(Project,pid);p.members=[*(p.members or []),reader.id];db.commit()
    token=client.post('/api/auth/login',json={'username':'reader','password':'Review123!'}).json()['access_token'];headers={'Authorization':'Bearer '+token}
    state=view(client,headers,pid);assert not state['permissions']['submit'] and not state['permissions']['manage']
    submitted=next(row for row in state['documents'] if row['id']==document['id'])
    assert client.get('/api'+submitted['latest_version']['download_url'],headers=headers).content==b'evidence'
    assert upload(client,headers,pid,[document['id']]).status_code==403
    assert client.put(f'/api/projects/{pid}/documents/{document["id"]}',headers=headers,json=versioned(document["id"], {'status':'waived','reason':'只读不能写'}, 'ProjectDocument')).status_code==403


def rule_body(row,**changes):
    from app.services.document_requirements import RequirementInput
    fields=RequirementInput.model_fields.keys()
    return {key:value for key,value in row.items() if key in fields}|changes


def test_edit_unsubmitted_rule_updates_pending_but_preserves_submitted_snapshot(client,admin_headers):
    pid=create(client,admin_headers)['id'];state=view(client,admin_headers,pid)
    pending=next(row for row in state['documents'] if row['code']=='4.1')
    definition=next(row for row in state['definitions'] if row['id']==pending['definition_id'])
    changed=client.put(f'/api/projects/{pid}/documents/requirements/{definition["id"]}',headers=admin_headers,
        json=rule_body(definition,name='本项目新的资料名',required=False))
    assert changed.status_code==200,changed.text
    updated=next(row for row in view(client,admin_headers,pid)['documents'] if row['id']==pending['id'])
    assert updated['name']=='本项目新的资料名' and not updated['required'] and updated['version']>pending['version']
    assert upload(client,admin_headers,pid,[pending['id']]).status_code==200
    assert client.put(f'/api/projects/{pid}/documents/requirements/{definition["id"]}',headers=admin_headers,
        json=rule_body(changed.json(),name='仅后续新资料采用',required=True)).status_code==200
    kept=next(row for row in view(client,admin_headers,pid)['documents'] if row['id']==pending['id'])
    assert kept['name']=='本项目新的资料名' and not kept['required'] and kept['versions']
    assert client.put(f'/api/projects/{pid}/documents/requirements/{definition["id"]}',headers=admin_headers,
        json=rule_body(changed.json(),source_metadata={'fake':True})).status_code==422


def legacy_fixture(pid):
    from app.db import SessionLocal
    from app.models import ProjectDocumentDefinition,ProjectDocument
    with SessionLocal() as db:
        definition=db.query(ProjectDocumentDefinition).filter_by(project_id=pid,code='4.1').one()
        document=db.query(ProjectDocument).filter_by(definition_id=definition.id).one()
        source={'kind':'task_deliverable','task':{'id':'original-task','version':3},'original':{'name':'原交付物','required':True,'done':False}}
        for row in (definition,document):
            row.stage_key='legacy';row.code='legacy.'+'a'*32;row.name='原交付物';row.timing_type='on_demand';row.timing_event=None;row.timing_offset_hours=None;row.source_metadata=source
        document.version=1  # Frozen pre-T22 migration records begin at version 1.
        db.commit();return document.id,definition.id,source


def test_explicit_legacy_classification_keeps_identity_versions_and_source(client,admin_headers):
    pid=create(client,admin_headers)['id'];did,rid,source=legacy_fixture(pid)
    assert upload(client,admin_headers,pid,[did]).status_code==200
    current=next(row for row in view(client,admin_headers,pid)['definitions'] if row['id']==rid)
    body=rule_body(current,stage_key='6',code='6.99',name='明确分类的原交付物')
    assert client.put(f'/api/projects/{pid}/documents/requirements/{rid}',headers=admin_headers,json=body).status_code==200
    state=view(client,admin_headers,pid);classified=[row for row in state['documents'] if row['definition_id']==rid]
    assert len(classified)==1 and classified[0]['id']==did and classified[0]['stage_key']=='6'
    assert classified[0]['status']=='submitted' and classified[0]['versions'] and classified[0]['source_metadata']==source
    assert client.post(f'/api/projects/{pid}/documents/requirements',headers=admin_headers,
        json=rule_body(current)).status_code==400


def test_legacy_offline_queue_translation_is_idempotent_and_keeps_rejected_drafts(client,admin_headers):
    from app.db import SessionLocal
    from app.models import DocumentVersion,ChangeLog,OfflineOpReceipt
    pid=create(client,admin_headers)['id'];did,_,_=legacy_fixture(pid)
    asset=upload(client,admin_headers,pid).json()['file']
    payload={'project_id':pid,'task_id':'original-task','version':3,'deliverables':[{'name':'原交付物','done':True,'attachment_id':asset['id']}]}
    operation={'type':'deliverable','client_uuid':'preserved-original-uuid','payload':payload}
    result=client.post('/api/worker/sync',headers=admin_headers,json={'ops':[operation]})
    assert result.status_code==200 and result.json()['accepted']==['preserved-original-uuid'],result.text
    assert client.post('/api/worker/sync',headers=admin_headers,json={'ops':[operation]}).json()['accepted']==['preserved-original-uuid']
    unknown=copy.deepcopy(operation);unknown['client_uuid']='unknown-original-uuid';unknown['payload']['deliverables'][0]['name']='不能猜对应'
    refused=client.post('/api/worker/sync',headers=admin_headers,json={'ops':[unknown]})
    assert refused.json()['rejected']==['unknown-original-uuid'] and '保留' in refused.json()['rejected_detail']['unknown-original-uuid']
    with SessionLocal() as db:
        assert db.query(DocumentVersion).filter_by(project_document_id=did).count()==1
        log=db.query(ChangeLog).filter_by(entity_type='legacy_offline_document').one()
        assert log.new_value['client_uuid']=='preserved-original-uuid' and log.new_value['original_payload']==payload
        assert db.query(OfflineOpReceipt).filter_by(client_uuid='unknown-original-uuid').count()==0


def test_retention_delete_restore_and_legacy_no_auto_purge_preserve_new_document_evidence(client,admin_headers):
    from datetime import timedelta
    from app.db import SessionLocal
    from app.models import Project,ProjectArchive,DocumentVersion,Attachment,ProjectDocument
    from app.services.project_retention import project_retention_snapshot
    from app.services.report_scheduler import purge_soft_deleted_projects
    from app.time_utils import utcnow
    pid=create(client,admin_headers)['id'];document=view(client,admin_headers,pid)['documents'][0]
    result=upload(client,admin_headers,pid,[document['id']]);assert result.status_code==200
    version_id=result.json()['documents'][0]['latest_version']['id']
    with SessionLocal() as db:
        snapshot=project_retention_snapshot(db,db.get(Project,pid))
        assert snapshot['records']['document_versions'][0]['id']==version_id
        assert snapshot['binary_files']['manifest'][0]['sha256']==hashlib.sha256(b'evidence').hexdigest()
    current=client.get('/api/projects/'+pid,headers=admin_headers).json()
    deleted=client.delete('/api/projects/'+pid,headers=admin_headers,params={'version':current['version'],'reason':'验证新资料保全'})
    assert deleted.status_code==200,deleted.text
    row=next(row for row in client.get('/api/project-recycle-bin/projects',headers=admin_headers).json() if row['id']==pid)
    restored=client.post(f'/api/project-recycle-bin/projects/{pid}/restore',headers=admin_headers,json={'version':row['version'],'deletion_id':row['deletion_id']})
    assert restored.status_code==200,restored.text
    assert next(row for row in view(client,admin_headers,pid)['documents'] if row['id']==document['id'])['latest_version']['id']==version_id
    # #85: the retired automatic purge entrypoint cannot delete old recycled evidence.
    other=create(client,admin_headers)['id'];other_doc=view(client,admin_headers,other)['documents'][0]
    assert upload(client,admin_headers,other,[other_doc['id']]).status_code==200
    with SessionLocal() as db:
        p=db.get(Project,other);p.deleted_at=utcnow()-timedelta(days=60);db.commit()
        result=purge_soft_deleted_projects(db)
        assert result['purged']==0 and result['disabled'];db.commit()
        assert db.get(Project,other).deleted_at is not None
        assert not db.query(ProjectArchive).filter_by(project_id=other).first()
        assert db.query(DocumentVersion).filter_by(project_document_id=other_doc['id']).count()==1
        assert db.query(Attachment).filter_by(project_id=other).count()==1
        retained=project_retention_snapshot(db,db.get(Project,other))
        assert len(retained['records']['document_versions'])==1
        assert retained['binary_files']['manifest'][0]['available_at_purge']


def test_completed_document_gate_excludes_future_stages_and_requires_real_evidence(client,admin_headers):
    from app.db import SessionLocal
    from app.models import ProjectDocument
    from app.services.lifecycle_gates import _pending_required_documents
    pid=create(client,admin_headers)['id']
    with SessionLocal() as db:
        rows=db.query(ProjectDocument).filter_by(project_id=pid).all()
        for row in rows:
            row.status='pending' if row.stage_key in {'8','9'} else 'waived'
            row.waived_reason='本项目明确不适用' if row.status=='waived' else ''
        db.commit()
        assert _pending_required_documents(db,pid)==[]
        assert _pending_required_documents(db,pid,archiving=True)
        missing=next(row for row in rows if row.stage_key=='4')
        missing.status='submitted';missing.waived_reason='';db.commit()
        assert any(missing.code in text for text in _pending_required_documents(db,pid))
        missing.stage_key='legacy';db.commit()
        assert _pending_required_documents(db,pid)==[]


def test_agent_document_read_has_scope_and_project_isolation(client,admin_headers):
    from .test_agent_api import token,call
    pid=create(client,admin_headers)['id'];document=view(client,admin_headers,pid)['documents'][0]
    assert upload(client,admin_headers,pid,[document['id']]).status_code==200
    _,headers=token(client,admin_headers,scopes=['projects:read','files:read'])
    response=call(client,headers,'documents_read',project_id=pid)
    assert response.status_code==200 and response.json()['ok'],response.text
    rows=response.json()['data']['documents']
    assert next(row for row in rows if row['id']==document['id'])['latest_version']['sha256']==hashlib.sha256(b'evidence').hexdigest()
    _,limited=token(client,admin_headers,scopes=['projects:read'])
    denied=call(client,limited,'documents_read',project_id=pid)
    assert denied.status_code==403 or denied.json().get('ok') is False
    _,other=token(client,admin_headers,user='user-worker-001',scopes=['projects:read','files:read'])
    assert call(client,other,'documents_read',project_id=pid).json()['ok'] is False


def test_acceptance_scope_survives_structural_migration_but_detects_real_changes(client,admin_headers):
    from .test_startup_targets import setup_project
    from app.db import SessionLocal
    from app.models import Task,ProjectDocument
    from app.services.acceptance import acceptance_scope_fingerprint
    pid,tasks=setup_project(client,admin_headers)
    with SessionLocal() as db:
        task=db.get(Task,tasks[0]['id']);original=[{'name':'原验收交付要求','required':True,'doc_type':'PDF','note':'保留定义'}]
        task.deliverables=original;db.flush()
        before=acceptance_scope_fingerprint(db,pid)
        task.deliverables=[]
        moved=ProjectDocument(project_id=pid,instance_key='task:'+task.id+':0',stage_key='legacy',code='legacy.'+'b'*32,
            name='原验收交付要求',scope='project',timing_type='on_demand',required=True,
            source_metadata={'kind':'task_deliverable','task':{'id':task.id},'index':0,'original':original[0]})
        db.add(moved);db.flush()
        assert acceptance_scope_fingerprint(db,pid)==before
        task.name+='真实范围修改';db.flush()
        assert acceptance_scope_fingerprint(db,pid)!=before


def test_old_template_import_keeps_original_sources_as_legacy_without_old_writes(client,admin_headers):
    from app.db import SessionLocal
    from app.models import Task,TemplatePreparationItem
    tid=client.post('/api/templates',headers=admin_headers,json={'name':'历史模板导入'}).json()['id']
    payload={'phases':[{'name':'原第1阶段','modules':[{'name':'原模块','tasks':[{'name':'原任务','quantity_type':'manual',
        'deliverables':[{'name':'原自定义交付要求','required':False,'original_extra':{'keep':'原字段'}}]}]}]}]}
    assert client.put(f'/api/templates/{tid}/tree',headers=admin_headers,json=payload).status_code==400
    imported=client.post(f'/api/templates/{tid}/import/json',headers=admin_headers,json=payload)
    assert imported.status_code==200,imported.text
    rows=client.get(f'/api/templates/{tid}/document-requirements',headers=admin_headers).json()
    row=next(row for row in rows if row['name']=='原自定义交付要求')
    assert row['stage_key']=='legacy' and not row['required']
    facts=row['source_metadata']['original']
    assert facts['original']['original_extra']=={'keep':'原字段'} and facts['phase']['name']=='原第1阶段'
    with SessionLocal() as db:
        assert not any(task.deliverables for task in db.query(Task).filter_by(template_id=tid))
        assert not db.query(TemplatePreparationItem).filter_by(template_id=tid).first()


def test_concurrent_real_submissions_keep_every_accepted_version_and_file(client,admin_headers):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier
    from app.db import SessionLocal,engine
    from app.models import Attachment,DocumentVersion,ChangeLog
    pid=create(client,admin_headers)['id'];document=view(client,admin_headers,pid)['documents'][0]
    barrier=Barrier(2)
    def submit(index):
        barrier.wait(timeout=10)
        return upload(client,admin_headers,pid,[document['id']],data=f'original-{index}'.encode())
    with ThreadPoolExecutor(max_workers=2) as pool:results=list(pool.map(submit,range(2)))
    assert all(result.status_code in {200,409} for result in results),[(r.status_code,r.text) for r in results]
    accepted=sum(result.status_code==200 for result in results)
    assert accepted>=1
    if engine.dialect.name=='postgresql':assert accepted==2
    with SessionLocal() as db:
        assert db.query(DocumentVersion).filter_by(project_document_id=document['id']).count()==accepted
        assert db.query(Attachment).filter_by(project_id=pid).count()==accepted
        assert db.query(ChangeLog).filter_by(project_id=pid,entity_type='project_document',action='submit').count()==accepted
    current=next(row for row in view(client,admin_headers,pid)['documents'] if row['id']==document['id'])
    assert current['version']==document['version']+accepted
    assert {client.get('/api'+row['download_url'],headers=admin_headers).content for row in current['versions']} <= {b'original-0',b'original-1'}
