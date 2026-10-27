"""原小数填报是完整只读来源，不进入本周完成或过期风险。"""
import io
from datetime import date
from decimal import Decimal
from openpyxl import load_workbook
from app.db import SessionLocal
from app.models import Project,DailyProgress,Task,FieldOperation
from tests.test_project_lifecycle_boundaries import project,task_entry


def test_legacy_decimal_reports_preserved_without_becoming_current_completion(client,pm_headers):
    pid=project(client,pm_headers);tid,eid,_=task_entry(pid)
    with SessionLocal() as db:
        actor=db.get(Project,pid).manager
        for index in range(3):db.add(DailyProgress(project_id=pid,task_id=tid,schedule_entry_id=eid,user_id=actor,
            work_date=date(2026,9,1),completed_quantity=.1,client_uuid=f'original-decimal-{index}'))
        db.commit()
    response=client.get(f'/api/projects/{pid}/export/excel',headers=pm_headers);assert response.status_code==200,response.text
    workbook=load_workbook(io.BytesIO(response.content))
    assert '交付概览' in workbook.sheetnames and '日程排程' not in workbook.sheetnames
    assert not any(cell.value==.30000000000000004 for sheet in workbook for row in sheet for cell in row)
    snapshot=client.get(f'/api/projects/{pid}/export/json',headers=pm_headers).json()
    assert len(snapshot['progress'])==3
    assert sum(Decimal(str(row['completed_quantity'])) for row in snapshot['progress'])==Decimal('.3')
    report=client.get(f'/api/projects/{pid}/report',headers=pm_headers,params={'ref_date':'2026-09-02'}).json()
    assert report['summary']['actual_records']==report['summary']['completed_work']==0 and not report['completed_items']
    assert not report['risks']  # Old plan quantity variance is not a current document deadline.
    with SessionLocal() as db:
        assert not db.get(Task,tid).manually_completed and db.query(FieldOperation).filter_by(project_id=pid).count()==0
        assert db.query(DailyProgress).filter_by(project_id=pid).count()==3
