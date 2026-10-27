"""Report forms: 生命周期表 / 质量问题反馈单 / 内部联系单（物品申请）端到端测试."""

from tests.lifecycle_helpers import versioned
import io
import zipfile
from datetime import datetime

from openpyxl import Workbook, load_workbook


def _project(client, headers, name="表单测试项目"):
    r = client.post("/api/projects", headers=headers, json={
        "name": name, "start_date": "2026-08-01", "end_date": "2026-12-31"})
    assert r.status_code == 200, r.text
    return r.json()["id"]


# ===== 生命周期表 =====

def test_lifecycle_export(client, pm_headers):
    pid = _project(client, pm_headers)
    unit = client.post(f"/api/projects/{pid}/vehicle-units", headers=pm_headers, json={
        "asset_code": "AGV-001", "fleet_no": "7", "current_stage": "planned"}).json()
    client.post(f"/api/projects/{pid}/vehicle-units/{unit['id']}/events", headers=pm_headers, json={
        "event_type": "fault", "event_date": "2026-08-12", "title": "驱动轮异响",
        "description": "行驶中左轮异响", "action_taken": "更换轴承", "handler": "施泽辉",
        "replaced_parts": [], "is_closed": False, "status_after": "fault"})

    r = client.get(f"/api/projects/{pid}/export/lifecycle", headers=pm_headers)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    headers = [ws.cell(row=2, column=c).value for c in range(1, 12)]
    assert headers[1] == "项目名称" and headers[5] == "出现过的问题"
    row = [ws.cell(row=3, column=c).value for c in range(1, 12)]
    assert row[1] == "表单测试项目"          # 项目名称
    assert row[3] == "AGV-001"              # 车辆编号
    assert row[4] == "7"                    # 车号
    assert row[5] == "驱动轮异响"            # 出现过的问题
    assert row[6] == "行驶中左轮异响"         # 问题描述
    assert row[7] == "更换轴承"              # 解决方法
    assert row[8] == "施泽辉"               # 记录人
    assert row[9] == "2026-08-12"           # 记录时间
    assert row[10] in (None, "")            # 问题分类（未填）


def _lifecycle_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["序号", "项目名称", "车辆型号", "车辆编号", "车号", "出现过的问题", "问题描述", "解决方法", "记录人", "记录时间", "问题分类"])
    ws.append(["1", "旧项目", "潜伏式", "77BC00180", "2",
               "光通信接线接错\n安全激光损坏", "现象A\n现象B", "更换光通信\n更换激光",
               "施泽辉", "2023/5/4\n2023-05-21 00:00:00", "硬件问题\n硬件问题"])
    ws.append(["2", "旧项目", "AGVLSYP", "77BC00192", "3", "安全激光损坏",
               "行驶中激光损坏", "更换激光", "裘辉", "2023-06-01", "硬件"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def test_lifecycle_import(client, pm_headers):
    pid = _project(client, pm_headers)
    r = client.post(f"/api/projects/{pid}/import/lifecycle", headers=pm_headers,
                    files={"file": ("生命周期表.xlsx", _lifecycle_xlsx(),
                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    assert r.json()["events"] == 3  # 第一行两行拆分 + 第二行一条

    units = client.get(f"/api/projects/{pid}/vehicle-units", headers=pm_headers).json()
    assert len(units) == 2
    unit1 = next(u for u in units if u["asset_code"] == "77BC00180")
    assert unit1["fleet_no"] == "2"
    detail = client.get(f"/api/projects/{pid}/vehicle-units/{unit1['id']}", headers=pm_headers).json()
    assert len(detail["events"]) == 2
    # 按 event_date 倒序：2023-05-21（现象B）在前，2023-05-04（现象A）在后
    assert detail["events"][0]["extra"]["category"] == "硬件问题"
    assert detail["events"][0]["description"] == "现象B"
    assert detail["events"][1]["description"] == "现象A"


def test_lifecycle_import_idempotent_matching_unit(client, pm_headers):
    pid = _project(client, pm_headers)
    client.post(f"/api/projects/{pid}/vehicle-units", headers=pm_headers, json={
        "asset_code": "77BC00192", "fleet_no": "3", "current_stage": "planned"})
    r = client.post(f"/api/projects/{pid}/import/lifecycle", headers=pm_headers,
                    files={"file": ("生命周期表.xlsx", _lifecycle_xlsx(),
                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    assert r.json()["units"] == 1  # 只新建 77BC00180，77BC00192 复用现有档案
    units = client.get(f"/api/projects/{pid}/vehicle-units", headers=pm_headers).json()
    assert len(units) == 2


def test_lifecycle_export_model_name(client, pm_headers):
    """车辆型号列导出 model_name（AGVLSYP），而非 vehicle_type（潜伏式）。"""
    pid = _project(client, pm_headers)
    vehicle = client.post(f"/api/projects/{pid}/vehicles", headers=pm_headers, json={
        "vehicle_type": "潜伏式", "model_name": "AGVLSYP", "quantity": 1}).json()
    unit = client.post(f"/api/projects/{pid}/vehicle-units", headers=pm_headers, json={
        "asset_code": "AGV-001", "vehicle_id": vehicle["id"], "fleet_no": "7",
        "current_stage": "planned"}).json()
    client.post(f"/api/projects/{pid}/vehicle-units/{unit['id']}/events", headers=pm_headers, json={
        "event_type": "fault", "event_date": "2026-08-12", "title": "驱动轮异响",
        "description": "行驶中左轮异响", "action_taken": "更换轴承", "handler": "施泽辉",
        "replaced_parts": [], "is_closed": False, "status_after": "fault"})

    r = client.get(f"/api/projects/{pid}/export/lifecycle", headers=pm_headers)
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    assert ws.cell(row=2, column=3).value == "车辆型号"
    assert ws.cell(row=3, column=3).value == "AGVLSYP"  # model_name


def test_lifecycle_import_links_vehicle_by_model_name(client, pm_headers):
    """导入按「车辆型号」列匹配 model_name 关联车型档案。"""
    pid = _project(client, pm_headers)
    client.post(f"/api/projects/{pid}/vehicles", headers=pm_headers, json={
        "vehicle_type": "潜伏式", "model_name": "AGVLSYP", "quantity": 1}).json()
    r = client.post(f"/api/projects/{pid}/import/lifecycle", headers=pm_headers,
                    files={"file": ("生命周期表.xlsx", _lifecycle_xlsx(),
                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    units = client.get(f"/api/projects/{pid}/vehicle-units", headers=pm_headers).json()
    unit2 = next(u for u in units if u["asset_code"] == "77BC00192")
    detail = client.get(f"/api/projects/{pid}/vehicle-units/{unit2['id']}", headers=pm_headers).json()
    assert detail["model_name"] == "AGVLSYP"      # 车辆型号列按 model_name 关联到车型
    assert detail["vehicle_id"] is not None
    unit1 = next(u for u in units if u["asset_code"] == "77BC00180")
    detail1 = client.get(f"/api/projects/{pid}/vehicle-units/{unit1['id']}", headers=pm_headers).json()
    assert detail1["vehicle_id"] is None          # 车辆型号「潜伏式」无对应 model_name，不建档


def test_lifecycle_import_idempotent_events(client, pm_headers):
    """重复导入同一文件整体幂等：事件不翻倍，并报告跳过数。"""
    pid = _project(client, pm_headers)
    files = {"file": ("生命周期表.xlsx", _lifecycle_xlsx(),
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r1 = client.post(f"/api/projects/{pid}/import/lifecycle", headers=pm_headers, files=files)
    assert r1.status_code == 200, r1.text
    assert r1.json()["events"] == 3 and r1.json()["skipped"] == 0
    r2 = client.post(f"/api/projects/{pid}/import/lifecycle", headers=pm_headers, files=files)
    assert r2.status_code == 200, r2.text
    assert r2.json()["events"] == 0 and r2.json()["skipped"] == 3
    units = client.get(f"/api/projects/{pid}/vehicle-units", headers=pm_headers).json()
    assert len(units) == 2  # 车辆档案也不重复
    detail = client.get(f"/api/projects/{pid}/vehicle-units/{units[0]['id']}", headers=pm_headers).json()
    assert len(detail["events"]) == 2  # 事件总数不变


def test_lifecycle_browser_record_export_reimport_preserves_resolution(client, pm_headers):
    pid = _project(client, pm_headers)
    unit = client.post(f"/api/projects/{pid}/vehicle-units", headers=pm_headers,
                       json={"asset_code": "QA-001"}).json()
    created = client.post(f"/api/projects/{pid}/vehicle-units/{unit['id']}/events", headers=pm_headers,
                          json={"event_type": "fault", "event_date": "2026-10-05", "title": "传感器报警",
                                "action_taken": "紧固连接器", "verification": "连续试运行正常",
                                "fault_part": "传感器", "is_closed": True, "status_after": "normal"})
    assert created.status_code == 200, created.text
    export = client.get(f"/api/projects/{pid}/export/lifecycle", headers=pm_headers)
    for _ in range(2):
        result = client.post(f"/api/projects/{pid}/import/lifecycle", headers=pm_headers,
                             files={"file": ("生命周期表.xlsx", export.content,
                                             "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert result.status_code == 200, result.text
        assert result.json()["events"] == 0 and result.json()["skipped"] == 1
    detail = client.get(f"/api/projects/{pid}/vehicle-units/{unit['id']}", headers=pm_headers).json()
    assert len(detail["events"]) == 1 and detail["open_issue_count"] == 0
    assert detail["events"][0]["is_closed"] is True
    assert detail["events"][0]["verification"] == "连续试运行正常"
    assert detail["events"][0]["fault_part"] == "传感器"


def test_lifecycle_import_bad_date_400(client, pm_headers):
    """无法解析的记录时间明确报错，不再静默用当天日期。"""
    pid = _project(client, pm_headers)
    wb = Workbook()
    ws = wb.active
    ws.append(["序号", "项目名称", "车辆型号", "车辆编号", "车号", "出现过的问题", "问题描述", "解决方法", "记录人", "记录时间", "问题分类"])
    ws.append(["1", "旧项目", "AGVLSYP", "77BC00180", "2", "光通信接线接错", "现象A", "更换光通信", "施泽辉", "？？？", "硬件"])
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post(f"/api/projects/{pid}/import/lifecycle", headers=pm_headers,
                    files={"file": ("生命周期表.xlsx", buf.getvalue(),
                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 400, r.text
    assert "记录时间" in r.json()["detail"]
    # 记录时间留空同样报错
    wb2 = Workbook()
    ws2 = wb2.active
    ws2.append(["序号", "项目名称", "车辆型号", "车辆编号", "车号", "出现过的问题", "问题描述", "解决方法", "记录人", "记录时间", "问题分类"])
    ws2.append(["1", "旧项目", "AGVLSYP", "77BC00180", "2", "光通信接线接错", "现象A", "更换光通信", "施泽辉", "", "硬件"])
    buf2 = io.BytesIO()
    wb2.save(buf2)
    r2 = client.post(f"/api/projects/{pid}/import/lifecycle", headers=pm_headers,
                     files={"file": ("生命周期表.xlsx", buf2.getvalue(),
                                     "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r2.status_code == 400, r2.text
    assert "记录时间" in r2.json()["detail"]


def test_lifecycle_import_unmatched_model_name_displayed(client, pm_headers):
    """未匹配车型档案时，导入的车辆型号仍持久化并在列表/详情展示。"""
    pid = _project(client, pm_headers)
    r = client.post(f"/api/projects/{pid}/import/lifecycle", headers=pm_headers,
                    files={"file": ("生命周期表.xlsx", _lifecycle_xlsx(),
                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    units = client.get(f"/api/projects/{pid}/vehicle-units", headers=pm_headers).json()
    unit1 = next(u for u in units if u["asset_code"] == "77BC00180")   # 车型「潜伏式」无对应档案
    assert unit1["vehicle_id"] is None
    assert unit1["model_name"] == "潜伏式"        # 列表展示导入的车辆型号
    detail = client.get(f"/api/projects/{pid}/vehicle-units/{unit1['id']}", headers=pm_headers).json()
    assert detail["model_name"] == "潜伏式"        # 单车详情展示导入的车辆型号
    unit2 = next(u for u in units if u["asset_code"] == "77BC00192")
    assert unit2["model_name"] == "AGVLSYP"        # 另一台未匹配档案同样持久化展示


def test_import_key_index_migration(client, upgrade_legacy_schema):
    """存量库缺 import_key 索引时，Alembic 版本迁移补齐（SQLite）。"""
    from sqlalchemy import text
    from app.db import SessionLocal, engine
    from sqlalchemy import inspect
    db = SessionLocal()
    try:
        db.execute(text("DROP INDEX IF EXISTS ix_vehicle_lifecycle_events_import_key"))
        db.commit()
        upgrade_legacy_schema()
        from app.db import _alembic_sync_schema
        _alembic_sync_schema()
        # A pooled sqlite3 handle can retain its old schema cache after another
        # handle performs DDL. Read the table catalog, then reflect the index
        # on that same physical connection (also valid for PostgreSQL).
        with engine.connect() as connection:
            reflection = inspect(connection)
            assert 'vehicle_lifecycle_events' in reflection.get_table_names()
            assert 'ix_vehicle_lifecycle_events_import_key' in {
                index['name'] for index in reflection.get_indexes('vehicle_lifecycle_events')
            }
    finally:
        db.close()


# ===== 质量问题反馈单 =====

def _quality_form_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.merge_cells("A1:H3")
    ws["A1"] = "产品质量问题反馈单据"
    ws["A4"] = "问题类型 Type"; ws["C4"] = "项目 Project"; ws["G4"] = "反馈人 Name"
    ws["A5"] = "硬件"; ws["C5"] = "天津环晟"; ws["G5"] = "孙晓杰"
    ws["A8"] = "车型 Model"; ws["D8"] = "车辆编号 No."; ws["G8"] = "问题来源 From"
    ws["A9"] = "AGVLSYP"; ws["D9"] = "77BC00255"; ws["G9"] = "项目现场"
    ws["A12"] = "问题部件 Parts"; ws["G12"] = "提交日期 Date"
    ws["A13"] = "伺服电机编码器"; ws["G13"] = "2024-05-08"
    ws["A16"] = "问题简述（或现象）Description"
    ws["A17"] = "行驶时一侧突然断连，另一侧轮子锁死，按急停无效"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def test_quality_form_import(client, pm_headers):
    pid = _project(client, pm_headers)
    r = client.post(f"/api/projects/{pid}/import/quality-form", headers=pm_headers,
                    files={"file": ("质量问题反馈单.xlsx", _quality_form_xlsx(),
                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    issue_id = r.json()["id"]
    issue = client.get(f"/api/projects/{pid}/issues", headers=pm_headers).json()
    assert len(issue) == 1
    assert issue[0]["issue_type"] == "hardware"
    assert issue[0]["description"].startswith("行驶时一侧突然断连")
    assert issue[0]["extra"]["vehicle_model"] == "AGVLSYP"
    assert issue[0]["extra"]["vehicle_no"] == "77BC00255"
    assert issue[0]["extra"]["part"] == "伺服电机编码器"
    assert issue[0]["extra"]["reporter_name"] == "孙晓杰"
    assert issue[0]["extra"]["feedback_date"] == "2024-05-08"


def test_quality_form_export(client, pm_headers):
    pid = _project(client, pm_headers)
    issue = client.post(f"/api/projects/{pid}/issues", headers=pm_headers, json={
        "title": "编码器断连", "description": "行驶时一侧突然断连，另一侧轮子锁死",
        "issue_type": "hardware", "severity": "high",
        "extra": {"vehicle_model": "AGVLSYP", "vehicle_no": "77BC00255", "source": "项目现场",
                  "part": "伺服电机编码器", "feedback_date": "2024-05-08", "reporter_name": "孙晓杰"},
    }).json()
    r = client.get(f"/api/projects/{pid}/issues/{issue['id']}/export/quality-form", headers=pm_headers)
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    assert ws.title == "77BC00255"
    assert ws["A1"].value == "产品质量问题反馈单据"
    assert ws["A5"].value == "硬件"
    assert ws["D9"].value == "77BC00255"
    assert ws["A13"].value == "伺服电机编码器"
    assert ws["G13"].value == datetime(2024, 5, 8)


def test_quality_form_export_matches_original_template(client, pm_headers):
    """导出版式逐项对齐原始附件 质量问题反馈单.xlsx：
    表头字体/下边框、合并单元格、字段位置、列宽、字体、边框、打印页边距。"""
    pid = _project(client, pm_headers)
    issue = client.post(f"/api/projects/{pid}/issues", headers=pm_headers, json={
        "title": "编码器断连", "description": "行驶时一侧突然断连，另一侧轮子锁死",
        "issue_type": "hardware", "severity": "high",
        "extra": {"vehicle_model": "AGVLSYP", "vehicle_no": "77BC00255", "source": "项目现场",
                  "part": "伺服电机编码器", "feedback_date": "2024-05-08", "reporter_name": "孙晓杰"},
    }).json()
    r = client.get(f"/api/projects/{pid}/issues/{issue['id']}/export/quality-form", headers=pm_headers)
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active

    # 合并单元格与原版一一对应
    expected_merges = {
        "A1:H3", "A4:B4", "C4:F4", "G4:H4", "A5:B7", "C5:F7", "G5:H7",
        "A8:C8", "D8:F8", "G8:H8", "A9:C11", "D9:F11", "G9:H11",
        "A12:D12", "E12:F15", "G12:H12", "A13:D15", "G13:H15",
        "A16:H16", "A17:H20", "A21:H21",
        "A22:D22", "E22:H22", "A23:D38", "E23:H38",
        "A39:D39", "E39:H39", "A40:D55", "E40:H55",
    }
    actual_merges = {str(rng).replace(" ", "") for rng in ws.merged_cells.ranges}
    assert actual_merges == expected_merges, f"合并单元格不一致: {actual_merges - expected_merges}"

    # 表头：黑体 20，仅下边框
    assert ws["A1"].font.name == "黑体" and ws["A1"].font.size == 20
    assert ws["A1"].border.top is None and ws["A1"].border.bottom.style == "thin"
    # 标题对齐：水平居中、垂直未设置（原版默认底部）
    assert ws["A1"].alignment.horizontal == "center"
    assert ws["A1"].alignment.vertical is None

    # 标签宋体 12 左对齐、值宋体 11 居中、图片区块标题楷体 12
    assert ws["A4"].font.name == "宋体" and ws["A4"].font.size == 12
    assert ws["A4"].alignment.horizontal == "left"
    assert ws["A5"].font.name == "宋体" and ws["A5"].font.size == 11
    assert ws["A5"].alignment.horizontal == "center"
    assert ws["A22"].font.name == "楷体" and ws["A22"].font.size == 12

    # 内容区全部细边框（标签、值、空图片区）；A21 分隔空行无边框
    for coord in ("A4", "A5", "C5", "A9", "D9", "G13", "A17", "A22", "E39", "E12", "A23", "E40"):
        assert ws[coord].border.left.style == "thin", f"{coord} 缺左边框"
        assert ws[coord].border.right.style == "thin", f"{coord} 缺右边框"
        assert ws[coord].border.top.style == "thin", f"{coord} 缺上边框"
        assert ws[coord].border.bottom.style == "thin", f"{coord} 缺下边框"
    for side in ("left", "right", "top", "bottom"):
        assert getattr(ws["A21"].border, side).style is None, f"A21 不应有 {side} 边框"
    # 空图片区对齐：原版 A23/E23/A40/E40 中心/中心
    for coord in ("A23", "E23", "A40", "E40"):
        assert ws[coord].alignment.horizontal == "center", f"{coord} 水平应居中"
        assert ws[coord].alignment.vertical == "center", f"{coord} 垂直应居中"

    # 列宽与原版一致（A–H 全部 10.5）
    for col in ("A", "B", "C", "D", "E", "F", "G", "H"):
        assert ws.column_dimensions[col].width == 10.5, f"{col} 列宽应为 10.5"

    # 行高与原版一致（标题合并区 2-3 行 14、A21 分隔空行 5，其余自适应）
    assert ws.row_dimensions[2].height == 14
    assert ws.row_dimensions[3].height == 14
    assert ws.row_dimensions[21].height == 5

    # 打印页边距与原版一致
    assert round(ws.page_margins.left, 2) == 0.7
    assert round(ws.page_margins.right, 2) == 0.7
    assert round(ws.page_margins.top, 2) == 0.75
    assert round(ws.page_margins.bottom, 2) == 0.75
    assert round(ws.page_margins.header, 2) == 0.3
    assert round(ws.page_margins.footer, 2) == 0.3

    # 字段位置
    assert ws["A4"].value == "问题类型 Type"
    assert ws["C4"].value == "项目 Project"
    assert ws["G4"].value == "反馈人 Name"
    assert ws["C5"].value == "表单测试项目"      # 项目名称回填
    assert ws["A8"].value == "车型 Model"
    assert ws["D8"].value == "车辆编号 No."
    assert ws["G8"].value == "问题来源 From"
    assert ws["A9"].value == "AGVLSYP"
    assert ws["G9"].value == "项目现场"
    assert ws["A12"].value == "问题部件 Parts"
    assert ws["G12"].value == "提交日期 Date"
    assert ws["A16"].value == "问题简述（或现象）Description"
    assert ws["A17"].value == "行驶时一侧突然断连，另一侧轮子锁死"
    assert ws["A22"].value == "问题图片（一）"
    assert ws["E22"].value == "问题图片（二）"
    assert ws["A39"].value == "问题图片（三）"
    assert ws["E39"].value == "问题图片（四）"


def test_quality_form_export_import_roundtrip(client, pm_headers):
    """系统导出的质量反馈单再导入：字段一一回填，日期为真实日期值可解析。"""
    pid = _project(client, pm_headers)
    issue = client.post(f"/api/projects/{pid}/issues", headers=pm_headers, json={
        "title": "编码器断连", "description": "行驶时一侧突然断连，另一侧轮子锁死",
        "issue_type": "hardware", "severity": "high",
        "extra": {"vehicle_model": "AGVLSYP", "vehicle_no": "77BC00255", "source": "项目现场",
                  "part": "伺服电机编码器", "feedback_date": "2024-05-08", "reporter_name": "孙晓杰"},
    }).json()
    r = client.get(f"/api/projects/{pid}/issues/{issue['id']}/export/quality-form", headers=pm_headers)
    assert r.status_code == 200, r.text
    imp = client.post(f"/api/projects/{pid}/import/quality-form", headers=pm_headers,
                      files={"file": ("质量问题反馈单.xlsx", io.BytesIO(r.content),
                                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert imp.status_code == 200, imp.text
    created = client.get(f"/api/projects/{pid}/issues", headers=pm_headers).json()
    assert len(created) == 2
    new_issue = next(i for i in created if i["id"] != issue["id"])
    assert new_issue["issue_type"] == "hardware"
    assert new_issue["description"].startswith("行驶时一侧突然断连")
    assert new_issue["extra"]["vehicle_model"] == "AGVLSYP"
    assert new_issue["extra"]["vehicle_no"] == "77BC00255"
    assert new_issue["extra"]["part"] == "伺服电机编码器"
    assert new_issue["extra"]["reporter_name"] == "孙晓杰"
    assert new_issue["extra"]["feedback_date"] == "2024-05-08"


def test_quality_form_import_invalid_layout_400(client, pm_headers):
    """错版表格（非质量问题反馈单）明确报错。"""
    pid = _project(client, pm_headers)
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "随便一张表"
    ws["A2"] = "数据"
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post(f"/api/projects/{pid}/import/quality-form", headers=pm_headers,
                    files={"file": ("错版.xlsx", buf.getvalue(),
                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 400, r.text
    assert "版式" in r.json()["detail"]


def test_quality_form_import_bad_date_400(client, pm_headers):
    """质量问题反馈单提交日期无法解析时明确报错。"""
    pid = _project(client, pm_headers)
    wb = Workbook()
    ws = wb.active
    ws.merge_cells("A1:H3")
    ws["A1"] = "产品质量问题反馈单据"
    ws["A4"] = "问题类型 Type"; ws["C4"] = "项目 Project"; ws["G4"] = "反馈人 Name"
    ws["A5"] = "硬件"; ws["C5"] = "天津环晟"; ws["G5"] = "孙晓杰"
    ws["A8"] = "车型 Model"; ws["D8"] = "车辆编号 No."; ws["G8"] = "问题来源 From"
    ws["A9"] = "AGVLSYP"; ws["D9"] = "77BC00255"; ws["G9"] = "项目现场"
    ws["A12"] = "问题部件 Parts"; ws["G12"] = "提交日期 Date"
    ws["A13"] = "伺服电机编码器"; ws["G13"] = "上周五"
    ws["A16"] = "问题简述（或现象）Description"
    ws["A17"] = "行驶时一侧突然断连"
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post(f"/api/projects/{pid}/import/quality-form", headers=pm_headers,
                    files={"file": ("质量问题反馈单.xlsx", buf.getvalue(),
                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 400, r.text
    assert "提交日期" in r.json()["detail"]


def test_quality_form_import_empty_date_400(client, pm_headers):
    """质量问题反馈单提交日期留空同样明确报错，不静默入库。"""
    pid = _project(client, pm_headers)
    wb = Workbook()
    ws = wb.active
    ws.merge_cells("A1:H3")
    ws["A1"] = "产品质量问题反馈单据"
    ws["A4"] = "问题类型 Type"; ws["C4"] = "项目 Project"; ws["G4"] = "反馈人 Name"
    ws["A5"] = "硬件"; ws["C5"] = "天津环晟"; ws["G5"] = "孙晓杰"
    ws["A8"] = "车型 Model"; ws["D8"] = "车辆编号 No."; ws["G8"] = "问题来源 From"
    ws["A9"] = "AGVLSYP"; ws["D9"] = "77BC00255"; ws["G9"] = "项目现场"
    ws["A12"] = "问题部件 Parts"; ws["G12"] = "提交日期 Date"
    ws["A13"] = "伺服电机编码器"; ws["G13"] = ""          # 提交日期留空
    ws["A16"] = "问题简述（或现象）Description"
    ws["A17"] = "行驶时一侧突然断连"
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post(f"/api/projects/{pid}/import/quality-form", headers=pm_headers,
                    files={"file": ("质量问题反馈单.xlsx", buf.getvalue(),
                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 400, r.text
    assert "提交日期" in r.json()["detail"]
    # 未入库：问题数保持为 0
    assert client.get(f"/api/projects/{pid}/issues", headers=pm_headers).json() == []


# ===== 内部联系单（物品申请） =====

def test_item_application_crud(client, pm_headers):
    pid = _project(client, pm_headers)
    created = client.post(f"/api/projects/{pid}/item-applications", headers=pm_headers, json={
        "title": "更换三支点螺栓需要吊车", "item_name": "吊车", "requirement": "更换三支点螺栓需要使用吊车，共计3000元",
        "reason": "三支点设计缺陷", "contract_no": "HCAGV-2025-170", "applicant": "孙晓杰",
        "department": "项目技术部", "apply_date": "2026-01-07",
    })
    assert created.status_code == 200, created.text
    app_id = created.json()["id"]

    listed = client.get(f"/api/projects/{pid}/item-applications", headers=pm_headers).json()
    assert len(listed) == 1
    assert listed[0]["applicant"] == "孙晓杰"

    updated = client.put(f"/api/projects/{pid}/item-applications/{app_id}", headers=pm_headers, json={
        "title": "更换三支点螺栓需要吊车", "item_name": "吊车", "requirement": "更换三支点螺栓需要使用吊车，共计3000元",
        "reason": "三支点设计缺陷", "contract_no": "HCAGV-2025-170", "applicant": "孙晓杰",
        "department": "项目技术部", "apply_date": "2026-01-07",
        "approver": "王工", "status": "approved"})
    assert updated.status_code == 200, updated.text
    assert updated.json()["approver"] == "王工"
    assert updated.json()["status"] == "approved"

    deleted = client.delete(f"/api/projects/{pid}/item-applications/{app_id}", headers=pm_headers)
    assert deleted.status_code == 200
    assert client.get(f"/api/projects/{pid}/item-applications", headers=pm_headers).json() == []


def _screenshot_item_application_payload(**overrides):
    """截图等价数据：申请部门为空、审批人=11、日期/状态枚举、金额、图号含顿号分隔。"""
    payload = {
        "title": "更换左右支架",
        "item_name": "钣金",
        "requirement": "111",
        "reason": "越南黄轮项目4.5T顶升式AGV，顶升机构上固定接近开光的支架被现场人工叉车撞击导致无法正常使用，故左右两边各需要一个支架进行更换。",
        "drawing_no": "LQPN450-780600-000、LQPN450-780400-...",
        "contract_no": "111",
        "amount": "3000元",
        "applicant": "李世豪",
        "approver": "11",
        "department": "",
        "apply_date": "2026-09-01",
        "status": "pending",
        "note": "111",
    }
    payload.update(overrides)
    return payload


def test_item_application_screenshot_equivalent_submission(client, pm_headers):
    """截图场景：等价提交数据必须成功且只创建一条记录（合法数据不 500、不落两条）。"""
    pid = _project(client, pm_headers)
    r = client.post(f"/api/projects/{pid}/item-applications", headers=pm_headers,
                    json=_screenshot_item_application_payload())
    assert r.status_code == 200, r.text
    created = r.json()
    assert created["department"] == ""       # 申请部门为空原样保存
    assert created["approver"] == "11"
    assert created["status"] == "pending"
    assert created["apply_date"] == "2026-09-01"
    assert "、" in created["drawing_no"]     # 图号含顿号分隔原样保存

    listed = client.get(f"/api/projects/{pid}/item-applications", headers=pm_headers).json()
    assert len(listed) == 1                  # 只落一条
    assert listed[0]["applicant"] == "李世豪"

    # 编辑（PUT）同样成功
    upd = client.put(f"/api/projects/{pid}/item-applications/{created['id']}",
                     headers=pm_headers, json=_screenshot_item_application_payload())
    assert upd.status_code == 200, upd.text
    assert upd.json()["item_name"] == "钣金"


def test_item_application_invalid_status_422_not_500(client, pm_headers):
    """截图状态显示「待审批」；若前端传中文标签而非枚举值，接口须返回 422 而非 500，且不落库。"""
    pid = _project(client, pm_headers)
    r = client.post(f"/api/projects/{pid}/item-applications", headers=pm_headers,
                    json=_screenshot_item_application_payload(status="待审批"))
    assert r.status_code == 422, r.text      # 明确 4xx，不是 500
    assert client.get(f"/api/projects/{pid}/item-applications", headers=pm_headers).json() == []


def test_item_application_extra_fields_422_not_500(client, pm_headers):
    """编辑弹窗 setFieldsValue({...app}) 若把 id/project_id/created_at/updated_at 也带回请求体，
    extra=forbid 返回 422（明确 4xx），后端不允许静默接受额外字段，也不能因此抛 500。"""
    pid = _project(client, pm_headers)
    payload = _screenshot_item_application_payload()
    payload.update({
        "id": "should-not-accept",
        "project_id": pid,
        "created_at": "2026-09-01T00:00:00",
        "updated_at": "2026-09-01T00:00:00",
    })
    r = client.post(f"/api/projects/{pid}/item-applications", headers=pm_headers, json=payload)
    assert r.status_code == 422, r.text
    assert client.get(f"/api/projects/{pid}/item-applications", headers=pm_headers).json() == []


def test_item_application_field_length_limits_422_not_500(client, pm_headers):
    """边界字段：图号/合同号超 100 字符、备注超 5000 字符 → 422 且不落库；申请部门为空 → 200。"""
    pid = _project(client, pm_headers)
    overlong = "x" * 101
    for override in ({"drawing_no": overlong}, {"contract_no": overlong}, {"note": "y" * 5001}):
        r = client.post(f"/api/projects/{pid}/item-applications", headers=pm_headers,
                        json=_screenshot_item_application_payload(**override))
        assert r.status_code == 422, r.text
    assert client.get(f"/api/projects/{pid}/item-applications", headers=pm_headers).json() == []
    ok = client.post(f"/api/projects/{pid}/item-applications", headers=pm_headers,
                     json=_screenshot_item_application_payload())
    assert ok.status_code == 200, ok.text


def test_item_application_docx_export(client, pm_headers):
    pid = _project(client, pm_headers)
    app = client.post(f"/api/projects/{pid}/item-applications", headers=pm_headers, json={
        "title": "更换三支点螺栓需要吊车", "item_name": "吊车",
        "requirement": "更换三支点螺栓需要使用吊车，共计3000元", "reason": "三支点设计缺陷",
        "contract_no": "HCAGV-2025-170", "applicant": "孙晓杰", "department": "项目技术部",
        "apply_date": "2026-01-07"}).json()
    r = client.get(f"/api/projects/{pid}/item-applications/{app['id']}/export", headers=pm_headers)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml")
    assert zipfile.is_zipfile(io.BytesIO(r.content))


def _item_application_docx() -> bytes:
    from docx import Document
    doc = Document()
    doc.add_paragraph("内部联系单——呼伦贝尔驰宏矿业项目")
    doc.add_paragraph("需求：更换三支点螺栓需要使用吊车，共计3000元")
    doc.add_paragraph("原因：三支点设计缺陷")
    doc.add_paragraph("图号：DWG-01")
    doc.add_paragraph("合同号：HCAGV-2025-170")
    doc.add_paragraph("申请人：孙晓杰")
    doc.add_paragraph("审批人：")
    doc.add_paragraph("项目技术部")
    doc.add_paragraph("2026/1/7")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


def test_item_application_docx_import(client, pm_headers):
    pid = _project(client, pm_headers)
    r = client.post(f"/api/projects/{pid}/import/item-application", headers=pm_headers,
                    files={"file": ("内部联系单-吊车.docx", _item_application_docx(),
                                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert r.status_code == 200, r.text
    app_id = r.json()["id"]
    apps = client.get(f"/api/projects/{pid}/item-applications", headers=pm_headers).json()
    assert len(apps) == 1
    app = apps[0]
    assert app["requirement"] == "更换三支点螺栓需要使用吊车，共计3000元"
    assert app["reason"] == "三支点设计缺陷"
    assert app["contract_no"] == "HCAGV-2025-170"
    assert app["applicant"] == "孙晓杰"
    assert app["department"] == "项目技术部"
    assert app["amount"] == "3000元"
    assert app["item_name"] == "吊车"
    assert app["apply_date"] == "2026-01-07"
    assert app_id == app["id"]


def test_item_application_docx_roundtrip(client, pm_headers):
    """系统生成 Word → 回灌导入：申请事项/申请物品/金额/申请日期等字段不丢。"""
    pid = _project(client, pm_headers)
    created = client.post(f"/api/projects/{pid}/item-applications", headers=pm_headers, json={
        "title": "更换三支点螺栓需要吊车", "item_name": "吊车",
        "requirement": "更换三支点螺栓需要使用吊车，共计3000元", "reason": "三支点设计缺陷",
        "drawing_no": "DWG-01", "contract_no": "HCAGV-2025-170", "amount": "3000元",
        "applicant": "孙晓杰", "approver": "王工", "department": "项目技术部",
        "apply_date": "2026-01-07"}).json()
    r = client.get(f"/api/projects/{pid}/item-applications/{created['id']}/export", headers=pm_headers)
    assert r.status_code == 200, r.text
    imp = client.post(f"/api/projects/{pid}/import/item-application", headers=pm_headers,
                      files={"file": ("内部联系单.docx", io.BytesIO(r.content),
                                      "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert imp.status_code == 200, imp.text
    app = client.get(f"/api/projects/{pid}/item-applications/{imp.json()['id']}", headers=pm_headers).json()
    assert app["title"] == "更换三支点螺栓需要吊车"      # 申请事项
    assert app["item_name"] == "吊车"                   # 申请物品/服务
    assert app["requirement"] == "更换三支点螺栓需要使用吊车，共计3000元"
    assert app["reason"] == "三支点设计缺陷"
    assert app["drawing_no"] == "DWG-01"
    assert app["contract_no"] == "HCAGV-2025-170"
    assert app["amount"] == "3000元"                    # 金额
    assert app["applicant"] == "孙晓杰"
    assert app["approver"] == "王工"
    assert app["department"] == "项目技术部"
    assert app["apply_date"] == "2026-01-07"            # 申请日期


def test_legacy_repair_keeps_already_applied_event_schema_and_dates(client, pm_headers, upgrade_legacy_schema):
    from app.db import engine
    from sqlalchemy import text
    pid = _project(client, pm_headers)
    changed = client.put(f'/api/projects/{pid}', headers=pm_headers, json=versioned(pid, {'dispatch_platform': 'NDC'}))
    assert changed.status_code == 200, changed.text
    dated = client.put(f'/api/projects/{pid}/events/ship', headers=pm_headers,
                       json={'planned_date': '2026-10-27'})
    assert dated.status_code == 200, dated.text
    def state():
        with engine.connect() as connection:
            return (connection.execute(text('SELECT version_num FROM alembic_version')).scalar_one(),
                    connection.execute(text('SELECT dispatch_platform FROM projects WHERE id=:id'), {'id': pid}).scalar_one(),
                    connection.execute(text('SELECT event_key, planned_date FROM project_events WHERE project_id=:id'), {'id': pid}).all())
    before = state()
    upgrade_legacy_schema()
    assert state() == before
