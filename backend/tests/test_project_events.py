"""T20 actual-event authority, date precision and complete-fleet arrival evidence."""
from tests.lifecycle_helpers import versioned
from datetime import date, datetime

import pytest


def project(client, headers, **extra):
    response = client.post('/api/projects', headers=headers, json={
        'name': '关键事件项目', 'start_date': '2026-01-01', 'end_date': '2026-12-31', **extra,
    })
    assert response.status_code == 200, response.text
    return response.json()


def events(client, headers, pid):
    response = client.get(f'/api/projects/{pid}/events', headers=headers)
    assert response.status_code == 200, response.text
    return {item['event_key']: item for item in response.json()['events']}


def fleet(client, headers, quantity=2):
    p = project(client, headers, timezone='Asia/Shanghai')
    r = client.post(f"/api/projects/{p['id']}/vehicles", headers=headers,
                    json={'vehicle_type': '潜伏式', 'model_name': 'AGV', 'quantity': quantity})
    assert r.status_code == 200, r.text
    return p, r.json()['id']


def unit(client, headers, pid, vid, code):
    r = client.post(f'/api/projects/{pid}/vehicle-units', headers=headers,
                   json={'asset_code': code, 'vehicle_id': vid})
    assert r.status_code == 200, r.text
    return r.json()['id']


def test_plans_and_platform_are_one_source_and_do_not_move_status(client, pm_headers):
    p = project(client, pm_headers, dispatch_platform='NDC', event_plans={
        'internal_kickoff': '2026-01-02', 'ship': '2026-01-10', 'arrival_all': '2026-01-12'})
    assert p['dispatch_platform'] == 'NDC'
    assert p['planned_entry_date'] is None
    e = events(client, pm_headers, p['id'])
    assert len(e) == 9 and e['ship']['planned_date'] == '2026-01-10'
    assert all(x['actual_at'] is None for x in e.values())
    r = client.put(f"/api/projects/{p['id']}", headers=pm_headers,
                   json=versioned(p['id'], {'dispatch_platform': 'integrated', 'event_plans': {'ship': None}}))
    assert r.status_code == 200, r.text
    assert r.json()['dispatch_platform'] == 'integrated'
    assert events(client, pm_headers, p['id'])['ship']['planned_date'] is None
    assert r.json()['status'] == 'draft'


@pytest.mark.parametrize('body', [
    {'planned_entry_date': '2026-01-10'}, {'dispatch_platform': 'invalid'},
    {'event_plans': {'unknown': '2026-01-10'}},
])
def test_invalid_or_retired_dates_rejected_on_create(client, pm_headers, body):
    r = client.post('/api/projects', headers=pm_headers, json={
        'name': '拒绝无效日期', 'start_date': '2026-01-01', 'end_date': '2026-12-31', **body})
    assert r.status_code == 422, r.text


def test_manual_event_keeps_instant_and_clear_requires_audited_reason(client, pm_headers):
    from app.db import SessionLocal
    from app.models import ChangeLog, ProjectEvent
    p = project(client, pm_headers)
    path = f"/api/projects/{p['id']}/events/ship"
    assert client.put(path, headers=pm_headers, json={'actual_at': '2026-10-20T14:32:21+08:00'}).status_code == 200
    assert events(client, pm_headers, p['id'])['ship']['actual_at'] == '2026-10-20T06:32:21+00:00'
    assert client.put(path, headers=pm_headers, json={'actual_at': None, 'reason': '  '}).status_code == 400
    assert events(client, pm_headers, p['id'])['ship']['actual_at'] is not None
    assert client.put(path, headers=pm_headers, json={'actual_at': None, 'reason': '误登记发车'}).status_code == 200
    with SessionLocal() as db:
        row = db.query(ProjectEvent).filter_by(project_id=p['id'], event_key='ship').one()
        assert row.actual_at is None and row.recorded_by
        audit = db.query(ChangeLog).filter_by(project_id=p['id'], entity_type='project_event').all()
        assert any('误登记发车' in x.detail and x.old_value['actual_at'] for x in audit)
    assert client.get(f"/api/projects/{p['id']}", headers=pm_headers).json()['status'] == 'draft'
    assert client.put(path, headers=pm_headers, json={'actual_at': '2026-10-20T14:32:21'}).status_code == 422


def test_partial_registry_cannot_become_all_arrived_and_clear_is_atomic(client, pm_headers):
    p, vid = fleet(client, pm_headers)
    pid = p['id']; first = unit(client, pm_headers, pid, vid, 'AGV-1')
    path = f'/api/projects/{pid}/vehicle-arrivals'
    r = client.put(path, headers=pm_headers, json={'unit_ids': [first], 'arrived_at': '2026-10-01'})
    assert r.status_code == 200, r.text
    assert r.json()['arrival_summary']['registry_complete'] is False
    assert events(client, pm_headers, pid)['arrival_all']['actual_at'] is None
    second = unit(client, pm_headers, pid, vid, 'AGV-2')
    r = client.put(path, headers=pm_headers, json={'unit_ids': [second], 'arrived_at': '2026-10-03'})
    assert r.status_code == 200, r.text
    all_arrived = events(client, pm_headers, pid)['arrival_all']
    assert all_arrived['actual_date'] == '2026-10-03'
    assert all_arrived['actual_at'] == '2026-10-02T16:00:00+00:00'
    assert all_arrived['derived'] and all_arrived['actual_precision'] == 'date'
    assert client.put(f'/api/projects/{pid}/events/arrival_all', headers=pm_headers,
                      json=versioned(pid, {'actual_at': None, 'reason': '绕过逐台记录'})).status_code == 400
    assert client.put(path, headers=pm_headers, json={'unit_ids': [first, second], 'arrived_at': None}).status_code == 400
    assert all(x['arrived_at'] for x in client.get(f'/api/projects/{pid}/vehicle-units', headers=pm_headers).json())
    assert client.put(path, headers=pm_headers, json={'unit_ids': [first, second], 'arrived_at': None, 'reason': '误录到场'}).status_code == 200
    assert events(client, pm_headers, pid)['arrival_all']['actual_at'] is None


def test_fleet_expansion_invalidates_derived_actual_without_changing_project_state(client, pm_headers):
    p, vid = fleet(client, pm_headers, quantity=1); pid = p['id']
    uid = unit(client, pm_headers, pid, vid, 'AGV-1')
    client.put(f'/api/projects/{pid}/vehicle-arrivals', headers=pm_headers, json=versioned(pid, {'unit_ids': [uid], 'arrived_at': '2026-10-01'}))
    assert events(client, pm_headers, pid)['arrival_all']['actual_at']
    r = client.put(f'/api/projects/{pid}/vehicles/{vid}', headers=pm_headers, json={'quantity': 2})
    assert r.status_code == 200, r.text
    assert events(client, pm_headers, pid)['arrival_all']['actual_at'] is None
    assert client.get(f'/api/projects/{pid}', headers=pm_headers).json()['status'] == 'draft'


def test_batch_does_not_write_foreign_vehicle_and_member_cannot_change_events(client, pm_headers, worker_headers):
    p, vid = fleet(client, pm_headers, quantity=1); pid = p['id']
    own = unit(client, pm_headers, pid, vid, 'AGV-1')
    other = project(client, pm_headers)
    foreign = unit(client, pm_headers, other['id'], None, 'AGV-other')
    r = client.put(f'/api/projects/{pid}/vehicle-arrivals', headers=pm_headers,
                   json=versioned(pid, {'unit_ids': [own, foreign], 'arrived_at': '2026-10-01'}))
    assert r.status_code == 404
    assert client.get(f'/api/projects/{pid}/vehicle-units', headers=pm_headers).json()[0]['arrived_at'] is None
    client.put(f'/api/projects/{pid}', headers=pm_headers, json=versioned(pid, {'members': ['user-worker-001']}))
    assert client.put(f'/api/projects/{pid}/events/ship', headers=worker_headers,
                      json=versioned(pid, {'actual_at': '2026-10-01T10:30:00Z'})).status_code == 403
    assert client.put(f'/api/projects/{pid}/vehicle-arrivals', headers=worker_headers,
                      json=versioned(pid, {'unit_ids': [own], 'arrived_at': '2026-10-01'})).status_code == 403
    for payload in ({'planned_entry_date': None}, {'entry_plan': {'entry_date': '2026-10-01'}}):
        assert client.put(f'/api/projects/{pid}', headers=pm_headers, json=versioned(pid, payload)).status_code == 422


def test_portable_json_and_excel_preserve_event_precision_and_vehicle_arrival(client, pm_headers):
    from io import BytesIO
    from openpyxl import load_workbook
    p, vid = fleet(client, pm_headers, quantity=1); pid = p['id']
    client.put(f'/api/projects/{pid}', headers=pm_headers, json=versioned(pid, {'dispatch_platform': 'GRACE', 'event_plans': {'ship': '2026-09-30'}}))
    uid = unit(client, pm_headers, pid, vid, 'AGV-001')
    client.put(f'/api/projects/{pid}/vehicle-arrivals', headers=pm_headers, json=versioned(pid, {'unit_ids': [uid], 'arrived_at': '2026-10-03'}))
    client.put(f'/api/projects/{pid}/events/ship', headers=pm_headers, json=versioned(pid, {'actual_at': '2026-10-01T14:30:00+08:00'}))
    snapshot = client.get(f'/api/projects/{pid}/export/json', headers=pm_headers).json()
    assert snapshot['vehicle_units'][0]['arrived_at'] == '2026-10-03'
    assert snapshot['project']['dispatch_platform'] == 'GRACE'
    r = client.post('/api/projects/import/json', headers=pm_headers, json=snapshot)
    assert r.status_code == 200, r.text
    restored = r.json()['id']
    before, after = events(client, pm_headers, pid), events(client, pm_headers, restored)
    for key in before:
        assert (after[key]['planned_date'], after[key]['actual_at']) == (before[key]['planned_date'], before[key]['actual_at'])
    workbook = client.get(f'/api/projects/{pid}/export/excel', headers=pm_headers)
    assert workbook.status_code == 200, workbook.text[:100] if workbook.status_code != 200 else ''
    book = load_workbook(BytesIO(workbook.content))
    assert '关键事件' in book.sheetnames
    assert any('2026-10-03' in tuple(row) for row in book['关键事件'].values)
    assert any('到场日期' in tuple(row) for row in book['现场车辆'].values)


def test_old_portable_dates_map_to_ship_and_arrival_without_actual_times(client, pm_headers):
    payload = {'type': 'agvpm-project', 'project': {'name': '旧格式', 'start_date': '2026-01-01', 'end_date': '2026-12-31',
        'planned_entry_date': '2026-10-01', 'entry_plan': {'entry_date': '2026-10-03', 'meeting_point': '东门'}}}
    r = client.post('/api/projects/import/json', headers=pm_headers, json=payload)
    assert r.status_code == 200, r.text
    value = events(client, pm_headers, r.json()['id'])
    assert value['ship']['planned_date'] == '2026-10-01'
    assert value['arrival_all']['planned_date'] == '2026-10-03'
    assert all(x['actual_at'] is None for x in value.values())


def test_plan_grant_can_record_arrival_without_receiving_event_update_authority(client, pm_headers, worker_headers):
    p, vid = fleet(client, pm_headers, quantity=1); pid = p['id']
    uid = unit(client, pm_headers, pid, vid, 'AGV-1')
    grant = client.put(f'/api/projects/{pid}', headers=pm_headers, json=versioned(pid, {
        'members': ['user-worker-001'], 'member_permissions': [{'user_id': 'user-worker-001', 'capabilities': ['project.manage_plan']}]}))
    assert grant.status_code == 200, grant.text
    assert client.put(f'/api/projects/{pid}/vehicle-arrivals', headers=worker_headers,
                      json=versioned(pid, {'unit_ids': [uid], 'arrived_at': '2026-10-03'})).status_code == 200
    assert client.put(f'/api/projects/{pid}/events/ship', headers=worker_headers,
                      json=versioned(pid, {'actual_at': '2026-10-01T10:30:00Z'})).status_code == 403
    overview = client.get(f'/api/projects/{pid}/overview', headers=worker_headers).json()
    assert overview['project_events']['events'] == list(events(client, worker_headers, pid).values())


def test_postgres_concurrent_batches_keep_both_dates_and_derive_last_arrival(client, pm_headers):
    from concurrent.futures import ThreadPoolExecutor
    from app.db import engine
    if engine.dialect.name != 'postgresql':
        pytest.skip('real PostgreSQL aggregate serialization gate')
    p, vid = fleet(client, pm_headers); pid = p['id']
    ids = [unit(client, pm_headers, pid, vid, f'AGV-{i}') for i in range(2)]
    with ThreadPoolExecutor(max_workers=2) as pool:
        futures = [pool.submit(client.put, f'/api/projects/{pid}/vehicle-arrivals', headers=pm_headers,
                               json={'unit_ids': [uid], 'arrived_at': f'2026-10-0{i+1}'}) for i, uid in enumerate(ids)]
        responses = [future.result(timeout=20) for future in futures]
    assert [response.status_code for response in responses] == [200, 200]
    units = client.get(f'/api/projects/{pid}/vehicle-units', headers=pm_headers).json()
    assert sorted(u['arrived_at'] for u in units) == ['2026-10-01', '2026-10-02']
    assert events(client, pm_headers, pid)['arrival_all']['actual_date'] == '2026-10-02'


def test_readonly_legacy_date_survives_handover_clear_and_clone_keeps_only_plans(client, pm_headers):
    r = client.post('/api/projects/import/json', headers=pm_headers, json={
        'type': 'agvpm-project', 'project': {'name': '历史计划', 'start_date': '2026-01-01', 'end_date': '2026-12-31', 'dispatch_platform': 'NDC',
                    'planned_entry_date': '2026-10-01', 'entry_plan': {'entry_date': '2026-10-03', 'meeting_point': '东门'}}})
    assert r.status_code == 200, r.text
    pid = r.json()['id']
    cleared = client.put(f'/api/projects/{pid}', headers=pm_headers, json=versioned(pid, {'entry_plan': None}))
    assert cleared.status_code == 200, cleared.text
    assert cleared.json()['entry_plan'] == {'entry_date': '2026-10-03'}
    client.put(f'/api/projects/{pid}/events/ship', headers=pm_headers, json=versioned(pid, {'actual_at': '2026-10-01T10:00:00Z'}))
    cloned = client.post(f'/api/projects/{pid}/clone', headers=pm_headers, json={'name': '只复制计划'})
    assert cloned.status_code == 200, cloned.text
    detail = client.get(f"/api/projects/{cloned.json()['id']}", headers=pm_headers).json()
    assert detail['dispatch_platform'] == 'NDC'
    assert next(e for e in detail['events'] if e['event_key'] == 'ship')['planned_date'] == '2026-10-01'
    assert all(e['actual_at'] is None for e in detail['events'])


def test_zero_quantity_legacy_fleet_never_claims_all_arrived(client, pm_headers):
    from app.db import SessionLocal
    from app.models import Vehicle
    p, vid = fleet(client, pm_headers, quantity=1)
    # An old or externally repaired zero-count row is not evidence of any arrival.
    with SessionLocal() as db:
        db.get(Vehicle, vid).quantity = 0
        db.commit()
    r = client.get(f"/api/projects/{p['id']}/events", headers=pm_headers)
    assert r.status_code == 200, r.text
    assert r.json()['arrival_summary']['registry_complete'] is False
    assert r.json()['arrival_summary']['all_arrived_date'] is None
