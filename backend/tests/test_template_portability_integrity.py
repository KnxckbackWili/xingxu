"""模板 JSON/Excel 往返及审计原子性。"""

import io
import pytest


def test_template_json_can_import_into_another_template(client, admin_headers):
    exported = client.get("/api/templates/default/export/json", headers=admin_headers)
    assert exported.status_code == 200
    created = client.post("/api/templates", headers=admin_headers, json={"name": "导入目标"})
    assert created.status_code == 200
    target_id = created.json()["id"]
    payload = exported.json()
    payload["replace_missing_phases"] = True
    imported = client.post(f"/api/templates/{target_id}/import/json", headers=admin_headers, json=payload)
    assert imported.status_code == 200, imported.text
    target = client.get(f"/api/templates/{target_id}/tree", headers=admin_headers).json()
    source = client.get("/api/templates/default/tree", headers=admin_headers).json()
    assert [phase["name"] for phase in target] == [phase["name"] for phase in source]
    source_ids = {task["id"] for phase in source for module in phase["modules"] for task in module["tasks"]}
    target_ids = {task["id"] for phase in target for module in phase["modules"] for task in module["tasks"]}
    assert target_ids and not target_ids.intersection(source_ids)
    # The ordinary save path still rejects IDs belonging to another template.
    rejected = client.put(f"/api/templates/{target_id}/tree", headers=admin_headers, json=payload)
    assert rejected.status_code == 400


def test_excel_roundtrip_keeps_task_identity_metadata_and_reference_docs(client, admin_headers):
    from app.db import SessionLocal
    from app.models import Module, Task
    created = client.post("/api/templates", headers=admin_headers, json={"name": "Excel完整性"})
    template_id = created.json()["id"]
    with SessionLocal() as db:
        task = db.query(Task).filter_by(template_id=template_id).first()
        task_id = task.id
        task.description = "不可在重新导入后丢失的施工规范"
        task.applicable_vehicle_types = ["潜伏式"]
        module = db.get(Module, task.module_id)
        module.preparation_conditions = ["need_wms"]
        db.commit()
    definition = client.post(f'/api/templates/{template_id}/document-requirements', headers=admin_headers,
        json={'stage_key': '6', 'code': '6.99', 'name': '签字验收单', 'scope': 'project', 'timing_type': 'on_demand'})
    assert definition.status_code == 200, definition.text
    base = f"/api/templates/{template_id}/tasks/{task_id}/reference-docs"
    uploaded = client.post(base, headers=admin_headers, files={"file": ("说明.txt", b"important reference", "text/plain")})
    assert uploaded.status_code == 200
    workbook = client.get("/api/templates/export/excel", headers=admin_headers, params={"template_id": template_id})
    imported = client.post("/api/templates/import/excel", headers=admin_headers, params={"template_id": template_id},
                           files={"file": ("roundtrip.xlsx", workbook.content)})
    assert imported.status_code == 200, imported.text
    with SessionLocal() as db:
        task = db.get(Task, task_id)
        assert task is not None, "Excel import must not break generated-project source identities"
        assert task.description == "不可在重新导入后丢失的施工规范"
        assert task.applicable_vehicle_types == ["潜伏式"]
        assert db.get(Module, task.module_id).preparation_conditions == ["need_wms"]
    requirements = client.get(f'/api/templates/{template_id}/document-requirements', headers=admin_headers).json()
    assert any(row['code'] == '6.99' and row['name'] == '签字验收单' for row in requirements)
    downloaded = client.get(f"{base}/{uploaded.json()['id']}/download", headers=admin_headers)
    assert downloaded.status_code == 200
    assert downloaded.content == b"important reference"


def test_excel_import_and_audit_are_one_transaction(client, admin_headers, monkeypatch):
    from app.db import SessionLocal
    from app.models import Task
    from app.routers import io_routes
    from openpyxl import load_workbook
    workbook = client.get("/api/templates/export/excel", headers=admin_headers)
    sheet = load_workbook(io.BytesIO(workbook.content))
    original = sheet.active.cell(3, 4).value
    sheet.active.cell(3, 4).value = "不会成功保存的改名"
    content = io.BytesIO()
    sheet.save(content)
    def reject_audit(*args, **kwargs):
        raise RuntimeError("simulated audit write failure")
    monkeypatch.setattr(io_routes, "log_change", reject_audit)
    with pytest.raises(RuntimeError, match="simulated audit write failure"):
        client.post("/api/templates/import/excel", headers=admin_headers,
                    files={"file": ("changed.xlsx", content.getvalue())})
    with SessionLocal() as db:
        assert db.query(Task).filter(Task.name == "不会成功保存的改名").count() == 0
        assert db.query(Task).filter(Task.name == original).count() > 0
