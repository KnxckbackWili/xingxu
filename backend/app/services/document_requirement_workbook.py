"""Portable document-definition sheets; binary chunks stay below Excel cell limits."""
import io
import json
import base64
import zipfile
from openpyxl.styles import Font, PatternFill

from app.services.document_requirements import parse_bundle

COLUMNS = (
    ('stage_key', '阶段'), ('code', '编号（文本）'), ('name', '资料名称'), ('template_no', '模板编号'),
    ('is_group', '是否分组'), ('parent_code', '父分组编号'), ('required', '是否必需'), ('scope', '范围'),
    ('platform', '平台'), ('timing_type', '输出时机'), ('timing_event', '事件'),
    ('timing_offset_hours', '偏移小时'), ('timing_weekday', '星期（周一1至周日7）'), ('trigger_type', '触发类型'),
    ('confidential_default', '默认保密'), ('external_visible_default', '默认对外可见'),
    ('external_never_visible', '禁止外部查看'), ('sort_order', '排序'), ('note', '说明'), ('blank_file_key', '空白模板文件键'),
)
BOOLS = {'is_group', 'required', 'confidential_default', 'external_visible_default', 'external_never_visible'}
NUMBERS = {'timing_offset_hours', 'timing_weekday', 'sort_order'}


def check_workbook_size(data):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 120 * 1024 * 1024:
                raise ValueError('工作簿解压内容过大，未修改现有模板')
    except zipfile.BadZipFile:
        raise ValueError('不是有效的xlsx工作簿') from None


def _text_cells(ws):
    for row in ws:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'


def add_sheets(wb, bundle):
    ws = wb.create_sheet('资料要求')
    ws.append(['资料要求 · 编号请保持文本；星期采用周一1至周日7；文件在隐藏表中保存'])
    ws.append([label for _, label in COLUMNS])
    for item in bundle['document_requirements']:
        ws.append([('是' if item.get(key) else '否') if key in BOOLS else item.get(key) for key, _ in COLUMNS])
    ws.freeze_panes = 'A3'
    ws.auto_filter.ref = f'A2:T{ws.max_row}'
    for cell in ws[2]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F4E79')
    for letter in ('A', 'B', 'E', 'G', 'H', 'I'):
        ws.column_dimensions[letter].width = 15
    ws.column_dimensions['C'].width = 42
    ws.column_dimensions['S'].width = 65
    _text_cells(ws)
    assets = wb.create_sheet('空白模板文件')
    assets.append(['key', 'filename', 'content_type', 'file_size', 'sha256', 'chunk_index', 'chunk_count', 'base64'])
    for key, item in bundle['document_template_files'].items():
        chunks = [item['data'][i:i + 30000] for i in range(0, len(item['data']), 30000)]
        for index, chunk in enumerate(chunks):
            assets.append([key, item['filename'], item['content_type'], item['file_size'], item['sha256'], index, len(chunks), chunk])
    assets.sheet_state = 'hidden'
    _text_cells(assets)
    provenance = wb.create_sheet('资料历史来源')
    provenance.append(['code','chunk_index','chunk_count','base64'])
    for definition in bundle['document_requirements']:
        if not definition.get('source_metadata'):continue
        encoded=base64.b64encode(json.dumps(definition['source_metadata'],ensure_ascii=False).encode()).decode()
        chunks=[encoded[i:i+30000] for i in range(0,len(encoded),30000)]
        for index,chunk in enumerate(chunks):provenance.append([definition['code'],index,len(chunks),chunk])
    provenance.sheet_state='hidden'
    _text_cells(provenance)


def parse_sheets(wb):
    if '资料要求' not in wb.sheetnames:
        return None
    ws = wb['资料要求']
    if tuple(cell.value for cell in ws[2]) != tuple(label for _, label in COLUMNS):
        raise ValueError('资料要求表头不完整或顺序已改变，未修改现有模板')
    if ws.max_row > 2002:
        raise ValueError('资料要求超过2000项')
    definitions = []
    for number, values in enumerate(ws.iter_rows(min_row=3, values_only=True), 3):
        if not any(value is not None for value in values):
            continue
        item = {}
        for (key, _), value in zip(COLUMNS, values):
            if key in BOOLS:
                if value not in ('是', '否', True, False):
                    raise ValueError(f'资料要求第{number}行{key}必须是“是”或“否”')
                value = value in ('是', True)
            elif key in NUMBERS:
                if value in (None, ''):
                    value = 0 if key == 'sort_order' else None
                elif isinstance(value, bool) or not isinstance(value, (int, float)) or int(value) != value:
                    raise ValueError(f'资料要求第{number}行{key}必须是整数')
                else:
                    value = int(value)
            elif key == 'stage_key':
                value = str(int(value)) if isinstance(value, (int, float)) and int(value) == value else str(value or '')
            elif key in ('parent_code', 'blank_file_key'):
                # Keep the cell type for the shared portable-reference validator.
                value = None if value == '' else value
            elif key in ('name', 'code', 'template_no', 'note'):
                value = str(value or '')
            else:
                value = str(value) if value not in (None, '') else None
            item[key] = value
        definitions.append(item)
    files = {}
    if '空白模板文件' in wb.sheetnames:
        sheet = wb['空白模板文件']
        if sheet.max_row > 10000:
            raise ValueError('空白模板文件分片过多')
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not any(value is not None for value in row):
                continue
            if len(row) != 8:
                raise ValueError('空白模板文件分片格式无效')
            key, name, content_type, size, sha, index, count, chunk = row
            if not isinstance(key, str) or type(index) is not int or type(count) is not int or not 0 <= index < count <= 1000 or not isinstance(chunk, str) or len(chunk) > 30000:
                raise ValueError('空白模板文件分片索引或大小无效')
            metadata = {'filename': name, 'content_type': content_type, 'file_size': size, 'sha256': sha}
            item = files.setdefault(key, {'metadata': metadata, 'count': count, 'chunks': {}})
            if metadata != item['metadata'] or count != item['count'] or index in item['chunks']:
                raise ValueError('空白模板文件分片重复或配置冲突')
            item['chunks'][index] = chunk
    packed = {}
    for key, item in files.items():
        if len(item['chunks']) != item['count']:
            raise ValueError('空白模板文件分片不完整')
        packed[key] = {**item['metadata'], 'data': ''.join(item['chunks'][index] for index in range(item['count']))}
    sources={}
    if '资料历史来源' in wb.sheetnames:
        sheet=wb['资料历史来源']
        if sheet.max_row>10000:raise ValueError('历史来源分片过多')
        for code,index,count,chunk in sheet.iter_rows(min_row=2,values_only=True):
            if not isinstance(code,str) or len(code)>40 or type(index) is not int or type(count) is not int or not 0<=index<count<=1000 or not isinstance(chunk,str) or len(chunk)>30000:
                raise ValueError('历史来源分片格式无效')
            source=sources.setdefault(code,{'count':count,'chunks':{}})
            if source['count']!=count or index in source['chunks']:raise ValueError('历史来源分片重复或冲突')
            source['chunks'][index]=chunk
        total=0
        by_code={row['code']:row for row in definitions}
        for code,source in sources.items():
            if code not in by_code or len(source['chunks'])!=source['count']:raise ValueError('历史来源分片无归属或不完整')
            encoded=''.join(source['chunks'][index] for index in range(source['count']))
            total+=len(encoded)
            if total>14*1024*1024:raise ValueError('历史来源内容过大')
            try:metadata=json.loads(base64.b64decode(encoded,validate=True))
            except (ValueError,TypeError):raise ValueError('历史来源编码无效') from None
            by_code[code]['source_metadata']=metadata
    return parse_bundle({'document_requirements': definitions, 'document_template_files': packed})
