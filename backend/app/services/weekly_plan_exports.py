"""A readable weekly workbook, with explicit draft/source/unknown labels."""
from collections import OrderedDict
from io import BytesIO
import json
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _text(ws, row, col, value):
    if isinstance(value,(dict,list)):value=json.dumps(value,ensure_ascii=False)
    cell = ws.cell(row, col, value)
    if isinstance(value, str):
        cell.data_type = 's'  # A project/task/note beginning '=' is text, never a formula.
    cell.alignment = Alignment(vertical='top', wrap_text=True)
    return cell


def scope_label(cell, units):
    scope = cell.get('scope') or {}
    if scope.get('label'):
        return scope['label']
    if scope.get('vehicle_type'):
        return scope['vehicle_type']
    ids = scope.get('vehicle_unit_ids') or []
    return '、'.join(units.get(identity, '原车号待核对') for identity in ids) if ids else '全部适用车辆' if cell['kind'] == 'vehicle' else '工作项'


def cell_text(cell):
    facts = cell['completion']
    goal = '进行中' if cell['mode'] == 'in_progress' else f"完成 {cell['target_count']} 台" if cell['kind'] == 'vehicle' and cell['target_count'] is not None else '完成台数待核对' if cell['kind'] == 'vehicle' else '完成工作'
    status = {'unknown':'范围待核对','completed':'已完成','overdue':'已逾期','in_progress':'进行中','planned':'计划中'}[facts['state']]
    people = '、'.join(person['name'] for person in cell['assignees']) or '负责人待认领/指派'
    detail = f'{goal} · {"已发布" if cell["published"] else "未发布"}\n{people}\n{status}'
    if cell['kind'] == 'vehicle' and facts.get('required_count') is not None:
        detail += f'（当前有效 {facts.get("completed_count", 0)} / 累计目标 {facts["required_count"]}）'
    if facts.get('unidentified_completed'):
        detail += '\n含迁移完成、车号未知；不代表逐台核验'
    if cell.get('notes'):
        detail += '\n' + cell['notes']
    return detail


def excel(project, data, *, phase_id=None):
    workbook = Workbook()
    sheet = workbook.active; sheet.title = '周计划'
    weeks = [row['week_start'] for row in data['weeks']]
    headers = ['阶段', '工作项', '范围'] + [f'{week}\n至 {data["weeks"][index]["week_end"]}' for index, week in enumerate(weeks)]
    _text(sheet, 1, 1, project.name + ' · 周计划')
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    _text(sheet, 2, 1, f'现场时区：{data["timezone"]}；计划负责人不代表实际执行人；进行中不判逾期。')
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    for index, header in enumerate(headers, 1):
        cell = _text(sheet, 3, index, header)
        cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='234C70')
    groups = OrderedDict()
    for cell in data['cells']:
        if phase_id and cell['phase_id'] != phase_id:
            continue
        groups.setdefault((cell['task_id'], cell['scope_key']), []).append(cell)
    units = {unit['id']: unit['label'] for unit in data['vehicles']}
    for number, cells in enumerate(groups.values(), 4):
        first = cells[0]
        for col, value in enumerate([first['phase_name'], first['task_name'], scope_label(first, units)], 1):
            _text(sheet, number, col, value)
        by_week = {cell['week_start']: cell for cell in cells}
        for index, week in enumerate(weeks, 4):
            value = cell_text(by_week[week]) if week in by_week else '未安排'
            _text(sheet, number, index, value)
        sheet.row_dimensions[number].height = 88
    sheet.freeze_panes = 'D4'
    sheet.auto_filter.ref = f'A3:{get_column_letter(len(headers))}{max(3,sheet.max_row)}'
    for col in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(col)].width = 24 if col <= 3 else 35
    sheet.print_title_rows = '1:3'
    sheet.page_setup.orientation = 'landscape'; sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.sheet_properties.pageSetUpPr.fitToPage = True; sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 0
    warnings = workbook.create_sheet('核对提示')
    for col, value in enumerate(['类型', '核对说明'], 1):
        _text(warnings, 1, col, value).font = Font(bold=True)
    allowed_cells = {cell['id'] for cell in data['cells'] if not phase_id or cell['phase_id'] == phase_id}
    allowed_work = {item['id'] for phase in data['work_items'] if not phase_id or phase['id'] == phase_id for item in phase['work_items']}
    selected_warnings = [warning for warning in data['warnings'] if
        (not warning.get('task_id') or warning['task_id'] in allowed_work) and
        (not warning.get('phase_id') or not phase_id or warning['phase_id'] == phase_id) and
        (not warning.get('cell_id') or warning['cell_id'] in allowed_cells) and
        (not warning.get('cell_ids') or bool(set(warning['cell_ids']) & allowed_cells))]
    for row, warning in enumerate(selected_warnings, 2):
        _text(warnings, row, 1, warning['code']); _text(warnings, row, 2, warning['message'])
    for week in data['weeks']:
        values = [('节假日', ', '.join(week['holiday_days'])), ('调休补班日', ', '.join(week['makeup_workdays']))]
        for name, text in values:
            if text:
                row = warnings.max_row + 1; _text(warnings, row, 1, name); _text(warnings, row, 2, text)
    warnings.column_dimensions['A'].width = 28; warnings.column_dimensions['B'].width = 95
    source = workbook.create_sheet('原日程来源')
    fields = [('id','原记录ID'),('task_id','原工作ID'),('work_date','原日期'),('assignee_id','原负责人ID'),
              ('planned_quantity','原计划量'),('is_confirmed','原发布状态'),('status','原状态'),('notes','原备注')]
    for col, (_, title) in enumerate(fields, 1):
        _text(source, 1, col, title).font = Font(bold=True)
    seen = set()
    for cell in data['cells']:
        if phase_id and cell['phase_id'] != phase_id:
            continue
        for original in (cell.get('source_metadata') or {}).get('original_entries', []):
            if original['id'] in seen:
                continue
            seen.add(original['id']); row = source.max_row + 1
            for col, (key, _) in enumerate(fields, 1):
                _text(source, row, col, original.get(key))
    source.freeze_panes='A2'
    for col in range(1,len(fields)+1):source.column_dimensions[get_column_letter(col)].width=28
    stream = BytesIO(); workbook.save(stream); return stream.getvalue()
