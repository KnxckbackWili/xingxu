"""Project report forms: 生命周期表 / 质量问题反馈单 / 内部联系单.

面向三个交付物：
- 生命周期表.xlsx：车辆问题履历的扁平表格（导出/导入 → VehicleLifecycleEvent）
- 质量问题反馈单.xlsx：单条质量反馈单据（导出/导入 → Issue）
- 内部联系单.docx：申请物品/服务的 Word 联系单（导出/导入 → ItemApplication）
"""

import hashlib
import io
import re
import zipfile
from datetime import date, datetime
from typing import Optional

from dateutil.parser import parse as parse_datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from sqlalchemy.orm import Session

from app.models import (
    Project, Vehicle, VehicleUnit, VehicleLifecycleEvent, Issue, ItemApplication, Attachment,
)

_ISSUE_TYPES = ("fault", "repair", "maintenance", "inspection", "other")

_QUALITY_TYPE_CN = {
    "hardware": "硬件", "software": "软件", "environment": "环境",
    "process": "流程", "safety": "安全", "other": "其他",
}
_QUALITY_TYPE_EN = {value: key for key, value in _QUALITY_TYPE_CN.items()}


def _beautify(ws, title, widths=None):
    """整表美化：标题行 + 表头 + 细边框 + 交替行色 + 列宽（与 import_export._beautify 同风格）。"""
    ncols = max(ws.max_column, 1)
    thin = Side(style="thin", color="FFD9E2F0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.insert_rows(1)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=title)
    c.font = Font(bold=True, size=13, color="FFFFFF")
    c.fill = PatternFill(start_color="FF1677FF", end_color="FF1677FF", fill_type="solid")
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    for cell in ws[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="FF1F4E79", end_color="FF1F4E79", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    ws.row_dimensions[2].height = 20

    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.border = border
            if cell.row % 2 == 1:
                cell.fill = PatternFill(start_color="FFF2F6FC", end_color="FFF2F6FC", fill_type="solid")

    ws.freeze_panes = "A3"
    if widths:
        for col, w in widths.items():
            ws.column_dimensions[col].width = w


def _parse_loose_date(value) -> Optional[date]:
    """解析生命周期表/质量反馈单中的宽松日期：2023/5/4、2023-05-21 00:00:00、datetime 等。"""
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    try:
        return parse_datetime(text, fuzzy=False).date()
    except (ValueError, OverflowError):
        return None


def _split_multiline(value) -> list[str]:
    """把一个单元格里多行/多问题拆成若干条。

    附件原表常以换行或连续空格分隔同一格内的多条内容（中文不含连续空格，安全）。
    """
    if value is None:
        return []
    text = str(value).replace("\r", "\n")
    return [line.strip() for line in re.split(r"\n+| {2,}", text) if line.strip()]


# ===== 生命周期表 =====

LIFECYCLE_HEADERS = ["序号", "项目名称", "车辆型号", "车辆编号", "车号", "出现过的问题", "问题描述", "解决方法", "记录人", "记录时间", "问题分类"]
LIFECYCLE_WIDTHS = {"A": 6, "B": 20, "C": 12, "D": 14, "E": 8, "F": 26, "G": 40, "H": 30, "I": 10, "J": 14, "K": 12}


def lifecycle_to_excel(db: Session, project_id: str, *, events=None, title=None) -> bytes:
    """导出项目车辆问题履历为「AGV生命周期表」格式。"""
    p = db.get(Project, project_id)
    if p is None:
        raise ValueError("项目不存在")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(LIFECYCLE_HEADERS)

    units = {u.id: u for u in db.query(VehicleUnit).filter_by(project_id=project_id).all()}
    vehicles = {v.id: v for v in db.query(Vehicle).filter_by(project_id=project_id).all()}
    if events is None: events = (db.query(VehicleLifecycleEvent)
              .filter(VehicleLifecycleEvent.project_id == project_id,
                      VehicleLifecycleEvent.event_type.in_(_ISSUE_TYPES))
              .order_by(VehicleLifecycleEvent.event_date, VehicleLifecycleEvent.created_at)
              .all())
    seq = 0
    for event in events:
        unit = units.get(event.vehicle_unit_id)
        vehicle = vehicles.get(unit.vehicle_id) if unit else None
        seq += 1
        extra = event.extra or {}
        ws.append([
            seq,
            p.name,
            vehicle.model_name if vehicle else (unit.model_name if unit else "") or extra.get("model_name", ""),
            unit.asset_code if unit else "",
            (unit.fleet_no if unit else "") or extra.get("fleet_no", ""),
            event.title,
            event.description or "",
            event.action_taken or "",
            event.handler or "",
            event.event_date.isoformat() if event.event_date else "",
            extra.get("category", ""),
        ])
    _beautify(ws, title or "AGV生命周期表", widths=LIFECYCLE_WIDTHS)
    for row in ws:
        for cell in row:
            if isinstance(cell.value,str):cell.data_type='s'
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def lifecycle_from_excel(db: Session, project_id: str, data: bytes, operator_id: str) -> dict:
    """导入「AGV生命周期表」到项目车辆问题履历。

    逐行解析；同格多行按「出现过的问题/问题描述/解决方法/记录人/记录时间/问题分类」对齐拆成多条事件。
    车辆按「车辆型号」（model_name）关联车型档案、按车辆编号（asset_code）建档/匹配；车号写入车辆档案。
    同一 项目+车辆+事件日期+标题 的事件写入 import_key，重复导入整体幂等（不会翻倍）。
    错版表格（无表头）与无法解析的事件日期明确报错，不再静默使用当天日期。
    """
    p = db.get(Project, project_id)
    if p is None:
        raise ValueError("项目不存在")
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:
        raise ValueError(f"Excel 解析失败：{exc}") from exc
    ws = wb.active
    # 兼容两种版式：本系统导出的「标题行+表头第2行」与外部直接给的「表头第1行」——按表头内容定位
    header_row = None
    for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        vals = [str(c or "").strip() for c in row[:11]]
        if "车辆编号" in vals and "车号" in vals and "出现过的问题" in vals:
            header_row = idx
            break
    if header_row is None:
        raise ValueError("未找到生命周期表表头（车辆编号/车号/出现过的问题），请使用系统导出的 AGV生命周期表")
    rows = list(ws.iter_rows(min_row=header_row + 1, values_only=True))

    vehicle_by_model = {}
    for v in db.query(Vehicle).filter_by(project_id=project_id).all():
        model = (v.model_name or "").strip()
        if model and model not in vehicle_by_model:
            vehicle_by_model[model] = v
    unit_by_code = {u.asset_code: u for u in db.query(VehicleUnit).filter_by(project_id=project_id).all()}
    # Manual browser records have no import_key. Match their natural identity too,
    # otherwise exporting and reimporting a resolved issue recreates it as open.
    existing_keys = set()
    for event in db.query(VehicleLifecycleEvent).filter(
            VehicleLifecycleEvent.project_id == project_id,
            VehicleLifecycleEvent.event_type.in_(_ISSUE_TYPES)).all():
        if event.import_key:
            existing_keys.add(event.import_key)
        existing_keys.add(hashlib.sha1(
            f"{project_id}|{event.vehicle_unit_id}|{event.event_date.isoformat()}|{event.title.strip()}".encode("utf-8")
        ).hexdigest())
    created_units = 0
    created_events = 0
    skipped_events = 0
    notes = []

    for row_idx, row in enumerate(rows, start=header_row + 1):
        if not row or len(row) < 4 or not str(row[3] or "").strip():
            continue
        # 列序：序号/项目名称/车辆型号/车辆编号/车号/出现过的问题/问题描述/解决方法/记录人/记录时间/问题分类
        model_name = str(row[2] or "").strip()
        asset_code = str(row[3] or "").strip()
        fleet_no = str(row[4] or "").strip()
        if not asset_code:
            continue

        unit = unit_by_code.get(asset_code)
        if unit is None:
            vehicle = vehicle_by_model.get(model_name)
            unit = VehicleUnit(
                project_id=project_id, vehicle_id=vehicle.id if vehicle else None,
                asset_code=asset_code, fleet_no=fleet_no, model_name=model_name,
                current_stage="planned", current_status="normal",
            )
            db.add(unit)
            db.flush()
            unit_by_code[asset_code] = unit
            created_units += 1
        else:
            # 导入的车辆型号持久化到档案：未匹配车型档案时供单车 API/列表展示；车号空时补录
            if model_name and not unit.model_name:
                unit.model_name = model_name
            if fleet_no and not unit.fleet_no:
                unit.fleet_no = fleet_no

        titles = _split_multiline(row[5])
        descriptions = _split_multiline(row[6])
        actions = _split_multiline(row[7])
        handlers = _split_multiline(row[8])
        dates = _split_multiline(row[9])
        categories = _split_multiline(row[10])
        if not titles:
            continue
        # 以「出现过的问题」为准定条数：附件中部分行的分类/日期列存在多余空格段，
        # 若按各列最大段数会造出伪事件，故统一按标题段数对齐。
        count = len(titles)
        for i in range(count):
            title = titles[i] if i < len(titles) else titles[0]
            if not title:
                continue
            date_text = dates[i] if i < len(dates) else None
            event_date = _parse_loose_date(date_text)
            if event_date is None:
                raise ValueError(
                    f"第 {row_idx} 行（车辆编号 {asset_code}，问题「{title}」）记录时间缺失或无法解析"
                    f"：{str(date_text or '').strip() or '（空）'}，请填写为 2026-08-12 或 2026/8/12 格式"
                )
            import_key = hashlib.sha1(
                f"{project_id}|{unit.id}|{event_date.isoformat()}|{title}".encode("utf-8")
            ).hexdigest()
            if import_key in existing_keys:
                skipped_events += 1
                continue
            existing_keys.add(import_key)
            db.add(VehicleLifecycleEvent(
                vehicle_unit_id=unit.id, project_id=project_id, event_type="fault",
                event_date=event_date,
                title=title,
                description=descriptions[i] if i < len(descriptions) else "",
                action_taken=actions[i] if i < len(actions) else "",
                handler=handlers[i] if i < len(handlers) else "",
                extra={"category": categories[i] if i < len(categories) else "",
                       "fleet_no": fleet_no, "model_name": model_name},
                import_key=import_key,
                operator_id=operator_id,
            ))
            created_events += 1

    db.flush()
    if created_units:
        notes.append(f"新建车辆档案 {created_units} 台")
    if skipped_events:
        notes.append(f"跳过重复事件 {skipped_events} 条")
    if not created_units and not skipped_events:
        notes.append("按车辆编号匹配到现有档案")
    return {
        "message": f"生命周期表导入成功：新增现场车辆问题 {created_events} 条（{'、'.join(notes) or '无新增'}）",
        "units": created_units,
        "events": created_events,
        "skipped": skipped_events,
    }


# ===== 质量问题反馈单 =====

def quality_form_to_excel(db: Session, project_id: str, issue_id: str) -> bytes:
    """把一条项目问题导出为与原版一致的「产品质量问题反馈单据」表格。

    严格对照原始附件 质量问题反馈单.xlsx 的版式，重点核对：
    表头（黑体 20、仅下边框）、合并单元格（A4:B4/C4:F4/G4:H4、A5:B7/C5:F7/G5:H7、
    A8:C8/D8:F8/G8:H8、A9:C11/D9:F11/G9:H11、A12:D12/E12:F15/G12:H12、
    A13:D15/G13:H15、A16:H16、A17:H20、A21:H21、图片区两行四块）、
    字段位置、列宽行高（A–H=10.5，行高自适应）、
    字体（标签宋体 12、值宋体 11、图片区块标题楷体 12）、
    边框（内容区全部细边框，A21 分隔空行无边框）、
    对齐（标题 A1 水平居中垂直默认、图片区 A23/E23/A40/E40 中心/中心）、
    打印设置（页边距 0.7/0.75/0.3 与原版一致）。
    """
    issue = db.get(Issue, issue_id)
    if issue is None or issue.project_id != project_id:
        raise ValueError("问题不存在")
    p = db.get(Project, project_id)
    extra = issue.extra or {}
    vehicle_no = str(extra.get("vehicle_no", "")).strip() or (issue.id[:8] if issue.id else "Sheet1")

    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    f_title = Font(name="黑体", size=20)
    f_label = Font(name="宋体", size=12)
    f_value = Font(name="宋体", size=11)
    f_image = Font(name="楷体", size=12)
    al_title = Alignment(horizontal="center")                                          # 标题：水平居中、垂直默认（原版未设置 vertical）
    al_center = Alignment(horizontal="center", vertical="center")
    al_left = Alignment(horizontal="left", vertical="center", indent=1)                # 原版标签/图片标题均为 indent=1
    al_desc = Alignment(horizontal="center", vertical="center", wrap_text=True)

    wb = Workbook()
    ws = wb.active
    ws.title = vehicle_no[:31]

    # 合并单元格（与原版一一对应）
    ws.merge_cells("A1:H3")
    ws.merge_cells("A4:B4"); ws.merge_cells("C4:F4"); ws.merge_cells("G4:H4")
    ws.merge_cells("A5:B7"); ws.merge_cells("C5:F7"); ws.merge_cells("G5:H7")
    ws.merge_cells("A8:C8"); ws.merge_cells("D8:F8"); ws.merge_cells("G8:H8")
    ws.merge_cells("A9:C11"); ws.merge_cells("D9:F11"); ws.merge_cells("G9:H11")
    ws.merge_cells("A12:D12"); ws.merge_cells("E12:F15"); ws.merge_cells("G12:H12")
    ws.merge_cells("A13:D15"); ws.merge_cells("G13:H15")
    ws.merge_cells("A16:H16")
    ws.merge_cells("A17:H20")
    ws.merge_cells("A21:H21")
    ws.merge_cells("A22:D22"); ws.merge_cells("E22:H22")
    ws.merge_cells("A23:D38"); ws.merge_cells("E23:H38")
    ws.merge_cells("A39:D39"); ws.merge_cells("E39:H39")
    ws.merge_cells("A40:D55"); ws.merge_cells("E40:H55")

    # 字段内容
    ws["A1"] = "产品质量问题反馈单据"
    ws["A4"] = "问题类型 Type"; ws["C4"] = "项目 Project"; ws["G4"] = "反馈人 Name"
    ws["A5"] = _QUALITY_TYPE_CN.get(issue.issue_type or "", issue.issue_type or "")
    ws["C5"] = extra.get("project_name") or (p.name if p else "")
    ws["G5"] = extra.get("reporter_name", "")
    ws["A8"] = "车型 Model"; ws["D8"] = "车辆编号 No."; ws["G8"] = "问题来源 From"
    ws["A9"] = extra.get("vehicle_model", "")
    ws["D9"] = vehicle_no
    ws["G9"] = extra.get("source", "")
    ws["A12"] = "问题部件 Parts"; ws["G12"] = "提交日期 Date"
    ws["A13"] = extra.get("part", "")
    feedback_date = _parse_loose_date(extra.get("feedback_date", ""))
    if feedback_date is None and issue.created_at:
        feedback_date = issue.created_at.date()
    ws["G13"] = feedback_date
    ws["A16"] = "问题简述（或现象）Description"
    ws["A17"] = issue.description or ""
    ws["A22"] = "问题图片（一）"; ws["E22"] = "问题图片（二）"
    ws["A39"] = "问题图片（三）"; ws["E39"] = "问题图片（四）"
    # 图片区保持空白（与原版一致，留待打印后手填照片）

    # 样式：标题黑体 20、仅下边框、水平居中（垂直默认，原版未设置）
    ws["A1"].font = f_title
    ws["A1"].alignment = al_title
    ws["A1"].border = Border(bottom=thin)

    # 标签：宋体 12 左对齐细边框；值：宋体 11 居中细边框
    for cell, font, align in (
        (ws["A4"], f_label, al_left), (ws["C4"], f_label, al_left), (ws["G4"], f_label, al_left),
        (ws["A8"], f_label, al_left), (ws["D8"], f_label, al_left), (ws["G8"], f_label, al_left),
        (ws["A12"], f_label, al_left), (ws["G12"], f_label, al_left), (ws["A16"], f_label, al_left),
    ):
        cell.font = font; cell.alignment = align; cell.border = border
    for cell, font, align in (
        (ws["A5"], f_value, al_center), (ws["C5"], f_value, al_center), (ws["G5"], f_value, al_center),
        (ws["A9"], f_value, al_center), (ws["D9"], f_value, al_center), (ws["G9"], f_value, al_center),
        (ws["A13"], f_value, al_center), (ws["G13"], f_value, al_center),
    ):
        cell.font = font; cell.alignment = align; cell.border = border
    ws["A17"].font = f_value
    ws["A17"].alignment = al_desc
    ws["A17"].border = border
    for cell in (ws["A22"], ws["E22"], ws["A39"], ws["E39"]):
        cell.font = f_image
        cell.alignment = al_left
        cell.border = border
    # 空合并区（E12:F15 中部、A23/E23/A40/E40 图片区）也要细边框；A21 分隔空行无边框（与原版一致）
    for cell in (ws["E12"], ws["A23"], ws["E23"], ws["A40"], ws["E40"]):
        cell.border = border
    # 空图片区对齐：原版 A23/E23/A40/E40 为中心/中心
    for cell in (ws["A23"], ws["E23"], ws["A40"], ws["E40"]):
        cell.alignment = al_center

    # 行高（与原版一致：标题合并区 2-3 行 14，A21 分隔空行 5，其余自适应）
    ws.row_dimensions[2].height = 14
    ws.row_dimensions[3].height = 14
    ws.row_dimensions[21].height = 5

    # 列宽（与原版一致：A–H 全部 10.5）
    for col in ("A", "B", "C", "D", "E", "F", "G", "H"):
        ws.column_dimensions[col].width = 10.5

    # 打印设置（与原版页边距一致）
    ws.page_margins.left = 0.7
    ws.page_margins.right = 0.7
    ws.page_margins.top = 0.75
    ws.page_margins.bottom = 0.75
    ws.page_margins.header = 0.3
    ws.page_margins.footer = 0.3

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def quality_form_from_excel(db: Session, project_id: str, data: bytes, operator_id: str) -> dict:
    """导入「产品质量问题反馈单据」为一条项目问题（Issue）。

    版式为位置固定：第 5 行是 问题类型/项目/反馈人，第 9 行 车型/车辆编号/问题来源，
    第 13 行 问题部件/提交日期，第 17 行开始是问题简述。
    先校验版式（首行标题 + 固定标签）与内容，错版/空表单明确报错，避免导入出空问题。
    """
    p = db.get(Project, project_id)
    if p is None:
        raise ValueError("项目不存在")
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:
        raise ValueError(f"Excel 解析失败：{exc}") from exc
    ws = wb.active

    def cell_value(row: int, col: int) -> str:
        value = ws.cell(row=row, column=col).value
        return str(value).strip() if value is not None else ""

    # 错版表格校验：首行标题与固定标签必须匹配「产品质量问题反馈单据」版式
    title_cell = str(ws.cell(row=1, column=1).value or "").strip()
    labels_ok = all(label in cell_value(r, c) for r, c, label in (
        (4, 1, "问题类型"), (8, 1, "车型"), (12, 1, "问题部件"), (16, 1, "问题简述")))
    if "质量问题反馈" not in title_cell or not labels_ok:
        raise ValueError("不是质量问题反馈单版式：首行标题/问题类型/车型/问题部件/问题简述 缺失，请使用系统导出的反馈单")

    issue_type_cn = cell_value(5, 1)
    project_name = cell_value(5, 3)
    reporter_name = cell_value(5, 7)
    vehicle_model = cell_value(9, 1)
    vehicle_no = cell_value(9, 4)
    source = cell_value(9, 7)
    part = cell_value(13, 1)
    feedback_date = cell_value(13, 7)
    description_parts = []
    for row in range(17, min(ws.max_row + 1, 22)):
        value = ws.cell(row=row, column=1).value
        if value not in (None, ""):
            description_parts.append(str(value).strip())
    description = "\n".join(description_parts).strip()

    if not (description or part or vehicle_no):
        raise ValueError("质量问题反馈单内容为空：问题简述、问题部件与车辆编号均未填写")

    issue_type = _QUALITY_TYPE_EN.get(issue_type_cn, issue_type_cn.lower() if issue_type_cn in {"硬件", "软件"} else "other")
    if issue_type_cn == "硬件":
        issue_type = "hardware"
    elif issue_type_cn == "软件":
        issue_type = "software"
    title = description[:60] or f"质量反馈：{part or vehicle_no or '未填写部件'}"
    if not feedback_date:
        raise ValueError("质量问题反馈单缺少提交日期，请填写后再导入")
    parsed_date = _parse_loose_date(feedback_date)
    if parsed_date is None:
        raise ValueError(f"提交日期无法解析：「{feedback_date}」，请填写为 2026-08-12 或 2026/8/12 格式")
    extra = {
        "project_name": project_name or p.name,
        "reporter_name": reporter_name,
        "vehicle_model": vehicle_model,
        "vehicle_no": vehicle_no,
        "source": source,
        "part": part,
        "feedback_date": parsed_date.isoformat() if parsed_date else "",
    }
    issue = Issue(
        project_id=project_id, title=title, description=description,
        issue_type=issue_type, severity="medium", status="open",
        nature='quality',
        reporter_id=operator_id, extra=extra,
    )
    db.add(issue)
    db.flush()
    from app.services.document_triggers import issue_event
    issue_event(db,p,issue,operator_id)
    return {"message": f"质量问题反馈单已导入为问题「{title}」", "id": issue.id}


# ===== 内部联系单（物品申请） =====

def item_application_to_docx(db: Session, project_id: str, app_id: str) -> bytes:
    """把一条物品申请记录生成「内部联系单」Word 文档。"""
    from docx import Document
    from docx.shared import Pt,RGBColor

    p = db.get(Project, project_id)
    app = db.get(ItemApplication, app_id)
    if app is None or app.project_id != project_id:
        raise ValueError("联系单不存在")

    doc = Document()
    # Keep the existing contact form fields and order, with enough room for
    # ordinary multi-line requests and the signature block on the same page.
    doc.styles['Normal'].font.size=Pt(11)
    doc.styles['Normal'].paragraph_format.space_after=Pt(4)
    doc.styles['Title'].font.color.rgb=RGBColor(0,0,0)
    for style in doc.styles:
        for border in style.element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
    title = doc.add_paragraph(style='Title')
    title.paragraph_format.space_after=Pt(12)
    title.paragraph_format.keep_with_next=True
    title.alignment = 1  # center
    run = title.add_run(f"内部联系单——{p.name if p else ''}")
    run.bold = True
    run.font.size = Pt(16)

    fields = [
        ("申请事项", app.title or ""),
        ("申请物品/服务", app.item_name or ""),
        ("需求", app.requirement or ""),
        ("原因", app.reason or ""),
        ("图号", app.drawing_no or ""),
        ("合同号", app.contract_no or ""),
        ("金额", app.amount or ""),
        ("申请人", app.applicant or ""),
        ("审批人", app.approver or ""),
    ]
    for label, value in fields:
        para = doc.add_paragraph()
        para.add_run(f"{label}：").bold = True
        para.add_run(str(value))

    para = doc.add_paragraph()
    para.add_run(f"申请部门：{app.department or ''}").bold = True
    if app.apply_date:
        doc.add_paragraph(f"申请日期：{app.apply_date.isoformat()}")
    signature=doc.add_paragraph("项目技术部")
    signature.paragraph_format.space_before=Pt(10)
    signature.paragraph_format.keep_with_next=True
    doc.add_paragraph("日期：")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


def item_application_from_docx(db: Session, project_id: str, data: bytes, operator_id: str) -> dict:
    """解析「内部联系单」Word 文档为一条物品申请记录。

    按段落文本匹配字段标签（申请事项/申请物品服务/需求/原因/图号/合同号/金额/申请人/审批人/申请部门/申请日期），
    与 item_application_to_docx 生成版式一一对应，支持系统生成 → 回灌；
    对无完整标签的旧版本文档，回退到「需要…，」/「共计…元」/独立日期段解析。
    """
    p = db.get(Project, project_id)
    if p is None:
        raise ValueError("项目不存在")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            xml = zf.read("word/document.xml").decode("utf-8")
    except Exception as exc:
        raise ValueError(f"Word 解析失败：{exc}") from exc

    xml = re.sub(r"<w:tab[^>]*/>", "\t", xml)
    xml = re.sub(r"<w:br[^>]*/>", "\n", xml)
    paragraphs = re.findall(r"<w:p[ >].*?</w:p>", xml, re.S)
    texts = []
    for para in paragraphs:
        pieces = re.findall(r"<w:t[^>]*>([^<]*)</w:t>", para)
        line = "".join(pieces).strip()
        if line:
            texts.append(line)

    def pick(label: str) -> str:
        for line in texts:
            match = re.match(rf"^{re.escape(label)}\s*[:：]\s*(.*)$", line)
            if match:
                return match.group(1).strip()
        return ""

    requirement = pick("需求")
    reason = pick("原因")
    drawing_no = pick("图号")
    contract_no = pick("合同号")
    applicant = pick("申请人")
    approver = pick("审批人")
    department = pick("申请部门") or pick("部门")
    title = pick("申请事项")
    item_name = pick("申请物品/服务")
    amount = pick("金额")

    project_name = ""
    apply_date = _parse_loose_date(pick("申请日期"))
    standalone: list[str] = []
    for line in texts:
        m = re.match(r"^内部联系单[\s\-—–:：]*(.*)$", line)
        if m and not project_name:
            project_name = m.group(1).strip()
            continue
        if re.match(r"^(需求|原因|图号|合同号|申请人|审批人|申请部门|部门|申请事项|申请物品/服务|金额|申请日期|日期)\s*[:：]", line):
            continue
        m = re.match(r"^(\d{4}[-/.]\d{1,2}[-/.]\d{1,2})$", line)
        if m:
            if apply_date is None:
                apply_date = _parse_loose_date(m.group(1))
            continue
        standalone.append(line)
    if not department and standalone:
        # 独立段落（如「项目技术部」）作为申请部门
        department = standalone[0]

    # 回退链：标签缺失时按需求文本推断物品与金额（兼容旧版/手工版式）
    if not item_name:
        item_match = re.search(r"需要\s*([^，。]+?)\s*，", requirement)
        if item_match:
            # 去掉「需要」后常见的动作前缀：需要使用吊车 → 吊车
            candidate = re.sub(r"^(使用|调用|租用|采购|申请|用)+", "", item_match.group(1)).strip()
            item_name = candidate
    if not amount:
        amount_match = re.search(r"共计\s*([\d.]+)\s*元|([\d.]+)\s*元", requirement)
        if amount_match:
            amount = f"{amount_match.group(1) or amount_match.group(2)}元"
    if not item_name:
        item_name = requirement[:40]
    if not title:
        title = requirement or (standalone[0] if standalone else "")

    app = ItemApplication(
        project_id=project_id, title=title[:300], item_name=item_name,
        requirement=requirement, reason=reason, drawing_no=drawing_no,
        contract_no=contract_no, amount=amount, applicant=applicant,
        approver=approver, department=department, apply_date=apply_date,
        status="pending", created_by=operator_id,
    )
    db.add(app)
    db.flush()
    return {
        "message": f"内部联系单已导入为物品申请「{title or requirement[:40]}」",
        "id": app.id,
        "title": app.title,
        "project_name": project_name or p.name,
    }
